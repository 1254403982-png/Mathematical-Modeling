from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import shutil
import time
from dataclasses import asdict, dataclass, replace
from datetime import date, datetime, time as datetime_time
from decimal import Decimal, localcontext
from pathlib import Path
from typing import Any, Sequence

try:
    import pulp
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError as exc:  # pragma: no cover - only used for a clearer CLI error
    raise SystemExit(
        "缺少运行依赖。请执行：python -m pip install pulp openpyxl"
    ) from exc


N_PERIODS = 144
DELTA_T_HOURS = 1.0 / 6.0
ETA_CHARGE = 0.90
ETA_DISCHARGE = 0.90
SOC_MIN_KWH = 1200.0
SOC_MAX_KWH = 10800.0
INITIAL_SOC_KWH = 6000.0
TERMINAL_SOC_KWH = 6000.0
MAX_ACTION_KWH = 5000.0 / 6.0
COST_RELATIVE_TOLERANCE = 1.0e-6
CHECK_TOLERANCE = 1.0e-6
ACTION_TOLERANCE = 1.0e-7
POLISH_ACTIVE_TOLERANCE = 1.0e-4
POLISH_BOUND_TOLERANCE = 1.0e-4
POLISH_SOC_ANCHOR_TOLERANCE = 1.0e-3
POLISH_DECIMAL_PRECISION = 50

INPUT_HEADERS = ["时间", "电价", "小区负载", "光伏发电预测功率"]
RESULT_SHEET_PURCHASE = "计划购电量"
RESULT_SHEET_STORAGE = "充放电量"
RESULT_SHEET_NAMES = [RESULT_SHEET_PURCHASE, RESULT_SHEET_STORAGE]
# 表1六个指定区间，按物理时间理解：
# 区间(10:00, 10:10]的右端点为附件1标签"10:10"（Excel第62行），即模型位置t=61。
TABLE1_PERIODS = [
    (61, "10:00-10:10"),
    (73, "12:00-12:10"),
    (85, "14:00-14:10"),
    (97, "16:00-16:10"),
    (109, "18:00-18:10"),
    (121, "20:00-20:10"),
]
# 旧理解下这些区间在官方模板中的位置，仅供报告中对照。
OLD_TABLE1_TEMPLATE_POSITIONS = [60, 72, 84, 96, 108, 120]
FOUR_HOUR_LABELS = [
    "0:00-4:00",
    "4:00-8:00",
    "8:00-12:00",
    "12:00-16:00",
    "16:00-20:00",
    "20:00-24:00",
]


@dataclass(frozen=True)
class DayData:
    source_time_labels: list[str]
    price: list[float]
    load_kw: list[float]
    pv_kw: list[float]
    load_kwh: list[float]
    pv_kwh: list[float]


@dataclass
class ModelBundle:
    problem: pulp.LpProblem
    purchase_cost: pulp.LpAffineExpression
    total_storage_action: pulp.LpAffineExpression
    grid: list[pulp.LpVariable]
    charge: list[pulp.LpVariable]
    discharge: list[pulp.LpVariable]
    curtailment: list[pulp.LpVariable]
    soc: list[pulp.LpVariable]
    mode: list[pulp.LpVariable]


@dataclass(frozen=True)
class StageReport:
    stage: str
    status_code: int
    status: str
    solution_status_code: int
    solution_status: str
    objective_value: float
    purchase_cost_yuan: float
    runtime_seconds: float
    solver_wallclock_seconds: float | None
    mip_gap: float | None
    enumerated_nodes: int | None
    solver_version: str | None
    log_file: str


@dataclass(frozen=True)
class Solution:
    grid: list[float]
    charge: list[float]
    discharge: list[float]
    curtailment: list[float]
    soc: list[float]
    mode: list[int]


def parse_args() -> argparse.Namespace:
    project_root = Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser(
        description="求解 C 题问题1第一问的两阶段混合整数线性规划。"
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=project_root / "data" / "附件" / "附件1.xlsx",
        help="附件1.xlsx 路径",
    )
    parser.add_argument(
        "--template",
        type=Path,
        default=project_root / "data" / "附件" / "附件5" / "result1.xlsx",
        help="官方 result1.xlsx 模板路径",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=project_root / "output",
        help="结果输出目录",
    )
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def display_time_label(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if isinstance(value, datetime_time):
        return value.strftime("%H:%M")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def checked_number(
    value: Any,
    *,
    field: str,
    excel_row: int,
    require_nonnegative: bool,
) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(
            f"附件1第 {excel_row} 行的“{field}”不是数值：{value!r}"
        )
    number = float(value)
    if not math.isfinite(number):
        raise ValueError(
            f"附件1第 {excel_row} 行的“{field}”不是有限数值：{value!r}"
        )
    if require_nonnegative and number < 0.0:
        raise ValueError(
            f"附件1第 {excel_row} 行的“{field}”为负数：{number}"
        )
    return number


def read_day_data(path: Path) -> DayData:
    if not path.is_file():
        raise FileNotFoundError(f"找不到输入文件：{path}")

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if len(workbook.sheetnames) != 1:
            raise ValueError(
                f"附件1应只有一个工作表，实际为：{workbook.sheetnames}"
            )
        worksheet = workbook[workbook.sheetnames[0]]
        headers = [worksheet.cell(1, column).value for column in range(1, 5)]
        if headers != INPUT_HEADERS:
            raise ValueError(
                f"附件1表头不符合预期：期望 {INPUT_HEADERS}，实际 {headers}"
            )
        if worksheet.max_row != N_PERIODS + 1:
            raise ValueError(
                f"附件1应有 {N_PERIODS} 个数据行，实际为 {worksheet.max_row - 1}"
            )

        source_time_labels: list[str] = []
        price: list[float] = []
        load_kw: list[float] = []
        pv_kw: list[float] = []
        for excel_row in range(2, N_PERIODS + 2):
            raw_time = worksheet.cell(excel_row, 1).value
            if raw_time is None:
                raise ValueError(f"附件1第 {excel_row} 行的时间为空")
            source_time_labels.append(display_time_label(raw_time))
            price.append(
                checked_number(
                    worksheet.cell(excel_row, 2).value,
                    field="电价",
                    excel_row=excel_row,
                    require_nonnegative=False,
                )
            )
            load_kw.append(
                checked_number(
                    worksheet.cell(excel_row, 3).value,
                    field="小区负载",
                    excel_row=excel_row,
                    require_nonnegative=True,
                )
            )
            pv_kw.append(
                checked_number(
                    worksheet.cell(excel_row, 4).value,
                    field="光伏发电预测功率",
                    excel_row=excel_row,
                    require_nonnegative=True,
                )
            )
    finally:
        workbook.close()

    if not all(
        len(values) == N_PERIODS
        for values in (source_time_labels, price, load_kw, pv_kw)
    ):
        raise AssertionError("内部错误：读取后的数组长度不是144")

    return DayData(
        source_time_labels=source_time_labels,
        price=price,
        load_kw=load_kw,
        pv_kw=pv_kw,
        load_kwh=[value * DELTA_T_HOURS for value in load_kw],
        pv_kwh=[value * DELTA_T_HOURS for value in pv_kw],
    )


def read_result_labels(template_path: Path) -> list[str]:
    if not template_path.is_file():
        raise FileNotFoundError(f"找不到结果模板：{template_path}")

    workbook = load_workbook(template_path, data_only=False)
    try:
        if workbook.sheetnames != RESULT_SHEET_NAMES:
            raise ValueError(
                "result1.xlsx 工作表结构不符合预期："
                f"期望 {RESULT_SHEET_NAMES}，实际 {workbook.sheetnames}"
            )
        purchase_sheet = workbook[RESULT_SHEET_PURCHASE]
        storage_sheet = workbook[RESULT_SHEET_STORAGE]
        if purchase_sheet.max_row != N_PERIODS + 1 or purchase_sheet.max_column != 2:
            raise ValueError(
                "“计划购电量”工作表应为145行2列，实际为"
                f"{purchase_sheet.max_row}行{purchase_sheet.max_column}列"
            )
        if storage_sheet.max_row != 7 or storage_sheet.max_column != 5:
            raise ValueError(
                "“充放电量”工作表应为7行5列，实际为"
                f"{storage_sheet.max_row}行{storage_sheet.max_column}列"
            )
        if purchase_sheet.cell(1, 1).value != "时间段" or purchase_sheet.cell(1, 2).value != "购电量":
            raise ValueError("“计划购电量”工作表表头不符合官方模板")
        expected_storage_headers = ["时间段", "充电量", "放电量", "时刻", "储电量"]
        actual_storage_headers = [storage_sheet.cell(1, column).value for column in range(1, 6)]
        if actual_storage_headers != expected_storage_headers:
            raise ValueError(
                f"“充放电量”工作表表头不符合预期：{actual_storage_headers}"
            )
        actual_four_hour_labels = [storage_sheet.cell(row, 1).value for row in range(2, 8)]
        if actual_four_hour_labels != FOUR_HOUR_LABELS:
            raise ValueError(
                f"模板六个4小时时段不符合预期：{actual_four_hour_labels}"
            )
        result_labels = [
            str(purchase_sheet.cell(row, 1).value)
            for row in range(2, N_PERIODS + 2)
        ]
        if any(label == "None" for label in result_labels):
            raise ValueError("模板“计划购电量”工作表存在空时间标签")
        return result_labels
    finally:
        workbook.close()


def physical_interval_labels() -> list[str]:
    """模型位置t的物理区间标签。

    附件1时间标签是区间右端点：位置t对应物理区间((t-1)*10, t*10]分钟。
    t=1为0:00-0:10，t=144为23:50-00:00+1（右端点24:00按附件1写法记为00:00+1）。
    """
    labels: list[str] = []
    for t in range(1, N_PERIODS + 1):
        start_minutes = (t - 1) * 10
        end_minutes = t * 10
        start = f"{start_minutes // 60}:{start_minutes % 60:02d}"
        if end_minutes == 1440:
            end = "00:00+1"
        else:
            end = f"{end_minutes // 60}:{end_minutes % 60:02d}"
        labels.append(f"{start}-{end}")
    return labels


def build_model(data: DayData) -> ModelBundle:
    problem = pulp.LpProblem("Q1_Day_Ahead_Energy_Scheduling", pulp.LpMinimize)
    indices = range(N_PERIODS)

    grid = [pulp.LpVariable(f"g_{i + 1:03d}", lowBound=0.0) for i in indices]
    charge = [pulp.LpVariable(f"c_{i + 1:03d}", lowBound=0.0) for i in indices]
    discharge = [pulp.LpVariable(f"d_{i + 1:03d}", lowBound=0.0) for i in indices]
    curtailment = [
        pulp.LpVariable(f"q_{i + 1:03d}", lowBound=0.0, upBound=data.pv_kwh[i])
        for i in indices
    ]
    soc = [
        pulp.LpVariable(
            f"E_{i + 1:03d}", lowBound=SOC_MIN_KWH, upBound=SOC_MAX_KWH
        )
        for i in indices
    ]
    mode = [pulp.LpVariable(f"z_{i + 1:03d}", cat=pulp.LpBinary) for i in indices]

    for i in indices:
        problem += (
            grid[i]
            + data.pv_kwh[i]
            - curtailment[i]
            + discharge[i]
            == data.load_kwh[i] + charge[i],
            f"energy_balance_{i + 1:03d}",
        )
        previous_soc = INITIAL_SOC_KWH if i == 0 else soc[i - 1]
        problem += (
            soc[i]
            == previous_soc
            + ETA_CHARGE * charge[i]
            - discharge[i] / ETA_DISCHARGE,
            f"soc_transition_{i + 1:03d}",
        )
        problem += (
            charge[i] <= MAX_ACTION_KWH * mode[i],
            f"charge_mode_limit_{i + 1:03d}",
        )
        problem += (
            discharge[i] <= MAX_ACTION_KWH * (1.0 - mode[i]),
            f"discharge_mode_limit_{i + 1:03d}",
        )

    problem += soc[-1] == TERMINAL_SOC_KWH, "terminal_soc"
    purchase_cost = pulp.lpSum(data.price[i] * grid[i] for i in indices)
    total_storage_action = pulp.lpSum(charge[i] + discharge[i] for i in indices)
    problem += purchase_cost, "stage_1_minimum_purchase_cost"

    return ModelBundle(
        problem=problem,
        purchase_cost=purchase_cost,
        total_storage_action=total_storage_action,
        grid=grid,
        charge=charge,
        discharge=discharge,
        curtailment=curtailment,
        soc=soc,
        mode=mode,
    )


def expression_value(expression: pulp.LpAffineExpression, label: str) -> float:
    value = pulp.value(expression)
    if value is None or not math.isfinite(float(value)):
        raise RuntimeError(f"无法读取{label}的有限数值")
    return float(value)


def variable_values(variables: Sequence[pulp.LpVariable], label: str) -> list[float]:
    values: list[float] = []
    for i, variable in enumerate(variables, start=1):
        value = variable.value()
        if value is None or not math.isfinite(float(value)):
            raise RuntimeError(f"无法读取 {label}[{i}] 的有限数值")
        values.append(float(value))
    return values


def parse_cbc_log(log_path: Path, status: str) -> dict[str, Any]:
    text = log_path.read_text(encoding="utf-8", errors="replace")

    def first_float(pattern: str) -> float | None:
        match = re.search(pattern, text, flags=re.IGNORECASE | re.MULTILINE)
        return float(match.group(1)) if match else None

    def first_int(pattern: str) -> int | None:
        match = re.search(pattern, text, flags=re.IGNORECASE | re.MULTILINE)
        return int(match.group(1)) if match else None

    version_match = re.search(r"^Version:\s*([^\r\n]+)", text, flags=re.MULTILINE)
    reported_gap = first_float(r"^Gap:\s*([-+0-9.eE]+)")
    mip_gap = 0.0 if status == "Optimal" else reported_gap
    return {
        "solver_version": version_match.group(1).strip() if version_match else None,
        "solver_wallclock_seconds": first_float(
            r"^Time \(Wallclock seconds\):\s*([-+0-9.eE]+)"
        ),
        "best_objective": first_float(
            r"Search completed - best objective\s+([-+0-9.eE]+)"
        ),
        "enumerated_nodes": first_int(r"Enumerated nodes:\s*(\d+)"),
        "mip_gap": mip_gap,
    }


def solve_stage(
    bundle: ModelBundle,
    *,
    stage: str,
    objective: pulp.LpAffineExpression,
    log_path: Path,
    warm_start: bool,
) -> StageReport:
    bundle.problem.setObjective(objective)
    solver = pulp.PULP_CBC_CMD(
        msg=False,
        gapRel=0.0,
        threads=1,
        warmStart=warm_start,
        logPath=str(log_path),
        options=["randomSeed 1"],
    )
    started = time.perf_counter()
    status_code = bundle.problem.solve(solver)
    runtime_seconds = time.perf_counter() - started
    status = pulp.LpStatus.get(status_code, f"Unknown({status_code})")
    solution_status_code = bundle.problem.sol_status
    solution_status = pulp.LpSolution.get(
        solution_status_code, f"Unknown({solution_status_code})"
    )
    if status != "Optimal":
        raise RuntimeError(
            f"{stage}未达到最优状态：status={status}, solution_status={solution_status}。"
            f"详见 {log_path}"
        )

    log_values = parse_cbc_log(log_path, status)
    extracted_objective = expression_value(objective, f"{stage}目标值")
    solver_objective = log_values["best_objective"]
    if solver_objective is None:
        solver_objective = extracted_objective
    purchase_cost = expression_value(bundle.purchase_cost, f"{stage}购电费")
    if objective is bundle.purchase_cost:
        purchase_cost = solver_objective
    return StageReport(
        stage=stage,
        status_code=status_code,
        status=status,
        solution_status_code=solution_status_code,
        solution_status=solution_status,
        objective_value=solver_objective,
        purchase_cost_yuan=purchase_cost,
        runtime_seconds=runtime_seconds,
        solver_wallclock_seconds=log_values["solver_wallclock_seconds"],
        mip_gap=log_values["mip_gap"],
        enumerated_nodes=log_values["enumerated_nodes"],
        solver_version=log_values["solver_version"],
        log_file=log_path.name,
    )


def extract_solution(bundle: ModelBundle) -> Solution:
    mode_values = variable_values(bundle.mode, "z")
    rounded_mode: list[int] = []
    for i, value in enumerate(mode_values, start=1):
        rounded = int(round(value))
        if rounded not in (0, 1) or abs(value - rounded) > CHECK_TOLERANCE:
            raise RuntimeError(f"z[{i}]不是有效二元解：{value}")
        rounded_mode.append(rounded)

    return Solution(
        grid=variable_values(bundle.grid, "g"),
        charge=variable_values(bundle.charge, "c"),
        discharge=variable_values(bundle.discharge, "d"),
        curtailment=variable_values(bundle.curtailment, "q"),
        soc=variable_values(bundle.soc, "E"),
        mode=rounded_mode,
    )


def solve_decimal_linear_system(
    coefficients: list[list[Decimal]], right_hand_side: list[Decimal]
) -> list[Decimal]:
    size = len(coefficients)
    if size == 0 or any(len(row) != size for row in coefficients):
        raise RuntimeError("高精度活动集恢复方程组不是非空方阵")
    if len(right_hand_side) != size:
        raise RuntimeError("高精度活动集恢复方程组左右维数不一致")

    augmented = [
        coefficients[row][:] + [right_hand_side[row]] for row in range(size)
    ]
    singular_tolerance = Decimal("1e-35")
    for column in range(size):
        pivot_row = max(
            range(column, size), key=lambda row: abs(augmented[row][column])
        )
        if abs(augmented[pivot_row][column]) <= singular_tolerance:
            raise RuntimeError("高精度活动集恢复方程组奇异，无法可靠恢复")
        augmented[column], augmented[pivot_row] = (
            augmented[pivot_row],
            augmented[column],
        )
        pivot = augmented[column][column]
        for entry in range(column, size + 1):
            augmented[column][entry] /= pivot
        for row in range(size):
            if row == column:
                continue
            factor = augmented[row][column]
            if factor == 0:
                continue
            for entry in range(column, size + 1):
                augmented[row][entry] -= factor * augmented[column][entry]

    return [augmented[row][-1] for row in range(size)]


def polish_solution(
    data: DayData,
    raw_solution: Solution,
    *,
    stage1_cost: float,
    stage2_solver_objective: float,
) -> tuple[Solution, dict[str, Any]]:
    """Recover CBC's optimal active-set solution beyond its text-file precision."""

    with localcontext() as decimal_context:
        decimal_context.prec = POLISH_DECIMAL_PRECISION
        zero = Decimal(0)
        eta_charge = Decimal(str(ETA_CHARGE))
        eta_discharge = Decimal(str(ETA_DISCHARGE))
        max_action = Decimal(5000) / Decimal(6)
        active_tolerance = Decimal(str(POLISH_ACTIVE_TOLERANCE))
        bound_tolerance = Decimal(str(POLISH_BOUND_TOLERANCE))
        anchor_tolerance = Decimal(str(POLISH_SOC_ANCHOR_TOLERANCE))

        price = [Decimal(str(value)) for value in data.price]
        load = [Decimal(str(value)) for value in data.load_kwh]
        pv = [Decimal(str(value)) for value in data.pv_kwh]
        raw_grid = [Decimal(str(value)) for value in raw_solution.grid]
        raw_charge = [Decimal(str(value)) for value in raw_solution.charge]
        raw_discharge = [Decimal(str(value)) for value in raw_solution.discharge]
        raw_curtailment = [
            Decimal(str(value)) for value in raw_solution.curtailment
        ]
        raw_soc = [Decimal(str(value)) for value in raw_solution.soc]

        if max(abs(value) for value in raw_curtailment) > active_tolerance:
            raise RuntimeError(
                "CBC解存在非零弃光，当前高精度活动集恢复不能安全假定q_t=0"
            )

        charge = [zero for _ in range(N_PERIODS)]
        discharge = [zero for _ in range(N_PERIODS)]
        unknowns: list[tuple[str, int]] = []
        for i in range(N_PERIODS):
            if (
                raw_charge[i] > active_tolerance
                and raw_discharge[i] > active_tolerance
            ):
                raise RuntimeError(f"CBC解在t={i + 1}同时充放电，不能恢复活动集")

            if raw_charge[i] <= active_tolerance:
                charge[i] = zero
            elif abs(raw_charge[i] - max_action) <= bound_tolerance:
                charge[i] = max_action
            elif raw_grid[i] <= active_tolerance:
                inferred_charge = max(zero, pv[i] - load[i])
                if abs(inferred_charge - raw_charge[i]) > Decimal("0.01"):
                    raise RuntimeError(
                        f"t={i + 1}的零购电充电活动集与能量平衡不一致"
                    )
                charge[i] = inferred_charge
            else:
                unknowns.append(("c", i))

            if raw_discharge[i] <= active_tolerance:
                discharge[i] = zero
            elif abs(raw_discharge[i] - max_action) <= bound_tolerance:
                discharge[i] = max_action
            elif raw_grid[i] <= active_tolerance:
                inferred_discharge = max(zero, load[i] - pv[i])
                if abs(inferred_discharge - raw_discharge[i]) > Decimal("0.01"):
                    raise RuntimeError(
                        f"t={i + 1}的零购电放电活动集与能量平衡不一致"
                    )
                discharge[i] = inferred_discharge
            else:
                unknowns.append(("d", i))

        state_classes: list[str | None] = []
        soc_min = Decimal(str(SOC_MIN_KWH))
        soc_max = Decimal(str(SOC_MAX_KWH))
        for value in raw_soc:
            if abs(value - soc_max) <= anchor_tolerance:
                state_classes.append("max")
            elif abs(value - soc_min) <= anchor_tolerance:
                state_classes.append("min")
            else:
                state_classes.append(None)

        anchors: list[tuple[int, Decimal, str]] = []
        previous_class: str | None = None
        for i, state_class in enumerate(state_classes):
            if state_class is not None and state_class != previous_class:
                target = soc_max if state_class == "max" else soc_min
                anchors.append((i, target, state_class))
            previous_class = state_class
        terminal_target = Decimal(str(TERMINAL_SOC_KWH))
        if anchors and anchors[-1][0] == N_PERIODS - 1:
            if anchors[-1][1] != terminal_target:
                raise RuntimeError("日末SOC活动边界与终端SOC要求冲突")
        else:
            anchors.append((N_PERIODS - 1, terminal_target, "terminal"))

        equation_rows: list[list[Decimal]] = []
        equation_rhs: list[Decimal] = []
        initial_soc = Decimal(str(INITIAL_SOC_KWH))
        inverse_discharge_efficiency = Decimal(1) / eta_discharge
        for anchor_index, target, _ in anchors:
            known_state = initial_soc
            for i in range(anchor_index + 1):
                known_state += (
                    eta_charge * charge[i]
                    - inverse_discharge_efficiency * discharge[i]
                )
            row = [zero for _ in unknowns]
            for column, (kind, variable_index) in enumerate(unknowns):
                if variable_index <= anchor_index:
                    row[column] = (
                        eta_charge
                        if kind == "c"
                        else -inverse_discharge_efficiency
                    )
            equation_rows.append(row)
            equation_rhs.append(target - known_state)

        cost_ceiling = (
            Decimal(1) + Decimal(str(COST_RELATIVE_TOLERANCE))
        ) * Decimal(str(stage1_cost))
        known_cost = zero
        cost_row = [zero for _ in unknowns]
        for i in range(N_PERIODS):
            if raw_grid[i] > active_tolerance:
                known_cost += price[i] * (
                    load[i] - pv[i] + charge[i] - discharge[i]
                )
        for column, (kind, variable_index) in enumerate(unknowns):
            cost_row[column] = price[variable_index] * (
                Decimal(1) if kind == "c" else Decimal(-1)
            )
        equation_rows.append(cost_row)
        equation_rhs.append(cost_ceiling - known_cost)

        if len(equation_rows) != len(unknowns):
            raise RuntimeError(
                "高精度活动集恢复维数不匹配："
                f"{len(unknowns)}个未知量、{len(equation_rows)}个独立条件"
            )
        recovered_values = solve_decimal_linear_system(
            equation_rows, equation_rhs
        )
        recovered_variables: list[dict[str, Any]] = []
        for (kind, i), value in zip(unknowns, recovered_values, strict=True):
            raw_value = raw_charge[i] if kind == "c" else raw_discharge[i]
            if value < -Decimal("1e-20") or value > max_action + Decimal("1e-20"):
                raise RuntimeError(
                    f"恢复后的{kind}[{i + 1}]越界：{value}"
                )
            if abs(value - raw_value) > Decimal("0.01"):
                raise RuntimeError(
                    f"恢复后的{kind}[{i + 1}]偏离CBC活动集过大："
                    f"raw={raw_value}, recovered={value}"
                )
            if kind == "c":
                charge[i] = value
            else:
                discharge[i] = value
            recovered_variables.append(
                {
                    "variable": f"{kind}_{i + 1:03d}",
                    "raw_value": float(raw_value),
                    "recovered_value": float(value),
                    "adjustment": float(value - raw_value),
                }
            )

        grid: list[Decimal] = []
        curtailment: list[Decimal] = []
        soc: list[Decimal] = []
        current_soc = initial_soc
        for i in range(N_PERIODS):
            current_soc += (
                eta_charge * charge[i]
                - inverse_discharge_efficiency * discharge[i]
            )
            soc.append(current_soc)
            net_grid = load[i] + charge[i] - pv[i] - discharge[i]
            grid.append(max(zero, net_grid))
            curtailment.append(max(zero, -net_grid))
            if curtailment[-1] > pv[i] + Decimal("1e-20"):
                raise RuntimeError(f"恢复后的q[{i + 1}]超过可用光伏电量")
            if raw_grid[i] > active_tolerance and grid[-1] <= 0:
                raise RuntimeError(f"恢复使t={i + 1}离开正购电活动集")
            if raw_grid[i] <= active_tolerance and grid[-1] > Decimal("1e-20"):
                raise RuntimeError(f"恢复使t={i + 1}离开零购电活动集")

        recovered_cost = sum(
            price[i] * grid[i] for i in range(N_PERIODS)
        )
        recovered_objective = sum(charge) + sum(discharge)
        cost_residual = recovered_cost - cost_ceiling
        if abs(cost_residual) > Decimal("1e-20"):
            raise RuntimeError(
                f"高精度活动集恢复后的成本等式残差过大：{cost_residual}"
            )
        if min(soc) < soc_min - Decimal("1e-20") or max(soc) > soc_max + Decimal("1e-20"):
            raise RuntimeError("高精度活动集恢复后的SOC越界")

        mode = list(raw_solution.mode)
        for i in range(N_PERIODS):
            if charge[i] > Decimal("1e-20"):
                mode[i] = 1
            elif discharge[i] > Decimal("1e-20"):
                mode[i] = 0

        polished = Solution(
            grid=[float(value) for value in grid],
            charge=[float(value) for value in charge],
            discharge=[float(value) for value in discharge],
            curtailment=[float(value) for value in curtailment],
            soc=[float(value) for value in soc],
            mode=mode,
        )
        diagnostics: dict[str, Any] = {
            "method": "CBC最优活动集的50位Decimal线性方程恢复",
            "reason": "CBC文本解约8位小数，直接导出不足以通过1e-6残差验收",
            "decimal_precision": POLISH_DECIMAL_PRECISION,
            "active_value_tolerance": POLISH_ACTIVE_TOLERANCE,
            "soc_anchor_tolerance": POLISH_SOC_ANCHOR_TOLERANCE,
            "unknown_count": len(unknowns),
            "equation_count": len(equation_rows),
            "soc_anchors": [
                {"t": i + 1, "target_soc_kwh": float(target), "type": kind}
                for i, target, kind in anchors
            ],
            "recovered_variables": recovered_variables,
            "raw_text_solution_purchase_cost_yuan": float(
                sum(price[i] * raw_grid[i] for i in range(N_PERIODS))
            ),
            "recovered_purchase_cost_yuan": float(recovered_cost),
            "recovered_cost_ceiling_residual_yuan": float(cost_residual),
            "raw_text_solution_storage_action_kwh": float(
                sum(raw_charge) + sum(raw_discharge)
            ),
            "recovered_storage_action_kwh": float(recovered_objective),
            "cbc_internal_objective_kwh": stage2_solver_objective,
            "recovered_minus_cbc_objective_kwh": float(
                recovered_objective - Decimal(str(stage2_solver_objective))
            ),
            "max_charge_adjustment_kwh": float(
                max(
                    abs(charge[i] - raw_charge[i]) for i in range(N_PERIODS)
                )
            ),
            "max_discharge_adjustment_kwh": float(
                max(
                    abs(discharge[i] - raw_discharge[i])
                    for i in range(N_PERIODS)
                )
            ),
            "max_grid_adjustment_kwh": float(
                max(abs(grid[i] - raw_grid[i]) for i in range(N_PERIODS))
            ),
            "max_soc_adjustment_kwh": float(
                max(abs(soc[i] - raw_soc[i]) for i in range(N_PERIODS))
            ),
        }
        return polished, diagnostics


def validate_solution(
    data: DayData,
    solution: Solution,
    *,
    stage1_cost: float,
) -> dict[str, Any]:
    balance_residuals: list[float] = []
    soc_residuals: list[float] = []
    charge_mode_violations: list[float] = []
    discharge_mode_violations: list[float] = []

    for i in range(N_PERIODS):
        balance_residuals.append(
            solution.grid[i]
            + data.pv_kwh[i]
            - solution.curtailment[i]
            + solution.discharge[i]
            - data.load_kwh[i]
            - solution.charge[i]
        )
        previous_soc = INITIAL_SOC_KWH if i == 0 else solution.soc[i - 1]
        soc_residuals.append(
            solution.soc[i]
            - previous_soc
            - ETA_CHARGE * solution.charge[i]
            + solution.discharge[i] / ETA_DISCHARGE
        )
        charge_mode_violations.append(
            max(0.0, solution.charge[i] - MAX_ACTION_KWH * solution.mode[i])
        )
        discharge_mode_violations.append(
            max(
                0.0,
                solution.discharge[i]
                - MAX_ACTION_KWH * (1.0 - solution.mode[i]),
            )
        )

    total_grid = sum(solution.grid)
    total_charge = sum(solution.charge)
    total_discharge = sum(solution.discharge)
    total_curtailment = sum(solution.curtailment)
    total_load = sum(data.load_kwh)
    total_pv = sum(data.pv_kwh)
    final_cost = sum(
        data.price[i] * solution.grid[i] for i in range(N_PERIODS)
    )
    efficiency_identity_residual = total_discharge - (
        ETA_CHARGE * ETA_DISCHARGE * total_charge
    )
    daily_energy_identity_residual = total_grid - (
        total_load
        - total_pv
        + total_curtailment
        + (1.0 - ETA_CHARGE * ETA_DISCHARGE) * total_charge
    )
    simultaneous_periods = [
        i + 1
        for i in range(N_PERIODS)
        if solution.charge[i] > ACTION_TOLERANCE
        and solution.discharge[i] > ACTION_TOLERANCE
    ]
    cost_ceiling = (1.0 + COST_RELATIVE_TOLERANCE) * stage1_cost

    metrics: dict[str, Any] = {
        "continuous_check_tolerance": CHECK_TOLERANCE,
        "storage_action_zero_tolerance": ACTION_TOLERANCE,
        "max_energy_balance_residual_kwh": max(
            abs(value) for value in balance_residuals
        ),
        "max_soc_transition_residual_kwh": max(
            abs(value) for value in soc_residuals
        ),
        "min_soc_kwh": min(solution.soc),
        "max_soc_kwh": max(solution.soc),
        "max_charge_kwh": max(solution.charge),
        "max_discharge_kwh": max(solution.discharge),
        "max_charge_mode_violation_kwh": max(charge_mode_violations),
        "max_discharge_mode_violation_kwh": max(discharge_mode_violations),
        "min_grid_purchase_kwh": min(solution.grid),
        "min_curtailment_kwh": min(solution.curtailment),
        "max_curtailment_upper_bound_violation_kwh": max(
            max(0.0, solution.curtailment[i] - data.pv_kwh[i])
            for i in range(N_PERIODS)
        ),
        "simultaneous_charge_discharge_periods": simultaneous_periods,
        "initial_soc_kwh": INITIAL_SOC_KWH,
        "terminal_soc_kwh": solution.soc[-1],
        "terminal_soc_residual_kwh": solution.soc[-1] - TERMINAL_SOC_KWH,
        "daily_efficiency_identity_residual_kwh": efficiency_identity_residual,
        "daily_energy_identity_residual_kwh": daily_energy_identity_residual,
        "stage1_optimal_cost_yuan": stage1_cost,
        "stage2_purchase_cost_yuan": final_cost,
        "stage2_cost_ceiling_yuan": cost_ceiling,
        "stage2_cost_ceiling_violation_yuan": max(0.0, final_cost - cost_ceiling),
        "total_load_kwh": total_load,
        "total_pv_kwh": total_pv,
        "total_grid_purchase_kwh": total_grid,
        "total_charge_kwh": total_charge,
        "total_discharge_kwh": total_discharge,
        "total_storage_action_kwh": total_charge + total_discharge,
        "total_curtailment_kwh": total_curtailment,
    }
    checks = {
        "energy_balance": metrics["max_energy_balance_residual_kwh"]
        <= CHECK_TOLERANCE,
        "soc_transition": metrics["max_soc_transition_residual_kwh"]
        <= CHECK_TOLERANCE,
        "soc_bounds": (
            metrics["min_soc_kwh"] >= SOC_MIN_KWH - CHECK_TOLERANCE
            and metrics["max_soc_kwh"] <= SOC_MAX_KWH + CHECK_TOLERANCE
        ),
        "charge_limit": metrics["max_charge_kwh"]
        <= MAX_ACTION_KWH + CHECK_TOLERANCE,
        "discharge_limit": metrics["max_discharge_kwh"]
        <= MAX_ACTION_KWH + CHECK_TOLERANCE,
        "charge_discharge_mutex": not simultaneous_periods,
        "mode_limits": (
            metrics["max_charge_mode_violation_kwh"] <= CHECK_TOLERANCE
            and metrics["max_discharge_mode_violation_kwh"] <= CHECK_TOLERANCE
        ),
        "grid_and_curtailment_bounds": (
            metrics["min_grid_purchase_kwh"] >= -CHECK_TOLERANCE
            and metrics["min_curtailment_kwh"] >= -CHECK_TOLERANCE
            and metrics["max_curtailment_upper_bound_violation_kwh"]
            <= CHECK_TOLERANCE
        ),
        "terminal_soc": abs(metrics["terminal_soc_residual_kwh"])
        <= CHECK_TOLERANCE,
        "daily_efficiency_identity": abs(efficiency_identity_residual)
        <= CHECK_TOLERANCE,
        "daily_energy_identity": abs(daily_energy_identity_residual)
        <= CHECK_TOLERANCE,
        "stage2_cost_ceiling": metrics["stage2_cost_ceiling_violation_yuan"]
        <= CHECK_TOLERANCE,
    }
    metrics["checks"] = checks
    metrics["all_checks_passed"] = all(checks.values())
    return metrics


def aggregate_four_hour(solution: Solution) -> list[dict[str, Any]]:
    aggregates: list[dict[str, Any]] = []
    for block, label in enumerate(FOUR_HOUR_LABELS):
        start = block * 24
        stop = start + 24
        aggregates.append(
            {
                "period": label,
                "start_t": start + 1,
                "end_t": stop,
                "charge_kwh": sum(solution.charge[start:stop]),
                "discharge_kwh": sum(solution.discharge[start:stop]),
            }
        )
    return aggregates


def compare_template_structure(source_path: Path, output_path: Path) -> None:
    source = load_workbook(source_path, data_only=False)
    output = load_workbook(output_path, data_only=False)
    editable_cells = {
        RESULT_SHEET_PURCHASE: {
            f"B{row}" for row in range(2, N_PERIODS + 2)
        },
        RESULT_SHEET_STORAGE: {
            *(f"B{row}" for row in range(2, 8)),
            *(f"C{row}" for row in range(2, 8)),
            "E2",
            "E3",
        },
    }
    try:
        if source.sheetnames != output.sheetnames:
            raise RuntimeError("输出模板的工作表名称或顺序发生变化")
        for sheet_name in source.sheetnames:
            source_sheet = source[sheet_name]
            output_sheet = output[sheet_name]
            if (
                source_sheet.max_row,
                source_sheet.max_column,
            ) != (
                output_sheet.max_row,
                output_sheet.max_column,
            ):
                raise RuntimeError(f"输出模板工作表“{sheet_name}”的尺寸发生变化")
            if list(source_sheet.merged_cells.ranges) != list(
                output_sheet.merged_cells.ranges
            ):
                raise RuntimeError(f"输出模板工作表“{sheet_name}”的合并单元格发生变化")
            for row in range(1, source_sheet.max_row + 1):
                if source_sheet.row_dimensions[row].height != output_sheet.row_dimensions[row].height:
                    raise RuntimeError(f"输出模板工作表“{sheet_name}”第{row}行高度发生变化")
                for column in range(1, source_sheet.max_column + 1):
                    source_cell = source_sheet.cell(row, column)
                    output_cell = output_sheet.cell(row, column)
                    if source_cell.style_id != output_cell.style_id:
                        raise RuntimeError(
                            f"输出模板单元格“{sheet_name}!{source_cell.coordinate}”样式发生变化"
                        )
                    if (
                        source_cell.coordinate not in editable_cells[sheet_name]
                        and source_cell.value != output_cell.value
                    ):
                        raise RuntimeError(
                            f"输出模板非目标单元格“{sheet_name}!{source_cell.coordinate}”内容发生变化"
                        )
            for column_letter, source_dimension in source_sheet.column_dimensions.items():
                output_dimension = output_sheet.column_dimensions[column_letter]
                if source_dimension.width != output_dimension.width:
                    raise RuntimeError(
                        f"输出模板工作表“{sheet_name}”第{column_letter}列宽发生变化"
                    )
    finally:
        source.close()
        output.close()


def write_result_copy(
    template_path: Path,
    output_path: Path,
    solution: Solution,
    four_hour: list[dict[str, Any]],
) -> None:
    if template_path.resolve() == output_path.resolve():
        raise ValueError("输出路径不得与官方模板路径相同")

    source_hash_before = sha256_file(template_path)
    shutil.copy2(template_path, output_path)
    workbook = load_workbook(output_path, data_only=False)
    try:
        purchase_sheet = workbook[RESULT_SHEET_PURCHASE]
        storage_sheet = workbook[RESULT_SHEET_STORAGE]
        for i, value in enumerate(solution.grid, start=2):
            purchase_sheet.cell(i, 2).value = value
        for block, aggregate in enumerate(four_hour, start=2):
            storage_sheet.cell(block, 2).value = aggregate["charge_kwh"]
            storage_sheet.cell(block, 3).value = aggregate["discharge_kwh"]
        storage_sheet["E2"] = INITIAL_SOC_KWH
        storage_sheet["E3"] = solution.soc[-1]
        workbook.save(output_path)
    finally:
        workbook.close()

    if sha256_file(template_path) != source_hash_before:
        raise RuntimeError("官方 result1.xlsx 模板在导出过程中被改变")
    compare_template_structure(template_path, output_path)


def write_result1_2(
    path: Path,
    solution: Solution,
    physical_labels: list[str],
) -> None:
    """导出按物理时间区间标注的购电量表。

    与官方模板不同：第一列直接写实际物理区间（首行0:00-0:10，
    末行23:50-00:00+1），第二列为对应位置的g_t，共144个数据行。
    """
    if len(physical_labels) != N_PERIODS:
        raise ValueError("物理区间标签数量不是144")
    workbook = Workbook()
    try:
        worksheet = workbook.active
        worksheet.title = RESULT_SHEET_PURCHASE
        worksheet.append(["时间段", "购电量"])
        for label, value in zip(
            physical_labels, solution.grid, strict=True
        ):
            worksheet.append([label, value])
        style_header(worksheet)
        worksheet.column_dimensions["A"].width = 18
        worksheet.column_dimensions["B"].width = 22
        for cell in worksheet["B"][1:]:
            cell.number_format = "0.0000000000"
        workbook.save(path)
    finally:
        workbook.close()

    check = load_workbook(path, data_only=True, read_only=True)
    try:
        if check.sheetnames != [RESULT_SHEET_PURCHASE]:
            raise RuntimeError(f"result1_2工作表结构异常：{check.sheetnames}")
        rows = list(
            check[RESULT_SHEET_PURCHASE].iter_rows(
                min_row=2, values_only=True
            )
        )
        if len(rows) != N_PERIODS:
            raise RuntimeError(f"result1_2数据行数不是144：{len(rows)}")
        for i, (label, value) in enumerate(rows):
            if label != physical_labels[i]:
                raise RuntimeError(
                    f"result1_2第{i + 1}行标签异常：{label!r} != "
                    f"{physical_labels[i]!r}"
                )
            if value is None or abs(float(value) - solution.grid[i]) > CHECK_TOLERANCE:
                raise RuntimeError(
                    f"result1_2第{i + 1}行购电量与解不一致：{value!r}"
                )
    finally:
        check.close()


def style_header(worksheet: Any) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions


def write_details_workbook(
    path: Path,
    data: DayData,
    result_labels: list[str],
    physical_labels: list[str],
    solution: Solution,
    table1: list[dict[str, Any]],
    four_hour: list[dict[str, Any]],
    stage_reports: list[StageReport],
    validation: dict[str, Any],
    polishing: dict[str, Any],
) -> None:
    workbook = Workbook()
    details = workbook.active
    details.title = "144时段明细"
    details.append(
        [
            "模型位置t",
            "附件1时间标签",
            "result1时间段",
            "物理区间(result1_2)",
            "电价p_t(元/kWh)",
            "负荷功率(kW)",
            "光伏预测功率(kW)",
            "负荷电量L_t(kWh)",
            "光伏电量R_t(kWh)",
            "购电量g_t(kWh)",
            "充电量c_t(kWh)",
            "放电量d_t(kWh)",
            "弃光量q_t(kWh)",
            "时段末储电量E_t(kWh)",
            "充放电状态z_t",
        ]
    )
    for i in range(N_PERIODS):
        details.append(
            [
                i + 1,
                data.source_time_labels[i],
                result_labels[i],
                physical_labels[i],
                data.price[i],
                data.load_kw[i],
                data.pv_kw[i],
                data.load_kwh[i],
                data.pv_kwh[i],
                solution.grid[i],
                solution.charge[i],
                solution.discharge[i],
                solution.curtailment[i],
                solution.soc[i],
                solution.mode[i],
            ]
        )
    style_header(details)
    details.column_dimensions["A"].width = 12
    details.column_dimensions["B"].width = 18
    details.column_dimensions["C"].width = 22
    details.column_dimensions["D"].width = 22
    for column in "EFGHIJKLMNO":
        details.column_dimensions[column].width = 22
        for cell in details[column][1:]:
            cell.number_format = "0.0000000000"
    details.column_dimensions["O"].width = 18

    summary = workbook.create_sheet("表1与汇总")
    summary.append(["表1指定时段", "模型位置t", "计划购电量(kWh)"])
    for item in table1:
        summary.append([item["period"], item["t"], item["grid_purchase_kwh"]])
    summary.append([])
    summary.append(["全天指标", "数值", "单位"])
    summary.append(["全天购电量", validation["total_grid_purchase_kwh"], "kWh"])
    summary.append(["全天购电费（第二阶段最终解）", validation["stage2_purchase_cost_yuan"], "元"])
    summary.append(["第一阶段最优购电费C1*", validation["stage1_optimal_cost_yuan"], "元"])
    summary.append(["全天充电量", validation["total_charge_kwh"], "kWh"])
    summary.append(["全天放电量", validation["total_discharge_kwh"], "kWh"])
    summary.append(["全天充放电动作", validation["total_storage_action_kwh"], "kWh"])
    summary.append(["全天弃光量", validation["total_curtailment_kwh"], "kWh"])
    summary.append([])
    summary.append(["4小时时段", "充电量(kWh)", "放电量(kWh)"])
    for item in four_hour:
        summary.append([item["period"], item["charge_kwh"], item["discharge_kwh"]])
    style_header(summary)
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 24
    summary.column_dimensions["C"].width = 24
    for row in summary.iter_rows(min_row=2):
        for cell in row[1:]:
            if isinstance(cell.value, float):
                cell.number_format = "0.0000000000"

    report_sheet = workbook.create_sheet("求解与验收报告")
    report_sheet.append(["类别", "指标", "数值"])
    for stage_report in stage_reports:
        stage_dict = asdict(stage_report)
        for key, value in stage_dict.items():
            report_sheet.append([stage_report.stage, key, value])
    for key, value in polishing.items():
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False)
        report_sheet.append(["高精度恢复", key, value])
    for key, value in validation.items():
        if key == "checks":
            continue
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False)
        report_sheet.append(["验收指标", key, value])
    for key, value in validation["checks"].items():
        report_sheet.append(["验收结论", key, "通过" if value else "未通过"])
    style_header(report_sheet)
    report_sheet.column_dimensions["A"].width = 22
    report_sheet.column_dimensions["B"].width = 48
    report_sheet.column_dimensions["C"].width = 32
    for cell in report_sheet["C"][1:]:
        if isinstance(cell.value, float):
            cell.number_format = "0.0000000000"

    workbook.save(path)
    workbook.close()


def relative_name(path: Path, base: Path) -> str:
    try:
        return path.resolve().relative_to(base.resolve()).as_posix()
    except ValueError:
        return str(path.resolve())


def main() -> None:
    args = parse_args()
    input_path = args.input.resolve()
    template_path = args.template.resolve()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    result_path = output_dir / "result1.xlsx"
    result1_2_path = output_dir / "result1_2.xlsx"
    details_path = output_dir / "q1_details.xlsx"
    report_path = output_dir / "q1_report.json"
    stage1_log = output_dir / "q1_stage1_cbc.log"
    stage2_log = output_dir / "q1_stage2_cbc.log"

    data = read_day_data(input_path)
    result_labels = read_result_labels(template_path)
    physical_labels = physical_interval_labels()
    for t, period in TABLE1_PERIODS:
        if physical_labels[t - 1] != period:
            raise RuntimeError(
                f"表1指定区间位置{t}的物理标签不一致："
                f"{period!r} != {physical_labels[t - 1]!r}"
            )
    input_hash = sha256_file(input_path)
    template_hash = sha256_file(template_path)

    bundle = build_model(data)
    stage1_report = solve_stage(
        bundle,
        stage="第一阶段：最小购电费",
        objective=bundle.purchase_cost,
        log_path=stage1_log,
        warm_start=False,
    )
    stage1_cost = stage1_report.objective_value
    bundle.problem += (
        bundle.purchase_cost
        <= (1.0 + COST_RELATIVE_TOLERANCE) * stage1_cost,
        "preserve_stage_1_optimal_cost",
    )
    stage2_report = solve_stage(
        bundle,
        stage="第二阶段：最小充放电动作",
        objective=bundle.total_storage_action,
        log_path=stage2_log,
        warm_start=False,
    )

    raw_solution = extract_solution(bundle)
    solution, polishing = polish_solution(
        data,
        raw_solution,
        stage1_cost=stage1_cost,
        stage2_solver_objective=stage2_report.objective_value,
    )
    validation = validate_solution(data, solution, stage1_cost=stage1_cost)
    stage2_report = replace(
        stage2_report,
        purchase_cost_yuan=validation["stage2_purchase_cost_yuan"],
    )
    four_hour = aggregate_four_hour(solution)
    table1 = [
        {
            "period": period,
            "t": t,
            "grid_purchase_kwh": solution.grid[t - 1],
        }
        for t, period in TABLE1_PERIODS
    ]

    write_result_copy(template_path, result_path, solution, four_hour)
    write_result1_2(result1_2_path, solution, physical_labels)
    write_details_workbook(
        details_path,
        data,
        result_labels,
        physical_labels,
        solution,
        table1,
        four_hour,
        [stage1_report, stage2_report],
        validation,
        polishing,
    )

    project_root = Path(__file__).resolve().parent.parent
    report = {
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "model": {
            "periods": N_PERIODS,
            "delta_t_hours": DELTA_T_HOURS,
            "eta_charge": ETA_CHARGE,
            "eta_discharge": ETA_DISCHARGE,
            "soc_min_kwh": SOC_MIN_KWH,
            "soc_max_kwh": SOC_MAX_KWH,
            "initial_soc_kwh": INITIAL_SOC_KWH,
            "terminal_soc_kwh": TERMINAL_SOC_KWH,
            "max_charge_or_discharge_kwh_per_period": MAX_ACTION_KWH,
            "stage2_cost_relative_tolerance": COST_RELATIVE_TOLERANCE,
        },
        "source_files": {
            "input": relative_name(input_path, project_root),
            "input_sha256": input_hash,
            "template": relative_name(template_path, project_root),
            "template_sha256": template_hash,
        },
        "label_mapping": {
            "attachment1_label_semantics": "附件1时间标签为该10分钟区间的右端点",
            "formula": "模型位置t的物理区间为((t-1)*10, t*10]分钟",
            "official_template_label_offset": "官方模板第t行标签比物理区间晚10分钟",
            "result1_2_first_row_label": physical_labels[0],
            "result1_2_last_row_label": physical_labels[-1],
            "table1_interpretation": "按物理时间理解",
            "table1_positions": [
                {
                    "period": period,
                    "t_physical": t,
                    "t_previous_template_anchor": previous_t,
                }
                for (t, period), previous_t in zip(
                    TABLE1_PERIODS,
                    OLD_TABLE1_TEMPLATE_POSITIONS,
                    strict=True,
                )
            ],
        },
        "output_files": {
            "official_template_copy": relative_name(result_path, project_root),
            "official_template_copy_sha256": sha256_file(result_path),
            "result1_2_physical_labels": relative_name(result1_2_path, project_root),
            "result1_2_physical_labels_sha256": sha256_file(result1_2_path),
            "details": relative_name(details_path, project_root),
            "details_sha256": sha256_file(details_path),
            "stage1_solver_log": relative_name(stage1_log, project_root),
            "stage2_solver_log": relative_name(stage2_log, project_root),
        },
        "stage_reports": [asdict(stage1_report), asdict(stage2_report)],
        "table1": table1,
        "four_hour_storage_aggregates": four_hour,
        "numerical_polishing": polishing,
        "acceptance": validation,
        "template_structure_preserved": True,
    }
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )

    print("问题1第一问求解完成")
    print(
        f"第一阶段：{stage1_report.status}，C1* = "
        f"{stage1_report.purchase_cost_yuan:.10f} 元，"
        f"MIP gap = {stage1_report.mip_gap:.3g}"
    )
    print(
        f"第二阶段：{stage2_report.status}，总充放电动作 = "
        f"{stage2_report.objective_value:.10f} kWh，"
        f"最终购电费 = {stage2_report.purchase_cost_yuan:.10f} 元，"
        f"MIP gap = {stage2_report.mip_gap:.3g}"
    )
    print(
        "验收检查："
        + ("全部通过" if validation["all_checks_passed"] else "存在未通过项")
    )
    print(f"官方模板副本：{result_path}")
    print(f"物理区间标签版：{result1_2_path}")
    print(f"144时段明细：{details_path}")
    print(f"求解报告：{report_path}")


if __name__ == "__main__":
    main()
