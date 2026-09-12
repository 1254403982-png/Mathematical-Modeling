"""问题4：波动电价下重算问题2与问题3。

权威规格：order/orderForQ4.md。附件2/4的时间标签均为区间右端点；模型位置
`t=1..144`分别表示物理区间00:00-00:10至23:50-24:00。官方result4宽表的
文字按物理意义填写，因此每个日期行展示“当日t=2..144 + 次日t=1”。

建议运行：
  python q4.py --self-test
  python q4.py --all --max-days 2 --output-dir ../output/q4_smoke
  python q4.py --all --output-dir ../output/q4

完整运行分三类产物：严格因果的预测/场景缓存、两条正式年度策略与敏感性、模板
回填和回读验收。正式工作簿的终值代理已按修复后的Q2/Q3对齐：v=下次凌晨
00:00-05:00窗口的场景均价/eta_c（充电边际）；eta_d*min、eta_d*median、v=0
三档只作敏感性，不按事后费用切换正式方案。
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import shutil
import sys
import time
from copy import copy
from dataclasses import asdict, dataclass
from datetime import date, datetime, time as datetime_time, timedelta
from pathlib import Path
from typing import Any, Iterable, Sequence

import numpy as np
import pulp
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import q2  # noqa: E402
import q3  # noqa: E402

T = q2.T
ETA_CHARGE = q2.ETA_CHARGE
ETA_DISCHARGE = q2.ETA_DISCHARGE
SOC_MIN_KWH = q2.SOC_MIN_KWH
SOC_MAX_KWH = q2.SOC_MAX_KWH
INITIAL_SOC_KWH = q2.INITIAL_SOC_KWH
MAX_ACTION_KWH = q2.MAX_ACTION_KWH
EMERGENCY_PRICE_FACTOR = q2.EMERGENCY_PRICE_FACTOR
MAX_SCENARIOS = q2.MAX_SCENARIOS
CHECK_TOLERANCE = q2.CHECK_TOLERANCE
SIMULTANEOUS_PRODUCT_TOLERANCE = q2.SIMULTANEOUS_PRODUCT_TOLERANCE
ACTION_ZERO_TOLERANCE = q2.ACTION_ZERO_TOLERANCE
FORECAST_START_DATE = q2.FORECAST_START_DATE
FORECAST_END_DATE = q2.FORECAST_END_DATE
BOUNDARY_PLAN_DATE = q2.BOUNDARY_PLAN_DATE
FOUR_HOUR_LABELS = q2.FOUR_HOUR_LABELS
ISSUE_HOURS = q3.ISSUE_HOURS
HOUR_INDEX = q3.HOUR_INDEX
KAPPA = q3.KAPPA

FORMAL_DAYS = 334
PRICE_HISTORY_DAYS = 35
MIN_REGRESSION_SAMPLES = 7
PRICE_EPSILON = 1.0e-4
NET_LOAD_SCALE_KWH = 1.0e5
SHORT_RHO_CANDIDATES = (0.0, 0.5, 0.8, 0.95)
PRICE_MODEL_CANDIDATES = ("ar", "net")
TERMINAL_MODES = ("overnight", "zero", "minimum", "median")
OVERNIGHT_WINDOW_SLOTS = 30  # 与Q2/Q3修复一致：物理00:00-05:00共30段
DP_GRID_STEP_KWH = 10.0
ADJUST_UP_PRICE_FACTOR = 1.5
ADJUST_DOWN_PRICE_FACTOR = 0.5
BRANCHES = ("4-2", "4-3")
BRANCH_INDEX = {name: i for i, name in enumerate(BRANCHES)}
RESULT42_SHEETS = ("计划购电量", "充放电量", "紧急购电量")
RESULT43_SHEETS = ("计划购电量", "调整购电量", "充放电量", "紧急购电量")
PAPER_INTERVALS = (
    (61, "10:00-10:10"),
    (73, "12:00-12:10"),
    (85, "14:00-14:10"),
    (97, "16:00-16:10"),
    (109, "18:00-18:10"),
    (121, "20:00-20:10"),
)
PAPER_EMERGENCY_DATES = (
    date(2025, 3, 20),
    date(2025, 6, 21),
    date(2025, 9, 23),
    date(2025, 12, 21),
)

DATA_PATHS = {
    "annual": Path("data/附件/附件2.xlsx"),
    "pv_forecast": Path("data/附件/附件3.xlsx"),
    "prices": Path("data/附件/附件4.xlsx"),
    "template42": Path("data/附件/附件5/result4-2.xlsx"),
    "template43": Path("data/附件/附件5/result4-3.xlsx"),
}

CACHE_FILES = {
    "scenario_price": "q4_scenario_price.npy",
    "scenario_net": "q4_scenario_net.npy",
    "scenario_count": "q4_scenario_count.npy",
    "point_price": "q4_point_price.npy",
    "point_load": "q4_point_load.npy",
    "point_pv": "q4_point_pv.npy",
    "boundary_price": "q4_boundary_price.npy",
    "boundary_net": "q4_boundary_net.npy",
    "boundary_count": "q4_boundary_count.npy",
    "boundary_point_price": "q4_boundary_point_price.npy",
    "metadata": "q4_prep_summary.json",
}


@dataclass(frozen=True)
class PriceData:
    dates: tuple[date, ...]
    right_end_labels: tuple[str, ...]
    values: np.ndarray
    date_to_index: dict[date, int]
    daily_level: np.ndarray


@dataclass(frozen=True)
class RawResourceEvent:
    issue_date: date
    issue_hour: int
    day_indices: np.ndarray
    t_abs: np.ndarray
    load_forecast: np.ndarray
    pv_historical: np.ndarray
    pv_formal: np.ndarray | None
    next_load_forecast: np.ndarray
    next_pv_historical: np.ndarray


@dataclass(frozen=True)
class ResourceEvent:
    issue_date: date
    issue_hour: int
    day_indices: np.ndarray
    t_abs: np.ndarray
    load_forecast: np.ndarray
    pv_forecast: np.ndarray
    pv_historical: np.ndarray
    pv_formal: np.ndarray | None
    current_net_energy_forecast_kwh: float
    next_net_energy_forecast_kwh: float

    @property
    def issue_datetime(self) -> datetime:
        return datetime.combine(self.issue_date, datetime_time(self.issue_hour))

    @property
    def close_datetime(self) -> datetime:
        return self.issue_datetime + timedelta(hours=24)


@dataclass(frozen=True)
class PriceForecast:
    branch: str
    issue_date: date
    issue_hour: int
    family: str
    short_rho: float
    final: np.ndarray
    base: np.ndarray
    correction: np.ndarray
    shape_window: np.ndarray
    current_level: float
    next_level: float
    intercept: float
    coefficient: float
    sample_count: int
    fit_status: str
    train_dates: tuple[date, ...]
    last_known_actual_price: float | None
    last_known_base_price: float | None


@dataclass(frozen=True)
class ResidualEvent:
    branch: str
    issue_date: date
    issue_hour: int
    close_datetime: datetime
    price_residual: np.ndarray
    load_residual: np.ndarray
    pv_residual: np.ndarray


@dataclass
class WindowModel:
    problem: pulp.LpProblem
    original_objective: pulp.LpAffineExpression
    z: list[pulp.LpVariable]
    charge: list[list[pulp.LpVariable]]
    discharge: list[list[pulp.LpVariable]]
    emergency: list[list[pulp.LpVariable]]
    unused: list[list[pulp.LpVariable]]
    soc: list[list[pulp.LpVariable]]
    delta_plus: list[pulp.LpVariable]
    delta_minus: list[pulp.LpVariable]
    ncur: int


@dataclass(frozen=True)
class WindowResult:
    issue_date: date
    issue_hour: int
    kind: str
    z: np.ndarray
    delta_plus: np.ndarray
    delta_minus: np.ndarray
    objective_value: float
    scenario_count: int
    status: str
    solve_seconds: float
    tiebreak_used: bool
    binary_fallback_used: bool
    max_balance_residual_kwh: float
    max_soc_residual_kwh: float
    max_simultaneous_product_kwh2: float
    max_delta_product_kwh2: float
    min_variable_value_kwh: float
    min_soc_kwh: float
    max_soc_kwh: float
    max_charge_kwh: float
    max_discharge_kwh: float
    highs_max_primal_infeasibility: float
    simplex_iterations: int
    ipm_iterations: int
    mip_gap: float | None


@dataclass(frozen=True)
class RunConfig:
    name: str
    branch: str
    terminal_mode: str
    update_hours: tuple[int, ...]
    adjust_purchases: bool
    refresh_value: bool
    main_output: bool
    description: str


RUN_CONFIGS = {
    "42-main": RunConfig(
        "42-main", "4-2", "overnight", ISSUE_HOURS, False, True, True,
        "4-2正式方案：0:00锁定普通购电，四次刷新储能价值；终值=次日凌晨窗口场景均价/eta_c",
    ),
    "43-main": RunConfig(
        "43-main", "4-3", "overnight", ISSUE_HOURS, True, True, True,
        "4-3正式方案：四次普通购电调整与价值刷新；终值=次日凌晨窗口场景均价/eta_c",
    ),
    "42-terminal-zero": RunConfig(
        "42-terminal-zero", "4-2", "zero", ISSUE_HOURS, False, True, False,
        "4-2终值代理v=0敏感性",
    ),
    "42-terminal-minimum": RunConfig(
        "42-terminal-minimum", "4-2", "minimum", ISSUE_HOURS, False, True, False,
        "4-2终值代理eta_d*min(场景均价)敏感性",
    ),
    "42-terminal-median": RunConfig(
        "42-terminal-median", "4-2", "median", ISSUE_HOURS, False, True, False,
        "4-2终值代理eta_d*median(场景均价)敏感性",
    ),
    "43-terminal-zero": RunConfig(
        "43-terminal-zero", "4-3", "zero", ISSUE_HOURS, True, True, False,
        "4-3终值代理v=0敏感性",
    ),
    "43-terminal-minimum": RunConfig(
        "43-terminal-minimum", "4-3", "minimum", ISSUE_HOURS, True, True, False,
        "4-3终值代理eta_d*min(场景均价)敏感性",
    ),
    "43-terminal-median": RunConfig(
        "43-terminal-median", "4-3", "median", ISSUE_HOURS, True, True, False,
        "4-3终值代理eta_d*median(场景均价)敏感性",
    ),
    "42-no-value-refresh": RunConfig(
        "42-no-value-refresh", "4-2", "overnight", (0,), False, False, False,
        "4-2反事实：0:00后不刷新储能价值",
    ),
    "43-no-adjustment": RunConfig(
        "43-no-adjustment", "4-3", "overnight", ISSUE_HOURS, False, True, False,
        "4-3反事实：普通购电锁定但四次刷新储能价值",
    ),
    "43-no-value-refresh": RunConfig(
        "43-no-value-refresh", "4-3", "overnight", ISSUE_HOURS, True, False, False,
        "4-3反事实：照常滚动普通购电但执行器冻结0:00价值",
    ),
}


@dataclass(frozen=True)
class MainRunArrays:
    g0: np.ndarray
    geff: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    emergency: np.ndarray
    unused: np.ndarray
    soc: np.ndarray
    initial_soc: np.ndarray
    terminal_soc: np.ndarray
    boundary_plan: np.ndarray
    boundary_price_mean: np.ndarray


@dataclass(frozen=True)
class ExportDay:
    plan_grid: np.ndarray
    effective_grid: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    emergency: np.ndarray
    initial_soc_kwh: float
    terminal_soc_kwh: float


def parse_args() -> argparse.Namespace:
    root = Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser(description="求解数学建模C题问题4：波动电价下重算Q2/Q3")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--all", action="store_true", help="准备缓存、运行全部方案并组装正式结果")
    mode.add_argument("--prep", action="store_true", help="只准备预测和场景缓存")
    mode.add_argument("--run-config", choices=tuple(RUN_CONFIGS), help="只运行一个年度配置")
    mode.add_argument("--assemble", action="store_true", help="从已完成配置组装Excel和汇总")
    mode.add_argument("--verify", action="store_true", help="独立回读已生成正式结果")
    mode.add_argument("--self-test", action="store_true", help="运行内置单元级检查")
    parser.add_argument("--root", type=Path, default=root)
    parser.add_argument("--output-dir", type=Path, default=None)
    parser.add_argument("--max-days", type=int, default=None, help="冒烟测试仅运行前N个正式日")
    parser.add_argument("--progress-every", type=int, default=5)
    parser.add_argument("--dp-grid-step", type=float, default=DP_GRID_STEP_KWH)
    parser.add_argument("--resume", action="store_true", help="复用匹配的缓存和已完成配置")
    parser.add_argument("--main-only", action="store_true", help="仅运行两条正式主方案，跳过敏感性")
    return parser.parse_args()


def get_output_dir(args: argparse.Namespace) -> Path:
    path = args.output_dir if args.output_dir is not None else args.root / "output" / "q4"
    path = path.resolve()
    path.mkdir(parents=True, exist_ok=True)
    return path


def input_path(args: argparse.Namespace, key: str) -> Path:
    return (args.root / DATA_PATHS[key]).resolve()


def write_csv(
    path: Path,
    rows: Iterable[dict[str, Any]],
    fieldnames: Sequence[str],
) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames, extrasaction="raise")
        writer.writeheader()
        writer.writerows(rows)


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as file:
        return list(csv.DictReader(file))


def formal_dates(count: int = FORMAL_DAYS) -> tuple[date, ...]:
    return tuple(FORECAST_START_DATE + timedelta(days=i) for i in range(count))


def issue_datetime(day: date, hour: int) -> datetime:
    return datetime.combine(day, datetime_time(hour))


def format_datetime(value: datetime) -> str:
    return value.isoformat(timespec="minutes")


def read_price_data(path: Path) -> PriceData:
    if not path.is_file():
        raise FileNotFoundError(f"找不到附件4：{path}")
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if tuple(workbook.sheetnames) != ("Sheet1",):
            raise ValueError(f"附件4工作表异常：{workbook.sheetnames}")
        sheet = workbook["Sheet1"]
        if sheet.max_row != 366 or sheet.max_column != T + 1:
            raise ValueError(f"附件4维度异常：{sheet.max_row}×{sheet.max_column}")
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        if len(header) != T + 1 or header[0] != "日期\\时间":
            raise ValueError(f"附件4首表头或列数异常：{header[0]!r}, columns={len(header)}")
        labels = tuple(q2.display_time_label(value) for value in header[1:])
        dates: list[date] = []
        values: list[list[float]] = []
        for excel_row, row in enumerate(rows, start=2):
            if len(row) < T + 1:
                raise ValueError(f"附件4第{excel_row}行列数不足")
            day = q2.as_date(row[0], f"附件4第{excel_row}行")
            dates.append(day)
            values.append(
                [
                    q2.checked_float(
                        value,
                        context=f"附件4第{excel_row}行第{column}列电价",
                        nonnegative=False,
                    )
                    for column, value in enumerate(row[1 : T + 1], start=2)
                ]
            )
    finally:
        workbook.close()
    expected = [date(2025, 1, 1) + timedelta(days=i) for i in range(365)]
    if dates != expected:
        raise ValueError("附件4日期不是2025-01-01至2025-12-31连续自然日")
    array = np.asarray(values, dtype=np.float64)
    if array.shape != (365, T) or not np.all(np.isfinite(array)):
        raise ValueError(f"附件4价格数组异常：{array.shape}")
    if float(np.min(array)) <= 0.0:
        raise ValueError(f"附件4含非正价格：min={float(np.min(array))}")
    if labels[0] not in ("00:10", "0:10") or labels[-1] != "0:00+1":
        raise ValueError(f"附件4右端点标签异常：首={labels[0]!r}, 末={labels[-1]!r}")
    return PriceData(
        tuple(dates), labels, array, {day: i for i, day in enumerate(dates)}, np.mean(array, axis=1)
    )


def validate_inputs(annual: q2.AnnualData, prices: PriceData) -> None:
    if annual.dates != prices.dates:
        raise ValueError("附件2与附件4日期集合不一致")
    if annual.right_end_labels != prices.right_end_labels:
        mismatches = [
            (i + 1, annual.right_end_labels[i], prices.right_end_labels[i])
            for i in range(T)
            if annual.right_end_labels[i] != prices.right_end_labels[i]
        ]
        raise ValueError(f"附件2与附件4右端点标签不一致：{mismatches[:5]}")
    for name, values in (("负荷", annual.load_kwh), ("光伏", annual.pv_kwh)):
        if values.shape != (365, T) or not np.all(np.isfinite(values)):
            raise ValueError(f"附件2{name}数组异常：{values.shape}")
        if float(np.min(values)) < 0.0:
            raise ValueError(f"附件2{name}含负值")


def make_q2_raw_event(
    day_idx: int,
    hour: int,
    current: q2.ForecastDraft,
    following: q2.ForecastDraft,
) -> RawResourceEvent | None:
    """构造4-2事件；函数签名和实现均不接收附件3。"""
    if current.load_forecast is None or current.pv_forecast is None:
        return None
    if following.load_forecast is None or following.pv_forecast is None:
        return None
    k = KAPPA[hour]
    if k == 0:
        load = current.load_forecast.copy()
        pv = current.pv_forecast.copy()
    else:
        load = np.concatenate((current.load_forecast[k:], following.load_forecast[:k]))
        pv = np.concatenate((current.pv_forecast[k:], following.pv_forecast[:k]))
    day_indices, t_abs = q3.window_positions(day_idx, hour)
    return RawResourceEvent(
        issue_date=current.target_date,
        issue_hour=hour,
        day_indices=day_indices,
        t_abs=t_abs,
        load_forecast=load,
        pv_historical=pv,
        pv_formal=None,
        next_load_forecast=following.load_forecast.copy(),
        next_pv_historical=following.pv_forecast.copy(),
    )


def make_q3_raw_event(
    day_idx: int,
    hour: int,
    current: q2.ForecastDraft,
    following: q2.ForecastDraft,
    attachment3: q3.Attachment3,
    annual: q2.AnnualData,
) -> RawResourceEvent | None:
    pending = q3.build_event(day_idx, hour, current, following, attachment3, annual)
    if pending is None or following.load_forecast is None or following.pv_forecast is None:
        return None
    return RawResourceEvent(
        issue_date=current.target_date,
        issue_hour=hour,
        day_indices=pending.day,
        t_abs=pending.t_abs,
        load_forecast=pending.load_forecast,
        pv_historical=pending.pv_historical,
        pv_formal=pending.pv_formal,
        next_load_forecast=following.load_forecast.copy(),
        next_pv_historical=following.pv_forecast.copy(),
    )


def build_raw_resource_events(
    annual: q2.AnnualData,
    attachment3: q3.Attachment3,
) -> tuple[
    dict[str, dict[tuple[date, int], RawResourceEvent]],
    q2.ForecastDraft,
    q2.ForecastDraft,
]:
    raw: dict[str, dict[tuple[date, int], RawResourceEvent]] = {branch: {} for branch in BRANCHES}
    history = q2.ForecastHistory.empty()
    for day_idx, day in enumerate(annual.dates):
        current = q2.forecast_one_day(day, history)
        following = q2.forecast_one_day(day + timedelta(days=1), history)
        for hour in ISSUE_HOURS:
            event42 = make_q2_raw_event(day_idx, hour, current, following)
            if event42 is not None:
                raw["4-2"][(day, hour)] = event42
            event43 = make_q3_raw_event(day_idx, hour, current, following, attachment3, annual)
            if event43 is not None:
                raw["4-3"][(day, hour)] = event43
        q2.append_observation(
            history, day, annual.load_kwh[day_idx], annual.pv_kwh[day_idx], current
        )

    boundary = q2.forecast_one_day(BOUNDARY_PLAN_DATE, history)
    boundary_next = q2.forecast_one_day(BOUNDARY_PLAN_DATE + timedelta(days=1), history)
    if any(
        value is None
        for value in (
            boundary.load_forecast,
            boundary.pv_forecast,
            boundary_next.load_forecast,
            boundary_next.pv_forecast,
        )
    ):
        raise RuntimeError("2026-01-01边界历史预测不足")
    day_indices = np.full(T, 365, dtype=np.int64)
    t_abs = np.arange(1, T + 1, dtype=np.int64)
    event42 = RawResourceEvent(
        issue_date=BOUNDARY_PLAN_DATE,
        issue_hour=0,
        day_indices=day_indices,
        t_abs=t_abs,
        load_forecast=boundary.load_forecast.copy(),
        pv_historical=boundary.pv_forecast.copy(),
        pv_formal=None,
        next_load_forecast=boundary_next.load_forecast.copy(),
        next_pv_historical=boundary_next.pv_forecast.copy(),
    )
    raw["4-2"][(BOUNDARY_PLAN_DATE, 0)] = event42
    # 附件3没有2026行；4-3边界按Q3既有规则退回历史光伏预测。
    raw["4-3"][(BOUNDARY_PLAN_DATE, 0)] = event42
    return raw, boundary, boundary_next


def actual_resource_path(
    raw: RawResourceEvent | ResourceEvent,
    annual: q2.AnnualData,
) -> tuple[np.ndarray, np.ndarray] | None:
    if np.any(raw.day_indices < 0) or np.any(raw.day_indices >= len(annual.dates)):
        return None
    return (
        annual.load_kwh[raw.day_indices, raw.t_abs - 1],
        annual.pv_kwh[raw.day_indices, raw.t_abs - 1],
    )


def calibrate_bg_weight(
    raw43: dict[tuple[date, int], RawResourceEvent],
    annual: q2.AnnualData,
) -> dict[str, Any]:
    cutoff = issue_datetime(FORECAST_START_DATE, 0)
    formal_errors: list[np.ndarray] = []
    historical_errors: list[np.ndarray] = []
    used: list[tuple[date, int]] = []
    for key, event in sorted(raw43.items()):
        if event.issue_date.year != 2025 or event.issue_date.month != 1:
            continue
        if event.pv_formal is None or issue_datetime(event.issue_date, event.issue_hour) + timedelta(hours=24) > cutoff:
            continue
        actual = actual_resource_path(event, annual)
        if actual is None:
            continue
        actual_pv = actual[1]
        formal_errors.append(actual_pv - event.pv_formal)
        historical_errors.append(actual_pv - event.pv_historical)
        used.append(key)
    if not used:
        raise RuntimeError("没有在2025-02-01 0:00前完整关闭的Q3组合校准事件")
    e_formal = np.concatenate(formal_errors)
    e_historical = np.concatenate(historical_errors)
    result = q3.bates_granger_weight(e_formal, e_historical)
    weight = float(result["w"])
    combined_errors: list[np.ndarray] = []
    for key in used:
        event = raw43[key]
        actual = actual_resource_path(event, annual)
        assert actual is not None and event.pv_formal is not None
        combined = q3.combine_pv(weight, event.pv_formal, event.pv_historical)
        combined_errors.append(actual[1] - combined)
    return {
        "scope": "events_closed_by_2025-02-01T00:00",
        "event_count": len(used),
        "position_count": int(e_formal.size),
        "event_keys": [f"{day.isoformat()} {hour:02d}:00" for day, hour in used],
        "w": weight,
        "w_raw": float(result["w_raw"]),
        "denominator": float(result["denominator"]),
        "clipped": bool(result["clipped"]),
        "formal": q3.error_triplet(-e_formal),
        "historical": q3.error_triplet(-e_historical),
        "combined": q3.error_triplet(-np.concatenate(combined_errors)),
    }


def finalize_resource_events(
    raw: dict[str, dict[tuple[date, int], RawResourceEvent]],
    annual: q2.AnnualData,
    bg_weight: float,
) -> dict[str, dict[tuple[date, int], ResourceEvent]]:
    finalized: dict[str, dict[tuple[date, int], ResourceEvent]] = {
        branch: {} for branch in BRANCHES
    }
    date_to_index = {day: i for i, day in enumerate(annual.dates)}
    for branch in BRANCHES:
        for key, event in raw[branch].items():
            if branch == "4-3" and event.pv_formal is not None:
                pv = q3.combine_pv(bg_weight, event.pv_formal, event.pv_historical)
            else:
                pv = event.pv_historical.copy()
            k = KAPPA[event.issue_hour]
            ncur = T - k
            if event.issue_date in date_to_index:
                day_idx = date_to_index[event.issue_date]
                observed = float(
                    np.sum(
                        annual.load_kwh[day_idx, :k] - annual.pv_kwh[day_idx, :k]
                    )
                )
            elif k == 0:
                observed = 0.0
            else:
                raise RuntimeError("边界事件不应包含日内发布")
            current_n = observed + float(
                np.sum(event.load_forecast[:ncur] - pv[:ncur])
            )
            next_pv = event.next_pv_historical.copy()
            next_load = event.next_load_forecast.copy()
            if branch == "4-3" and k > 0:
                next_pv[:k] = pv[ncur:]
                next_load[:k] = event.load_forecast[ncur:]
            next_n = float(np.sum(next_load - next_pv))
            finalized[branch][key] = ResourceEvent(
                issue_date=event.issue_date,
                issue_hour=event.issue_hour,
                day_indices=event.day_indices.copy(),
                t_abs=event.t_abs.copy(),
                load_forecast=event.load_forecast.copy(),
                pv_forecast=pv,
                pv_historical=event.pv_historical.copy(),
                pv_formal=None if event.pv_formal is None else event.pv_formal.copy(),
                current_net_energy_forecast_kwh=current_n,
                next_net_energy_forecast_kwh=next_n,
            )
    return finalized


def history_window(prices: PriceData, target: date) -> tuple[date, ...]:
    prior = [day for day in prices.dates if day < target]
    return tuple(prior[-PRICE_HISTORY_DAYS:])


def normalized_shape(prices: PriceData, window: Sequence[date]) -> np.ndarray:
    if not window:
        raise RuntimeError("没有完整历史日，无法构造价格形状")
    curves = np.stack(
        [
            prices.values[prices.date_to_index[day]]
            / prices.daily_level[prices.date_to_index[day]]
            for day in window
        ]
    )
    shape = np.mean(curves, axis=0)
    shape /= float(np.mean(shape))
    if shape.shape != (T,) or not np.all(np.isfinite(shape)) or float(np.min(shape)) <= 0.0:
        raise RuntimeError("历史价格形状异常")
    if abs(float(np.mean(shape)) - 1.0) > 1.0e-12:
        raise RuntimeError("历史价格形状均值不为1")
    return shape


def fallback_level(
    prices: PriceData,
    window: Sequence[date],
    reason: str,
) -> dict[str, Any]:
    selected = tuple(window[-MIN_REGRESSION_SAMPLES:])
    if not selected:
        raise RuntimeError("价格回退没有任何完整历史日")
    level = float(
        np.mean([prices.daily_level[prices.date_to_index[day]] for day in selected])
    )
    return {
        "intercept": max(PRICE_EPSILON, level),
        "coefficient": 0.0,
        "sample_count": len(selected),
        "status": f"fallback_recent_mean:{reason}",
        "train_dates": selected,
    }


def feature_nearly_constant(values: np.ndarray) -> bool:
    if values.size == 0:
        return True
    scale = max(1.0, float(np.max(np.abs(values))))
    return float(np.ptp(values)) <= 1.0e-10 * scale


def fit_ar_level(
    target: date,
    prices: PriceData,
    window: Sequence[date],
) -> dict[str, Any]:
    train: list[date] = []
    x_values: list[float] = []
    y_values: list[float] = []
    for day in window:
        previous = day - timedelta(days=1)
        if previous not in prices.date_to_index:
            continue
        train.append(day)
        x_values.append(float(prices.daily_level[prices.date_to_index[previous]]))
        y_values.append(float(prices.daily_level[prices.date_to_index[day]]))
    if len(train) < MIN_REGRESSION_SAMPLES:
        return fallback_level(prices, window, "ar_samples_lt_7")
    x = np.asarray(x_values)
    y = np.asarray(y_values)
    if feature_nearly_constant(x):
        return fallback_level(prices, window, "ar_feature_near_constant")
    centered = x - float(np.mean(x))
    rho_raw = float(np.dot(centered, y - float(np.mean(y))) / np.dot(centered, centered))
    rho = float(min(0.99, max(0.0, rho_raw)))
    intercept = float(np.mean(y) - rho * np.mean(x))
    return {
        "intercept": intercept,
        "coefficient": rho,
        "coefficient_raw": rho_raw,
        "sample_count": len(train),
        "status": "ols_bounded_rho",
        "train_dates": tuple(train),
    }


def fit_net_level(
    branch: str,
    hour: int,
    prices: PriceData,
    window: Sequence[date],
    resources: dict[str, dict[tuple[date, int], ResourceEvent]],
) -> dict[str, Any]:
    train: list[date] = []
    x_values: list[float] = []
    y_values: list[float] = []
    for day in window:
        event = resources[branch].get((day, hour))
        if event is None:
            continue
        train.append(day)
        x_values.append(event.current_net_energy_forecast_kwh / NET_LOAD_SCALE_KWH)
        y_values.append(float(prices.daily_level[prices.date_to_index[day]]))
    if len(train) < MIN_REGRESSION_SAMPLES:
        return fallback_level(prices, window, "net_samples_lt_7")
    x = np.asarray(x_values)
    y = np.asarray(y_values)
    if feature_nearly_constant(x):
        return fallback_level(prices, window, "net_feature_near_constant")
    centered = x - float(np.mean(x))
    beta = float(np.dot(centered, y - float(np.mean(y))) / np.dot(centered, centered))
    intercept = float(np.mean(y) - beta * np.mean(x))
    return {
        "intercept": intercept,
        "coefficient": beta,
        "coefficient_raw": beta,
        "sample_count": len(train),
        "status": "ols_net_load",
        "train_dates": tuple(train),
    }


def forecast_price_event(
    branch: str,
    event: ResourceEvent,
    family: str,
    short_rho: float,
    prices: PriceData,
    resources: dict[str, dict[tuple[date, int], ResourceEvent]],
) -> PriceForecast:
    if family not in PRICE_MODEL_CANDIDATES:
        raise ValueError(f"未知价格模型：{family}")
    window = history_window(prices, event.issue_date)
    shape = normalized_shape(prices, window)
    if family == "ar":
        fit = fit_ar_level(event.issue_date, prices, window)
        previous = event.issue_date - timedelta(days=1)
        if previous not in prices.date_to_index:
            current_level = float(fit["intercept"])
        else:
            current_level = float(fit["intercept"]) + float(fit["coefficient"]) * float(
                prices.daily_level[prices.date_to_index[previous]]
            )
        current_level = max(PRICE_EPSILON, current_level)
        next_level = max(
            PRICE_EPSILON,
            float(fit["intercept"]) + float(fit["coefficient"]) * current_level,
        )
    else:
        fit = fit_net_level(branch, event.issue_hour, prices, window, resources)
        current_level = max(
            PRICE_EPSILON,
            float(fit["intercept"])
            + float(fit["coefficient"])
            * event.current_net_energy_forecast_kwh
            / NET_LOAD_SCALE_KWH,
        )
        next_level = max(
            PRICE_EPSILON,
            float(fit["intercept"])
            + float(fit["coefficient"])
            * event.next_net_energy_forecast_kwh
            / NET_LOAD_SCALE_KWH,
        )

    base_current = np.maximum(PRICE_EPSILON, current_level * shape)
    base_next = np.maximum(PRICE_EPSILON, next_level * shape)
    k = KAPPA[event.issue_hour]
    ncur = T - k
    correction_current = np.zeros(ncur, dtype=np.float64)
    last_actual: float | None = None
    last_base: float | None = None
    if k > 0:
        if event.issue_date not in prices.date_to_index:
            raise RuntimeError("边界日期不应有日内价格修正")
        row = prices.date_to_index[event.issue_date]
        last_actual = float(prices.values[row, k - 1])
        last_base = float(base_current[k - 1])
        error = last_actual - last_base
        exponent = np.ceil(np.arange(1, ncur + 1, dtype=np.float64) / 6.0)
        correction_current = np.power(short_rho, exponent) * error
    current_part = np.maximum(
        PRICE_EPSILON, base_current[k:] + correction_current
    )
    if k == 0:
        base_window = base_current.copy()
        correction = np.zeros(T, dtype=np.float64)
        final = current_part
    else:
        base_window = np.concatenate((base_current[k:], base_next[:k]))
        correction = np.concatenate((correction_current, np.zeros(k, dtype=np.float64)))
        final = np.concatenate((current_part, base_next[:k]))
    shape_window = shape[event.t_abs - 1]
    if final.shape != (T,) or float(np.min(final)) < PRICE_EPSILON - 1.0e-15:
        raise RuntimeError("价格事件预测维度或下限异常")
    return PriceForecast(
        branch=branch,
        issue_date=event.issue_date,
        issue_hour=event.issue_hour,
        family=family,
        short_rho=short_rho,
        final=final,
        base=base_window,
        correction=correction,
        shape_window=shape_window,
        current_level=current_level,
        next_level=next_level,
        intercept=float(fit["intercept"]),
        coefficient=float(fit["coefficient"]),
        sample_count=int(fit["sample_count"]),
        fit_status=str(fit["status"]),
        train_dates=tuple(fit["train_dates"]),
        last_known_actual_price=last_actual,
        last_known_base_price=last_base,
    )


def actual_price_path(event: ResourceEvent, prices: PriceData) -> np.ndarray:
    result = np.full(T, np.nan, dtype=np.float64)
    valid = (event.day_indices >= 0) & (event.day_indices < len(prices.dates))
    result[valid] = prices.values[event.day_indices[valid], event.t_abs[valid] - 1]
    return result


def summarize_error(error: np.ndarray) -> dict[str, float]:
    if error.size == 0:
        return {"mae": float("nan"), "rmse": float("nan"), "bias": float("nan")}
    return {
        "mae": float(np.mean(np.abs(error))),
        "rmse": float(np.sqrt(np.mean(np.square(error)))),
        "bias": float(np.mean(error)),
    }


def select_price_model(
    resources: dict[str, dict[tuple[date, int], ResourceEvent]],
    prices: PriceData,
) -> tuple[str, list[dict[str, Any]]]:
    rows: list[dict[str, Any]] = []
    scores: dict[str, tuple[float, float]] = {}
    for family in PRICE_MODEL_CANDIDATES:
        branch_metrics: dict[str, dict[str, float]] = {}
        for branch in BRANCHES:
            errors: list[np.ndarray] = []
            event_count = 0
            for day_offset in range(17):
                day = date(2025, 1, 15) + timedelta(days=day_offset)
                event = resources[branch].get((day, 0))
                if event is None:
                    continue
                forecast = forecast_price_event(branch, event, family, 0.0, prices, resources)
                actual = actual_price_path(event, prices)
                mask = np.isfinite(actual)
                errors.append(forecast.final[mask] - actual[mask])
                event_count += 1
            if not errors:
                raise RuntimeError(f"{branch} {family}没有1月15-31日0:00回测")
            metrics = summarize_error(np.concatenate(errors))
            branch_metrics[branch] = metrics
            rows.append(
                {
                    "stage": "level_model_selection",
                    "candidate": family,
                    "branch": branch,
                    "event_count": event_count,
                    "position_count": int(sum(item.size for item in errors)),
                    "mae_yuan_per_kwh": metrics["mae"],
                    "rmse_yuan_per_kwh": metrics["rmse"],
                    "bias_forecast_minus_actual": metrics["bias"],
                    "selected": "",
                }
            )
        avg_mae = float(np.mean([branch_metrics[b]["mae"] for b in BRANCHES]))
        avg_rmse = float(np.mean([branch_metrics[b]["rmse"] for b in BRANCHES]))
        scores[family] = (avg_mae, avg_rmse)
        rows.append(
            {
                "stage": "level_model_selection",
                "candidate": family,
                "branch": "two_branch_average",
                "event_count": 34,
                "position_count": 34 * T,
                "mae_yuan_per_kwh": avg_mae,
                "rmse_yuan_per_kwh": avg_rmse,
                "bias_forecast_minus_actual": float(
                    np.mean([branch_metrics[b]["bias"] for b in BRANCHES])
                ),
                "selected": "",
            }
        )
    complexity_rank = {"ar": 0, "net": 1}
    selected = min(
        PRICE_MODEL_CANDIDATES,
        key=lambda family: (scores[family][0], scores[family][1], complexity_rank[family]),
    )
    for row in rows:
        row["selected"] = row["candidate"] == selected
    return selected, rows


def select_short_rho(
    family: str,
    resources: dict[str, dict[tuple[date, int], ResourceEvent]],
    prices: PriceData,
) -> tuple[float, list[dict[str, Any]]]:
    cutoff = issue_datetime(FORECAST_START_DATE, 0)
    rows: list[dict[str, Any]] = []
    scores: dict[float, tuple[float, float]] = {}
    for rho in SHORT_RHO_CANDIDATES:
        per_branch: dict[str, dict[str, float]] = {}
        branch_counts: dict[str, tuple[int, int]] = {}
        for branch in BRANCHES:
            errors: list[np.ndarray] = []
            count = 0
            for (day, hour), event in sorted(resources[branch].items()):
                if day.year != 2025 or day.month != 1 or hour == 0:
                    continue
                if event.close_datetime > cutoff:
                    continue
                forecast = forecast_price_event(branch, event, family, rho, prices, resources)
                actual = actual_price_path(event, prices)
                mask = np.isfinite(actual)
                if not np.all(mask):
                    continue
                errors.append(forecast.final - actual)
                count += 1
            if not errors:
                raise RuntimeError(f"{branch} rho={rho}没有合法日内回测路径")
            concatenated = np.concatenate(errors)
            per_branch[branch] = summarize_error(concatenated)
            branch_counts[branch] = (count, concatenated.size)
            rows.append(
                {
                    "stage": "intraday_rho_selection",
                    "candidate": rho,
                    "branch": branch,
                    "event_count": count,
                    "position_count": concatenated.size,
                    "mae_yuan_per_kwh": per_branch[branch]["mae"],
                    "rmse_yuan_per_kwh": per_branch[branch]["rmse"],
                    "bias_forecast_minus_actual": per_branch[branch]["bias"],
                    "selected": "",
                }
            )
        avg_mae = float(np.mean([per_branch[b]["mae"] for b in BRANCHES]))
        avg_rmse = float(np.mean([per_branch[b]["rmse"] for b in BRANCHES]))
        scores[rho] = (avg_mae, avg_rmse)
        rows.append(
            {
                "stage": "intraday_rho_selection",
                "candidate": rho,
                "branch": "two_branch_average",
                "event_count": sum(branch_counts[b][0] for b in BRANCHES),
                "position_count": sum(branch_counts[b][1] for b in BRANCHES),
                "mae_yuan_per_kwh": avg_mae,
                "rmse_yuan_per_kwh": avg_rmse,
                "bias_forecast_minus_actual": float(
                    np.mean([per_branch[b]["bias"] for b in BRANCHES])
                ),
                "selected": "",
            }
        )
    selected = min(SHORT_RHO_CANDIDATES, key=lambda rho: (scores[rho][0], scores[rho][1], rho))
    for row in rows:
        row["selected"] = float(row["candidate"]) == selected
    return selected, rows


def build_price_forecasts(
    family: str,
    short_rho: float,
    resources: dict[str, dict[tuple[date, int], ResourceEvent]],
    prices: PriceData,
) -> dict[str, dict[tuple[date, int], PriceForecast]]:
    forecasts: dict[str, dict[tuple[date, int], PriceForecast]] = {
        branch: {} for branch in BRANCHES
    }
    for branch in BRANCHES:
        for key, event in sorted(resources[branch].items()):
            forecasts[branch][key] = forecast_price_event(
                branch, event, family, short_rho, prices, resources
            )
    return forecasts


def write_price_forecast_audit(
    path: Path,
    forecasts: dict[str, dict[tuple[date, int], PriceForecast]],
    resources: dict[str, dict[tuple[date, int], ResourceEvent]],
    prices: PriceData,
) -> None:
    fields = [
        "record_type", "branch", "issue_date", "issue_hour", "family", "short_rho",
        "fit_status", "sample_count", "intercept", "coefficient", "current_level",
        "next_level", "train_date", "train_daily_level", "train_net_forecast_kwh",
        "window_position", "absolute_date", "absolute_t", "physical_interval",
        "shape", "base_price", "short_correction", "final_forecast_price",
        "last_known_actual_price", "last_known_base_price", "realized_price_post_event",
        "realized_price_role",
    ]
    labels = q2.physical_interval_labels()
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fields, extrasaction="raise")
        writer.writeheader()
        for branch in BRANCHES:
            for key, forecast in sorted(forecasts[branch].items()):
                event = resources[branch][key]
                for train_day in forecast.train_dates:
                    train_event = resources[branch].get((train_day, forecast.issue_hour))
                    writer.writerow(
                        {
                            "record_type": "training_day",
                            "branch": branch,
                            "issue_date": forecast.issue_date.isoformat(),
                            "issue_hour": forecast.issue_hour,
                            "family": forecast.family,
                            "short_rho": forecast.short_rho,
                            "fit_status": forecast.fit_status,
                            "sample_count": forecast.sample_count,
                            "intercept": forecast.intercept,
                            "coefficient": forecast.coefficient,
                            "current_level": forecast.current_level,
                            "next_level": forecast.next_level,
                            "train_date": train_day.isoformat(),
                            "train_daily_level": float(
                                prices.daily_level[prices.date_to_index[train_day]]
                            ),
                            "train_net_forecast_kwh": "" if train_event is None else train_event.current_net_energy_forecast_kwh,
                        }
                    )
                actual = actual_price_path(event, prices)
                for u in range(T):
                    absolute_day_index = int(event.day_indices[u])
                    absolute_day = (
                        prices.dates[absolute_day_index]
                        if 0 <= absolute_day_index < len(prices.dates)
                        else BOUNDARY_PLAN_DATE
                    )
                    writer.writerow(
                        {
                            "record_type": "event_interval",
                            "branch": branch,
                            "issue_date": forecast.issue_date.isoformat(),
                            "issue_hour": forecast.issue_hour,
                            "family": forecast.family,
                            "short_rho": forecast.short_rho,
                            "fit_status": forecast.fit_status,
                            "sample_count": forecast.sample_count,
                            "intercept": forecast.intercept,
                            "coefficient": forecast.coefficient,
                            "current_level": forecast.current_level,
                            "next_level": forecast.next_level,
                            "window_position": u + 1,
                            "absolute_date": absolute_day.isoformat(),
                            "absolute_t": int(event.t_abs[u]),
                            "physical_interval": labels[int(event.t_abs[u]) - 1],
                            "shape": float(forecast.shape_window[u]),
                            "base_price": float(forecast.base[u]),
                            "short_correction": float(forecast.correction[u]),
                            "final_forecast_price": float(forecast.final[u]),
                            "last_known_actual_price": "" if forecast.last_known_actual_price is None else forecast.last_known_actual_price,
                            "last_known_base_price": "" if forecast.last_known_base_price is None else forecast.last_known_base_price,
                            "realized_price_post_event": "" if not math.isfinite(float(actual[u])) else float(actual[u]),
                            "realized_price_role": "unavailable_boundary" if not math.isfinite(float(actual[u])) else "post_event_audit_only",
                        }
                    )


def build_residual_events(
    branch: str,
    resources: dict[str, dict[tuple[date, int], ResourceEvent]],
    forecasts: dict[str, dict[tuple[date, int], PriceForecast]],
    annual: q2.AnnualData,
    prices: PriceData,
) -> dict[int, list[ResidualEvent]]:
    pools: dict[int, list[ResidualEvent]] = {hour: [] for hour in ISSUE_HOURS}
    for key, event in sorted(resources[branch].items()):
        actual_resource = actual_resource_path(event, annual)
        actual_price = actual_price_path(event, prices)
        if actual_resource is None or not np.all(np.isfinite(actual_price)):
            continue
        load_actual, pv_actual = actual_resource
        forecast = forecasts[branch][key]
        pools[event.issue_hour].append(
            ResidualEvent(
                branch=branch,
                issue_date=event.issue_date,
                issue_hour=event.issue_hour,
                close_datetime=event.close_datetime,
                price_residual=actual_price - forecast.final,
                load_residual=load_actual - event.load_forecast,
                pv_residual=pv_actual - event.pv_forecast,
            )
        )
    for hour in ISSUE_HOURS:
        pools[hour].sort(key=lambda item: (item.issue_date, item.issue_hour))
    return pools


def recent_closed_residuals(
    pool: Sequence[ResidualEvent],
    current_issue: datetime,
) -> list[ResidualEvent]:
    selected = [item for item in pool if item.close_datetime <= current_issue]
    return selected[-MAX_SCENARIOS:]


def price_forecast_metric_rows(
    forecasts: dict[str, dict[tuple[date, int], PriceForecast]],
    resources: dict[str, dict[tuple[date, int], ResourceEvent]],
    prices: PriceData,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    actual_threshold = float(np.quantile(prices.values[31:], 0.9))
    for branch in BRANCHES:
        for hour in ISSUE_HOURS:
            for month in range(2, 13):
                errors: list[np.ndarray] = []
                actuals: list[np.ndarray] = []
                forecasts_used: list[np.ndarray] = []
                event_count = 0
                fallback_count = 0
                for (day, event_hour), forecast in forecasts[branch].items():
                    if day.year != 2025 or day.month != month or event_hour != hour:
                        continue
                    event = resources[branch][(day, hour)]
                    actual = actual_price_path(event, prices)
                    mask = np.isfinite(actual)
                    if not np.any(mask):
                        continue
                    errors.append(forecast.final[mask] - actual[mask])
                    actuals.append(actual[mask])
                    forecasts_used.append(forecast.final[mask])
                    event_count += 1
                    fallback_count += forecast.fit_status.startswith("fallback")
                if not errors:
                    continue
                error = np.concatenate(errors)
                actual = np.concatenate(actuals)
                predicted = np.concatenate(forecasts_used)
                metrics = summarize_error(error)
                peak = np.abs(error[actual >= actual_threshold])
                rows.append(
                    {
                        "branch": branch,
                        "issue_hour": hour,
                        "month": month,
                        "event_count": event_count,
                        "position_count": error.size,
                        "mae_yuan_per_kwh": metrics["mae"],
                        "rmse_yuan_per_kwh": metrics["rmse"],
                        "bias_forecast_minus_actual": metrics["bias"],
                        "peak_threshold_yuan_per_kwh": actual_threshold,
                        "peak_mae_yuan_per_kwh": float(np.mean(peak)) if peak.size else 0.0,
                        "max_absolute_error_yuan_per_kwh": float(np.max(np.abs(error))),
                        "forecast_floor_trigger_rate": float(np.mean(predicted <= PRICE_EPSILON + 1.0e-15)),
                        "fallback_event_count": fallback_count,
                    }
                )
    return rows


def save_array(out: Path, key: str, value: np.ndarray) -> Path:
    path = out / CACHE_FILES[key]
    np.save(path, value, allow_pickle=False)
    return path


def build_scenario_cache(
    out: Path,
    count: int,
    resources: dict[str, dict[tuple[date, int], ResourceEvent]],
    forecasts: dict[str, dict[tuple[date, int], PriceForecast]],
    annual: q2.AnnualData,
    prices: PriceData,
) -> dict[str, Any]:
    event_dates = formal_dates(count)
    event_count = count * len(ISSUE_HOURS)
    p_scenario = np.zeros((2, event_count, MAX_SCENARIOS, T), dtype=np.float64)
    n_scenario = np.zeros_like(p_scenario)
    m_array = np.zeros((2, event_count), dtype=np.int16)
    p_point = np.zeros((2, event_count, T), dtype=np.float64)
    l_point = np.zeros_like(p_point)
    v_point = np.zeros_like(p_point)
    bp = np.zeros((2, MAX_SCENARIOS, T), dtype=np.float64)
    bn = np.zeros_like(bp)
    bm = np.zeros(2, dtype=np.int16)
    bpp = np.zeros((2, T), dtype=np.float64)

    audit_fields = [
        "branch", "issue_date", "issue_hour", "scenario_rank", "weight",
        "residual_issue_date", "residual_issue_hour", "residual_close_time",
        "current_issue_time", "closed_by_issue", "price_residual_mae",
        "load_residual_mae_kwh", "pv_residual_mae_kwh", "scenario_price_floor_count",
    ]
    audit_path = out / "q4_scenario_audit.csv"
    total_floor = 0
    total_prices = 0
    min_count = MAX_SCENARIOS
    max_count = 0
    with audit_path.open("w", encoding="utf-8-sig", newline="") as audit_file:
        writer = csv.DictWriter(audit_file, fieldnames=audit_fields, extrasaction="raise")
        writer.writeheader()
        for branch in BRANCHES:
            bi = BRANCH_INDEX[branch]
            residual_pools = build_residual_events(
                branch, resources, forecasts, annual, prices
            )
            for di, day in enumerate(event_dates):
                for hour in ISSUE_HOURS:
                    key = (day, hour)
                    if key not in resources[branch] or key not in forecasts[branch]:
                        raise RuntimeError(f"{branch} {day} {hour}:00预测事件缺失")
                    event = resources[branch][key]
                    pf = forecasts[branch][key]
                    current_issue = event.issue_datetime
                    selected = recent_closed_residuals(
                        residual_pools[hour], current_issue
                    )
                    if not selected:
                        raise RuntimeError(f"{branch} {day} {hour}:00无合法已关闭残差场景")
                    m = len(selected)
                    ei = di * 4 + HOUR_INDEX[hour]
                    m_array[bi, ei] = m
                    p_point[bi, ei] = pf.final
                    l_point[bi, ei] = event.load_forecast
                    v_point[bi, ei] = event.pv_forecast
                    for rank, residual in enumerate(selected):
                        scenario_p = np.maximum(
                            PRICE_EPSILON, pf.final + residual.price_residual
                        )
                        scenario_l = np.maximum(
                            0.0, event.load_forecast + residual.load_residual
                        )
                        scenario_v = np.maximum(
                            0.0, event.pv_forecast + residual.pv_residual
                        )
                        p_scenario[bi, ei, rank] = scenario_p
                        n_scenario[bi, ei, rank] = scenario_l - scenario_v
                        floor_count = int(np.count_nonzero(scenario_p <= PRICE_EPSILON + 1.0e-15))
                        total_floor += floor_count
                        total_prices += T
                        closed = residual.close_datetime <= current_issue
                        if not closed:
                            raise RuntimeError("场景残差在当前发布时刻尚未关闭")
                        writer.writerow(
                            {
                                "branch": branch,
                                "issue_date": day.isoformat(),
                                "issue_hour": hour,
                                "scenario_rank": rank + 1,
                                "weight": 1.0 / m,
                                "residual_issue_date": residual.issue_date.isoformat(),
                                "residual_issue_hour": residual.issue_hour,
                                "residual_close_time": format_datetime(residual.close_datetime),
                                "current_issue_time": format_datetime(current_issue),
                                "closed_by_issue": closed,
                                "price_residual_mae": float(np.mean(np.abs(residual.price_residual))),
                                "load_residual_mae_kwh": float(np.mean(np.abs(residual.load_residual))),
                                "pv_residual_mae_kwh": float(np.mean(np.abs(residual.pv_residual))),
                                "scenario_price_floor_count": floor_count,
                            }
                        )
                    min_count = min(min_count, m)
                    max_count = max(max_count, m)

            boundary_event = resources[branch][(BOUNDARY_PLAN_DATE, 0)]
            boundary_pf = forecasts[branch][(BOUNDARY_PLAN_DATE, 0)]
            selected = recent_closed_residuals(
                residual_pools[0], boundary_event.issue_datetime
            )
            if not selected:
                raise RuntimeError(f"{branch}边界事件无合法残差场景")
            bm[bi] = len(selected)
            bpp[bi] = boundary_pf.final
            for rank, residual in enumerate(selected):
                scenario_p = np.maximum(
                    PRICE_EPSILON, boundary_pf.final + residual.price_residual
                )
                scenario_l = np.maximum(
                    0.0, boundary_event.load_forecast + residual.load_residual
                )
                scenario_v = np.maximum(
                    0.0, boundary_event.pv_forecast + residual.pv_residual
                )
                bp[bi, rank] = scenario_p
                bn[bi, rank] = scenario_l - scenario_v
                writer.writerow(
                    {
                        "branch": branch,
                        "issue_date": BOUNDARY_PLAN_DATE.isoformat(),
                        "issue_hour": 0,
                        "scenario_rank": rank + 1,
                        "weight": 1.0 / len(selected),
                        "residual_issue_date": residual.issue_date.isoformat(),
                        "residual_issue_hour": residual.issue_hour,
                        "residual_close_time": format_datetime(residual.close_datetime),
                        "current_issue_time": format_datetime(boundary_event.issue_datetime),
                        "closed_by_issue": True,
                        "price_residual_mae": float(np.mean(np.abs(residual.price_residual))),
                        "load_residual_mae_kwh": float(np.mean(np.abs(residual.load_residual))),
                        "pv_residual_mae_kwh": float(np.mean(np.abs(residual.pv_residual))),
                        "scenario_price_floor_count": int(np.count_nonzero(scenario_p <= PRICE_EPSILON + 1.0e-15)),
                    }
                )

    save_array(out, "scenario_price", p_scenario)
    save_array(out, "scenario_net", n_scenario)
    save_array(out, "scenario_count", m_array)
    save_array(out, "point_price", p_point)
    save_array(out, "point_load", l_point)
    save_array(out, "point_pv", v_point)
    save_array(out, "boundary_price", bp)
    save_array(out, "boundary_net", bn)
    save_array(out, "boundary_count", bm)
    save_array(out, "boundary_point_price", bpp)
    return {
        "event_count": event_count,
        "formal_days": count,
        "scenario_count_min": int(min_count),
        "scenario_count_max": int(max_count),
        "scenario_price_floor_trigger_rate": total_floor / total_prices if total_prices else 0.0,
        "cache_shapes": {
            "scenario_price": list(p_scenario.shape),
            "scenario_net": list(n_scenario.shape),
            "scenario_count": list(m_array.shape),
            "point_price": list(p_point.shape),
        },
    }


def prep_cache_matches(out: Path, count: int, args: argparse.Namespace) -> bool:
    path = out / CACHE_FILES["metadata"]
    if not path.is_file():
        return False
    try:
        metadata = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return False
    expected_hashes = {
        key: q2.sha256_file(input_path(args, key))
        for key in ("annual", "pv_forecast", "prices")
    }
    return (
        metadata.get("formal_days") == count
        and metadata.get("input_sha256") == expected_hashes
        and all((out / filename).is_file() for filename in CACHE_FILES.values())
    )


def run_prep(args: argparse.Namespace, count: int) -> dict[str, Any]:
    out = get_output_dir(args)
    if args.resume and prep_cache_matches(out, count, args):
        return json.loads((out / CACHE_FILES["metadata"]).read_text(encoding="utf-8"))
    started = time.perf_counter()
    annual = q2.read_annual_data(input_path(args, "annual"))
    prices = read_price_data(input_path(args, "prices"))
    validate_inputs(annual, prices)
    attachment3 = q3.read_attachment3(input_path(args, "pv_forecast"))
    raw, _, _ = build_raw_resource_events(annual, attachment3)
    bg = calibrate_bg_weight(raw["4-3"], annual)
    resources = finalize_resource_events(raw, annual, float(bg["w"]))
    family, model_rows = select_price_model(resources, prices)
    short_rho, rho_rows = select_short_rho(family, resources, prices)
    selection_rows = model_rows + rho_rows
    write_csv(
        out / "q4_price_model_selection.csv",
        selection_rows,
        [
            "stage", "candidate", "branch", "event_count", "position_count",
            "mae_yuan_per_kwh", "rmse_yuan_per_kwh", "bias_forecast_minus_actual",
            "selected",
        ],
    )
    forecasts = build_price_forecasts(family, short_rho, resources, prices)
    write_price_forecast_audit(
        out / "q4_price_forecast_audit.csv", forecasts, resources, prices
    )
    metric_rows = price_forecast_metric_rows(forecasts, resources, prices)
    write_csv(
        out / "q4_price_forecast_metrics.csv",
        metric_rows,
        list(metric_rows[0].keys()),
    )
    scenario_summary = build_scenario_cache(
        out, count, resources, forecasts, annual, prices
    )
    input_hashes = {
        key: q2.sha256_file(input_path(args, key))
        for key in ("annual", "pv_forecast", "prices")
    }
    shape_max_error = max(
        abs(float(np.mean(forecast.shape_window)) - 1.0)
        for branch in BRANCHES
        for forecast in forecasts[branch].values()
        if forecast.issue_hour == 0
    )
    metadata = {
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "formal_days": count,
        "selected_price_model_family": family,
        "selected_short_rho": short_rho,
        "selection_rule": "two-branch average 00:00 MAE, then RMSE, exact tie favors AR; rho uses MAE/RMSE/lower-rho",
        "price_history_days": PRICE_HISTORY_DAYS,
        "minimum_regression_samples": MIN_REGRESSION_SAMPLES,
        "price_epsilon": PRICE_EPSILON,
        "net_load_scale_kwh": NET_LOAD_SCALE_KWH,
        "bates_granger": bg,
        "scenario": scenario_summary,
        "max_zero_hour_shape_mean_error": shape_max_error,
        "input_sha256": input_hashes,
        "runtime_seconds": time.perf_counter() - started,
    }
    (out / CACHE_FILES["metadata"]).write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    print(
        f"[prep] 完成：模型={family}, rho_p={short_rho}, "
        f"事件={scenario_summary['event_count']}, 耗时={metadata['runtime_seconds']:.1f}s",
        flush=True,
    )
    return metadata


def overnight_window_positions(issue_hour: int) -> np.ndarray:
    """窗口内与“距发布时刻最近的下一个0:00–5:00物理窗口”对应的144段位置（0-based）。

    h=0 时为当前日 t=1..30（窗口位置0..29）；h>0 时为次日 t=1..30，
    即窗口位置 ncur..ncur+29，其中 ncur=T-KAPPA[h]。
    """
    if issue_hour == 0:
        return np.arange(0, OVERNIGHT_WINDOW_SLOTS, dtype=np.int64)
    ncur = T - KAPPA[issue_hour]
    return np.arange(ncur, ncur + OVERNIGHT_WINDOW_SLOTS, dtype=np.int64)


def terminal_value_coefficient(
    mean_price: np.ndarray,
    mode: str,
    issue_hour: int,
) -> tuple[float, float]:
    """返回(v, 参考窗口场景均价)。与修复后的Q2/Q3同结构：
    正式口径为“下次凌晨0:00–5:00窗口场景均价÷eta_c（充电边际）”；
    eta_d*min / eta_d*median 仅为orderForQ4规定的敏感性对照，零档用于隔离终值影响。"""
    if mean_price.shape != (T,):
        raise ValueError("终值场景均价维度不是144")
    if mode == "zero":
        return 0.0, 0.0
    if mode == "overnight":
        reference = float(np.mean(mean_price[overnight_window_positions(issue_hour)]))
        return reference / ETA_CHARGE, reference
    if mode == "minimum":
        reference = float(np.min(mean_price))
        return ETA_DISCHARGE * reference, reference
    if mode == "median":
        reference = float(np.median(mean_price))
        return ETA_DISCHARGE * reference, reference
    raise ValueError(f"未知终值代理模式：{mode}")


def build_window_model(
    label: str,
    mean_price: np.ndarray,
    scenario_price: np.ndarray,
    scenario_net: np.ndarray,
    initial_soc: float,
    terminal_value: float,
    ncur: int,
    g0: np.ndarray | None,
    *,
    adjust_purchases: bool,
    fixed_current: bool,
    binary_mutex: bool,
) -> WindowModel:
    scenarios = scenario_net.shape[0]
    if scenario_price.shape != scenario_net.shape or scenario_net.shape[1] != T:
        raise ValueError("窗口价格与净负荷场景维度不一致")
    problem = pulp.LpProblem(f"Q4_{label}", pulp.LpMinimize)
    z = [pulp.LpVariable(f"z_{u + 1:03d}", lowBound=0.0) for u in range(T)]
    charge = [
        [
            pulp.LpVariable(
                f"C_{w + 1:02d}_{u + 1:03d}", lowBound=0.0, upBound=MAX_ACTION_KWH
            )
            for u in range(T)
        ]
        for w in range(scenarios)
    ]
    discharge = [
        [
            pulp.LpVariable(
                f"D_{w + 1:02d}_{u + 1:03d}", lowBound=0.0, upBound=MAX_ACTION_KWH
            )
            for u in range(T)
        ]
        for w in range(scenarios)
    ]
    emergency = [
        [pulp.LpVariable(f"b_{w + 1:02d}_{u + 1:03d}", lowBound=0.0) for u in range(T)]
        for w in range(scenarios)
    ]
    unused = [
        [pulp.LpVariable(f"U_{w + 1:02d}_{u + 1:03d}", lowBound=0.0) for u in range(T)]
        for w in range(scenarios)
    ]
    soc = [
        [
            pulp.LpVariable(
                f"E_{w + 1:02d}_{u + 1:03d}", lowBound=SOC_MIN_KWH, upBound=SOC_MAX_KWH
            )
            for u in range(T)
        ]
        for w in range(scenarios)
    ]
    modes: list[list[pulp.LpVariable]] | None = None
    if binary_mutex:
        modes = [
            [pulp.LpVariable(f"m_{w + 1:02d}_{u + 1:03d}", cat=pulp.LpBinary) for u in range(T)]
            for w in range(scenarios)
        ]
    delta_plus: list[pulp.LpVariable] = []
    delta_minus: list[pulp.LpVariable] = []
    if adjust_purchases:
        delta_plus = [pulp.LpVariable(f"dp_{u + 1:03d}", lowBound=0.0) for u in range(ncur)]
        delta_minus = [pulp.LpVariable(f"dm_{u + 1:03d}", lowBound=0.0) for u in range(ncur)]

    for w in range(scenarios):
        for u in range(T):
            problem += (
                z[u] + emergency[w][u] + discharge[w][u]
                == float(scenario_net[w, u]) + charge[w][u] + unused[w][u],
                f"balance_{w + 1:02d}_{u + 1:03d}",
            )
            previous: float | pulp.LpVariable = initial_soc if u == 0 else soc[w][u - 1]
            problem += (
                soc[w][u]
                == previous + ETA_CHARGE * charge[w][u] - discharge[w][u] / ETA_DISCHARGE,
                f"soc_{w + 1:02d}_{u + 1:03d}",
            )
            if modes is not None:
                problem += (
                    charge[w][u] <= MAX_ACTION_KWH * modes[w][u],
                    f"charge_mutex_{w + 1:02d}_{u + 1:03d}",
                )
                problem += (
                    discharge[w][u] <= MAX_ACTION_KWH * (1.0 - modes[w][u]),
                    f"discharge_mutex_{w + 1:02d}_{u + 1:03d}",
                )

    if g0 is not None:
        k = T - ncur
        if fixed_current:
            for u in range(ncur):
                problem += (
                    z[u] == float(g0[k + u]),
                    f"fixed_current_{u + 1:03d}",
                )
        elif adjust_purchases:
            for u in range(ncur):
                problem += (
                    z[u] - float(g0[k + u]) == delta_plus[u] - delta_minus[u],
                    f"delta_{u + 1:03d}",
                )

    if adjust_purchases:
        plan_cost = pulp.lpSum(
            float(mean_price[u]) * (z[u] + 0.5 * (delta_plus[u] + delta_minus[u]))
            for u in range(ncur)
        ) + pulp.lpSum(float(mean_price[u]) * z[u] for u in range(ncur, T))
    else:
        plan_cost = pulp.lpSum(float(mean_price[u]) * z[u] for u in range(T))
    probability = 1.0 / scenarios
    recourse = probability * pulp.lpSum(
        pulp.lpSum(
            EMERGENCY_PRICE_FACTOR * float(scenario_price[w, u]) * emergency[w][u]
            for u in range(T)
        )
        - terminal_value * soc[w][-1]
        for w in range(scenarios)
    )
    original = plan_cost + recourse
    problem += original, "expected_total_cost_with_terminal_value"
    return WindowModel(
        problem, original, z, charge, discharge, emergency, unused, soc,
        delta_plus, delta_minus, ncur,
    )


def model_arrays(model: WindowModel) -> dict[str, np.ndarray]:
    dp = np.zeros(T, dtype=np.float64)
    dm = np.zeros(T, dtype=np.float64)
    if model.delta_plus:
        dp[: model.ncur] = q2.values_1d(model.delta_plus)
        dm[: model.ncur] = q2.values_1d(model.delta_minus)
    return {
        "z": q2.values_1d(model.z),
        "c": q2.values_2d(model.charge),
        "d": q2.values_2d(model.discharge),
        "b": q2.values_2d(model.emergency),
        "u": q2.values_2d(model.unused),
        "e": q2.values_2d(model.soc),
        "dp": dp,
        "dm": dm,
    }


def solve_window(
    issue_date_value: date,
    hour: int,
    kind: str,
    mean_price: np.ndarray,
    scenario_price: np.ndarray,
    scenario_net: np.ndarray,
    initial_soc: float,
    terminal_value: float,
    ncur: int,
    g0: np.ndarray | None,
    *,
    adjust_purchases: bool,
    fixed_current: bool,
) -> WindowResult:
    if mean_price.shape != (T,):
        raise ValueError("场景均价维度不是144")
    if scenario_price.shape != scenario_net.shape or scenario_net.ndim != 2 or scenario_net.shape[1] != T:
        raise ValueError("场景价格/净负荷维度异常")
    if scenario_net.shape[0] < 1:
        raise ValueError("窗口没有场景")
    if not (SOC_MIN_KWH - CHECK_TOLERANCE <= initial_soc <= SOC_MAX_KWH + CHECK_TOLERANCE):
        raise ValueError(f"窗口初始SOC越界：{initial_soc}")
    if g0 is None:
        if ncur != T or adjust_purchases or fixed_current:
            raise ValueError("0:00原计划参数组合错误")
    else:
        if g0.shape != (T,) or not (0 < ncur < T):
            raise ValueError("日内窗口g0/ncur异常")
        if adjust_purchases == fixed_current:
            raise ValueError("日内窗口必须且只能选择调整或固定当前购电")

    label = f"{issue_date_value.isoformat()}_{hour:02d}_{kind}"
    model = build_window_model(
        label, mean_price, scenario_price, scenario_net, initial_soc, terminal_value,
        ncur, g0, adjust_purchases=adjust_purchases, fixed_current=fixed_current,
        binary_mutex=False,
    )
    status, solve_seconds, info = q2.solve_current_model(model, mip=False)
    best_objective = float(pulp.value(model.original_objective))
    arrays = model_arrays(model)
    max_product = float(np.max(arrays["c"] * arrays["d"]))
    tiebreak = False
    binary = False
    if max_product > SIMULTANEOUS_PRODUCT_TOLERANCE:
        tiebreak = True
        tolerance = max(1.0e-7, abs(best_objective) * 1.0e-10)
        model.problem += (
            model.original_objective <= best_objective + tolerance,
            "preserve_primary_optimal_cost",
        )
        model.problem.setObjective(
            pulp.lpSum(
                variable
                for rows in (model.charge, model.discharge)
                for row in rows
                for variable in row
            )
        )
        tie_status, elapsed, info = q2.solve_current_model(model, mip=False)
        status = f"{status}; tiebreak={tie_status}"
        solve_seconds += elapsed
        arrays = model_arrays(model)
        max_product = float(np.max(arrays["c"] * arrays["d"]))
    if max_product > SIMULTANEOUS_PRODUCT_TOLERANCE:
        binary = True
        model = build_window_model(
            label, mean_price, scenario_price, scenario_net, initial_soc, terminal_value,
            ncur, g0, adjust_purchases=adjust_purchases, fixed_current=fixed_current,
            binary_mutex=True,
        )
        status, elapsed, info = q2.solve_current_model(model, mip=True)
        solve_seconds += elapsed
        status = f"binary_fallback={status}"
        arrays = model_arrays(model)
        max_product = float(np.max(arrays["c"] * arrays["d"]))

    z = arrays["z"]
    c = arrays["c"]
    d = arrays["d"]
    b = arrays["b"]
    unused = arrays["u"]
    soc = arrays["e"]
    previous = np.concatenate(
        (np.full((soc.shape[0], 1), initial_soc), soc[:, :-1]), axis=1
    )
    balance = z[None, :] + b + d - scenario_net - c - unused
    transition = soc - previous - ETA_CHARGE * c + d / ETA_DISCHARGE
    max_balance = float(np.max(np.abs(balance)))
    max_soc_residual = float(np.max(np.abs(transition)))
    max_delta_product = float(np.max(arrays["dp"] * arrays["dm"]))
    min_flow = float(min(np.min(z), np.min(c), np.min(d), np.min(b), np.min(unused)))
    checks = {
        "balance": max_balance <= CHECK_TOLERANCE,
        "soc_transition": max_soc_residual <= CHECK_TOLERANCE,
        "soc_bounds": float(np.min(soc)) >= SOC_MIN_KWH - CHECK_TOLERANCE
        and float(np.max(soc)) <= SOC_MAX_KWH + CHECK_TOLERANCE,
        "action_bounds": float(np.max(c)) <= MAX_ACTION_KWH + CHECK_TOLERANCE
        and float(np.max(d)) <= MAX_ACTION_KWH + CHECK_TOLERANCE,
        "nonnegative": min_flow >= -CHECK_TOLERANCE,
        "charge_discharge_mutex": max_product <= SIMULTANEOUS_PRODUCT_TOLERANCE,
        "delta_mutex": max_delta_product <= SIMULTANEOUS_PRODUCT_TOLERANCE,
    }
    if fixed_current and g0 is not None:
        k = T - ncur
        checks["fixed_current"] = float(np.max(np.abs(z[:ncur] - g0[k:]))) <= CHECK_TOLERANCE
    if not all(checks.values()):
        raise RuntimeError(f"{label}窗口LP验收失败：{checks}")
    cleaned = z.copy()
    cleaned[np.abs(cleaned) <= ACTION_ZERO_TOLERANCE] = 0.0
    if adjust_purchases and g0 is not None:
        k = T - ncur
        delta = cleaned[:ncur] - g0[k:]
        dp = np.zeros(T)
        dm = np.zeros(T)
        dp[:ncur] = np.maximum(delta, 0.0)
        dm[:ncur] = np.maximum(-delta, 0.0)
    else:
        dp = np.zeros(T)
        dm = np.zeros(T)
    mip_gap = getattr(info, "mip_gap", None) if binary else None
    return WindowResult(
        issue_date_value,
        hour,
        kind,
        cleaned,
        dp,
        dm,
        float(pulp.value(model.original_objective)),
        scenario_net.shape[0],
        status,
        solve_seconds,
        tiebreak,
        binary,
        max_balance,
        max_soc_residual,
        max_product,
        max_delta_product,
        min_flow,
        float(np.min(soc)),
        float(np.max(soc)),
        float(np.max(c)),
        float(np.max(d)),
        float(getattr(info, "max_primal_infeasibility", float("nan"))),
        int(getattr(info, "simplex_iteration_count", -1)),
        int(getattr(info, "ipm_iteration_count", -1)),
        None if mip_gap is None else float(mip_gap),
    )


def average_scenario_value_functions(
    scenario_net: np.ndarray,
    plan: np.ndarray,
    scenario_price: np.ndarray,
    terminal_value: float,
    grid: np.ndarray,
) -> tuple[np.ndarray, float]:
    if scenario_net.shape != scenario_price.shape or scenario_net.ndim != 2 or scenario_net.shape[1] != T:
        raise ValueError("DP场景价格与净负荷维度不一致")
    scenarios = scenario_net.shape[0]
    value_sum = np.zeros((T + 1, grid.size), dtype=np.float64)
    terminal = -terminal_value * grid
    value_sum[T] = scenarios * terminal
    max_convexity = 0.0
    for w in range(scenarios):
        next_value = terminal.copy()
        for u in range(T - 1, -1, -1):
            r = float(scenario_net[w, u] - plan[u])
            price = float(scenario_price[w, u])
            if r > 0.0:
                internal_span = min(MAX_ACTION_KWH, r) / ETA_DISCHARGE
                lower = np.maximum(SOC_MIN_KWH, grid - internal_span)
                upper = grid
                slope = EMERGENCY_PRICE_FACTOR * price * ETA_DISCHARGE
                transformed = next_value + slope * grid
                unconstrained = float(grid[int(np.argmin(transformed))])
                next_soc = np.clip(unconstrained, lower, upper)
                action = next_soc - grid
                current = (
                    EMERGENCY_PRICE_FACTOR
                    * price
                    * np.maximum(r + ETA_DISCHARGE * action, 0.0)
                    + np.interp(next_soc, grid, next_value)
                )
            else:
                internal_span = ETA_CHARGE * min(MAX_ACTION_KWH, -r)
                lower = grid
                upper = np.minimum(SOC_MAX_KWH, grid + internal_span)
                unconstrained = float(grid[int(np.argmin(next_value))])
                next_soc = np.clip(unconstrained, lower, upper)
                current = np.interp(next_soc, grid, next_value)
            violation = q2.convexity_violation(current, grid)
            max_convexity = max(max_convexity, violation)
            if violation > 1.0e-7:
                raise RuntimeError(
                    f"Q4 DP价值函数失去凸性：scenario={w + 1}, u={u + 1}, violation={violation}"
                )
            value_sum[u] += current
            next_value = current
    return value_sum / scenarios, max_convexity


def safe_candidate_minimizers(
    transformed_grid_values: np.ndarray,
    grid: np.ndarray,
    lower: float,
    upper: float,
    current_soc: float,
) -> list[float]:
    """候选点集合：可行区间端点、无约束最小值区带端点与当前SOC，全部强制裁剪到[lower, upper]。

    与q2.candidate_minimizers不同，q2的第五个候选在无约束最小值整体低于可行区间时
    会逃逸到lower以下（Q4逐场景低价时真实出现，导致单段放电超过上限）；本函数保证
    所有候选位于可行区间内。目标函数在区间上凸，最小值必位于无约束最小化集合到
    区间的投影处，因此裁剪不损失最优性。
    """
    minimum = float(np.min(transformed_grid_values))
    tie_tolerance = max(1.0e-10, abs(minimum) * 1.0e-12)
    tied = np.flatnonzero(transformed_grid_values <= minimum + tie_tolerance)
    q_low = float(grid[int(tied[0])])
    q_high = float(grid[int(tied[-1])])
    raw = [lower, upper, q_low, q_high, current_soc]
    clipped = [min(max(value, lower), upper) for value in raw]
    unique: list[float] = []
    for value in clipped:
        if not any(abs(value - existing) <= 1.0e-10 for existing in unique):
            unique.append(value)
    return unique


def execute_span(
    actual_prices: np.ndarray,
    actual_net: np.ndarray,
    plan: np.ndarray,
    k_from: int,
    k_to: int,
    future_values: np.ndarray,
    value_anchor: int,
    current_soc: float,
    grid: np.ndarray,
    charge: np.ndarray,
    discharge: np.ndarray,
    emergency: np.ndarray,
    unused: np.ndarray,
    soc: np.ndarray,
) -> float:
    for t in range(k_from, k_to):
        value_index = t - value_anchor + 1
        if not 1 <= value_index <= T:
            raise RuntimeError("执行器价值函数索引越界")
        future = future_values[value_index]
        r = float(actual_net[t] - plan[t])
        actual_price = float(actual_prices[t])
        if r > 0.0:
            lower = max(
                SOC_MIN_KWH,
                current_soc - min(MAX_ACTION_KWH, r) / ETA_DISCHARGE,
            )
            upper = current_soc
            transformed = future + EMERGENCY_PRICE_FACTOR * actual_price * ETA_DISCHARGE * grid
        else:
            lower = current_soc
            upper = min(
                SOC_MAX_KWH,
                current_soc + ETA_CHARGE * min(MAX_ACTION_KWH, -r),
            )
            transformed = future
        candidates = safe_candidate_minimizers(transformed, grid, lower, upper, current_soc)
        evaluated: list[tuple[float, float, float]] = []
        for next_soc in candidates:
            x = next_soc - current_soc
            psi = x / ETA_CHARGE if x >= 0.0 else ETA_DISCHARGE * x
            stage = EMERGENCY_PRICE_FACTOR * actual_price * max(r + psi, 0.0)
            objective = stage + float(np.interp(next_soc, grid, future))
            evaluated.append((objective, abs(x), next_soc))
        best = min(item[0] for item in evaluated)
        tolerance = max(1.0e-9, abs(best) * 1.0e-12)
        candidates_best = [item for item in evaluated if item[0] <= best + tolerance]
        _, _, next_soc = min(candidates_best, key=lambda item: (item[1], item[2]))
        x = next_soc - current_soc
        if abs(x) <= ACTION_ZERO_TOLERANCE:
            x = 0.0
            next_soc = current_soc
        if x >= 0.0:
            charge[t] = x / ETA_CHARGE
        else:
            discharge[t] = -ETA_DISCHARGE * x
        if charge[t] > MAX_ACTION_KWH + CHECK_TOLERANCE or discharge[t] > MAX_ACTION_KWH + CHECK_TOLERANCE:
            raise RuntimeError(
                f"执行器动作越界：t={t + 1}, r={r}, current_soc={current_soc}, "
                f"charge={charge[t]}, discharge={discharge[t]}"
            )
        psi = x / ETA_CHARGE if x >= 0.0 else ETA_DISCHARGE * x
        emergency[t] = max(r + psi, 0.0)
        unused[t] = max(-r - psi, 0.0)
        current_soc = min(max(next_soc, SOC_MIN_KWH), SOC_MAX_KWH)
        soc[t] = current_soc
    return current_soc


def load_cache(out: Path) -> dict[str, np.ndarray]:
    result: dict[str, np.ndarray] = {}
    for key in (
        "scenario_price", "scenario_net", "scenario_count", "point_price", "point_load",
        "point_pv", "boundary_price", "boundary_net", "boundary_count",
        "boundary_point_price",
    ):
        path = out / CACHE_FILES[key]
        if not path.is_file():
            raise FileNotFoundError(f"缺少Q4缓存：{path}，请先运行--prep")
        result[key] = np.load(path, mmap_mode="r", allow_pickle=False)
    return result


def config_paths(out: Path, config: RunConfig) -> dict[str, Path]:
    prefix = f"q4_{config.name}"
    if config.main_output:
        branch_prefix = "q4-2" if config.branch == "4-2" else "q4-3"
        return {
            "daily": out / f"{branch_prefix}_daily_metrics.csv",
            "solver": out / f"{branch_prefix}_solver_audit.csv",
            "interval": out / f"{branch_prefix}_interval_details.csv",
            "versions": out / f"{branch_prefix}_plan_versions.csv",
            "arrays": out / f"{branch_prefix}_main_arrays.npz",
            "summary": out / f"{prefix}_summary.json",
            "progress": out / f"{prefix}_progress.log",
        }
    return {
        "daily": out / f"{prefix}_daily_metrics.csv",
        "solver": out / f"{prefix}_solver_audit.csv",
        "summary": out / f"{prefix}_summary.json",
        "progress": out / f"{prefix}_progress.log",
    }


def configuration_complete(out: Path, config: RunConfig, count: int) -> bool:
    paths = config_paths(out, config)
    if not paths["summary"].is_file():
        return False
    try:
        summary = json.loads(paths["summary"].read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return False
    required = [paths["daily"], paths["solver"]]
    if config.main_output:
        required.extend([paths["interval"], paths["versions"], paths["arrays"]])
    return summary.get("formal_days") == count and all(path.is_file() for path in required)


def solver_audit_row(
    config: RunConfig,
    result: WindowResult,
    initial_soc: float,
    terminal_value: float,
    terminal_reference: float,
    mean_price: np.ndarray,
) -> dict[str, Any]:
    data = asdict(result)
    data.pop("z")
    data.pop("delta_plus")
    data.pop("delta_minus")
    data["issue_date"] = result.issue_date.isoformat()
    return {
        "config": config.name,
        "branch": config.branch,
        "initial_soc_kwh": initial_soc,
        "terminal_mode": config.terminal_mode,
        "terminal_value_coefficient": terminal_value,
        "terminal_reference_scenario_mean": terminal_reference,
        "terminal_efficiency_basis": "charge" if config.terminal_mode == "overnight" else "discharge_sensitivity",
        "scenario_mean_price_min": float(np.min(mean_price)),
        "scenario_mean_price_median": float(np.median(mean_price)),
        **data,
    }


def write_plan_version_rows(
    writer: csv.DictWriter,
    config: RunConfig,
    day: date,
    hour: int,
    result: WindowResult,
    g0: np.ndarray,
    point_price: np.ndarray,
    mean_price: np.ndarray,
    terminal_value: float,
) -> None:
    k = KAPPA[hour]
    ncur = T - k
    for u in range(T):
        absolute = k + u + 1
        offset = 1 if absolute > T else 0
        absolute_t = absolute - T if offset else absolute
        if hour == 0:
            variable_type = "g0"
            reference: float | str = ""
        elif u < ncur:
            variable_type = f"a{hour}" if config.adjust_purchases else "fixed_g"
            reference = float(g0[k + u])
        else:
            variable_type = "gtilde"
            reference = ""
        writer.writerow(
            {
                "branch": config.branch,
                "date": day.isoformat(),
                "issue_hour": hour,
                "window_position": u + 1,
                "absolute_date_offset": offset,
                "absolute_t": absolute_t,
                "variable_type": variable_type,
                "plan_kwh": float(result.z[u]),
                "delta_plus_kwh": float(result.delta_plus[u]) if hour > 0 and u < ncur else "",
                "delta_minus_kwh": float(result.delta_minus[u]) if hour > 0 and u < ncur else "",
                "g0_reference_kwh": reference,
                "point_price_forecast": float(point_price[u]),
                "scenario_mean_price": float(mean_price[u]),
                "terminal_value_coefficient": terminal_value,
            }
        )


def run_configuration(args: argparse.Namespace, config: RunConfig, count: int) -> dict[str, Any]:
    out = get_output_dir(args)
    if args.resume and configuration_complete(out, config, count):
        return json.loads(config_paths(out, config)["summary"].read_text(encoding="utf-8"))
    cache = load_cache(out)
    metadata = json.loads((out / CACHE_FILES["metadata"]).read_text(encoding="utf-8"))
    if int(metadata["formal_days"]) != count:
        raise RuntimeError("缓存正式日数与运行配置不一致")
    annual = q2.read_annual_data(input_path(args, "annual"))
    prices = read_price_data(input_path(args, "prices"))
    validate_inputs(annual, prices)
    bi = BRANCH_INDEX[config.branch]
    grid = q2.state_grid(args.dp_grid_step)
    dates = formal_dates(count)
    date_to_index = {day: i for i, day in enumerate(annual.dates)}
    paths = config_paths(out, config)
    started = time.perf_counter()

    main_shape = (count, T)
    if config.main_output:
        g0_all = np.zeros(main_shape)
        geff_all = np.zeros(main_shape)
        charge_all = np.zeros(main_shape)
        discharge_all = np.zeros(main_shape)
        emergency_all = np.zeros(main_shape)
        unused_all = np.zeros(main_shape)
        soc_all = np.zeros(main_shape)
        initial_soc_all = np.zeros(count)
        terminal_soc_all = np.zeros(count)
    else:
        g0_all = geff_all = charge_all = discharge_all = emergency_all = unused_all = soc_all = None
        initial_soc_all = terminal_soc_all = None

    daily_rows: list[dict[str, Any]] = []
    solver_rows: list[dict[str, Any]] = []
    dp_initial_soc = INITIAL_SOC_KWH
    analytical_initial_soc = INITIAL_SOC_KWH
    max_day_convexity = 0.0
    max_billing_identity = 0.0
    max_cross_day_soc = 0.0
    total_delta_abs = 0.0

    version_fields = [
        "branch", "date", "issue_hour", "window_position", "absolute_date_offset",
        "absolute_t", "variable_type", "plan_kwh", "delta_plus_kwh",
        "delta_minus_kwh", "g0_reference_kwh", "point_price_forecast",
        "scenario_mean_price", "terminal_value_coefficient",
    ]
    interval_fields = [
        "branch", "date", "t", "source_right_end_label", "physical_interval",
        "responsible_issue_hour", "point_price_forecast", "scenario_mean_planning_price",
        "actual_settlement_and_current_price", "actual_load_kwh", "actual_pv_kwh",
        "forecast_load_kwh", "forecast_pv_kwh", "g0_kwh", "geff_kwh",
        "delta_plus_kwh", "delta_minus_kwh", "dp_charge_kwh", "dp_discharge_kwh",
        "dp_emergency_kwh", "dp_unused_kwh", "dp_end_soc_kwh",
        "interval_plan_cost_yuan", "interval_adjustment_transaction_yuan",
        "interval_adjustment_surcharge_yuan", "interval_emergency_cost_yuan",
        "interval_total_bill_yuan",
    ]
    versions_file: Any = None
    versions_writer: csv.DictWriter | None = None
    interval_file: Any = None
    interval_writer: csv.DictWriter | None = None
    if config.main_output:
        versions_file = paths["versions"].open("w", encoding="utf-8-sig", newline="")
        versions_writer = csv.DictWriter(versions_file, fieldnames=version_fields, extrasaction="raise")
        versions_writer.writeheader()
        interval_file = paths["interval"].open("w", encoding="utf-8-sig", newline="")
        interval_writer = csv.DictWriter(interval_file, fieldnames=interval_fields, extrasaction="raise")
        interval_writer.writeheader()

    progress = paths["progress"]
    progress.write_text(
        f"{datetime.now().astimezone().isoformat(timespec='seconds')} started {config.name} {count} days\n",
        encoding="utf-8",
    )
    try:
        for di, day in enumerate(dates):
            day_idx = date_to_index[day]
            actual_price = prices.values[day_idx]
            actual_load = annual.load_kwh[day_idx]
            actual_pv = annual.pv_kwh[day_idx]
            actual_net = actual_load - actual_pv
            charge = np.zeros(T)
            discharge = np.zeros(T)
            emergency = np.zeros(T)
            unused = np.zeros(T)
            soc = np.zeros(T)
            responsible_hour = np.zeros(T, dtype=np.int16)
            point_price_used = np.zeros(T)
            mean_price_used = np.zeros(T)
            load_forecast_used = np.zeros(T)
            pv_forecast_used = np.zeros(T)
            day_initial_soc = dp_initial_soc
            current_soc = dp_initial_soc
            active_values: np.ndarray | None = None
            active_anchor = 0
            previous_kappa = 0
            geff: np.ndarray | None = None
            g0: np.ndarray | None = None
            day_convexity = 0.0
            solved_count = 0

            ordered = config.update_hours
            next_kappas = [KAPPA[h] for h in ordered[1:]] + [T]
            for update_index, hour in enumerate(ordered):
                k = KAPPA[hour]
                k_next = next_kappas[update_index]
                if previous_kappa < k:
                    if active_values is None or geff is None:
                        raise RuntimeError("执行时缺少上一发布事件的价值函数或计划")
                    current_soc = execute_span(
                        actual_price, actual_net, geff, previous_kappa, k,
                        active_values, active_anchor, current_soc, grid,
                        charge, discharge, emergency, unused, soc,
                    )
                event_index = di * 4 + HOUR_INDEX[hour]
                m = int(cache["scenario_count"][bi, event_index])
                if m <= 0:
                    raise RuntimeError(f"{config.name} {day} {hour}:00场景缺失")
                scenario_price = np.asarray(cache["scenario_price"][bi, event_index, :m])
                scenario_net = np.asarray(cache["scenario_net"][bi, event_index, :m])
                point_price = np.asarray(cache["point_price"][bi, event_index])
                point_load = np.asarray(cache["point_load"][bi, event_index])
                point_pv = np.asarray(cache["point_pv"][bi, event_index])
                mean_price = np.mean(scenario_price, axis=0)
                terminal_value, terminal_reference = terminal_value_coefficient(
                    mean_price, config.terminal_mode, hour
                )
                initial_for_solve = current_soc
                if hour == 0:
                    result = solve_window(
                        day, hour, "original", mean_price, scenario_price, scenario_net,
                        initial_for_solve, terminal_value, T, None,
                        adjust_purchases=False, fixed_current=False,
                    )
                    g0 = result.z.copy()
                    geff = g0.copy()
                else:
                    if g0 is None or geff is None:
                        raise RuntimeError("日内更新前缺少0:00原计划")
                    ncur = T - k
                    if config.adjust_purchases:
                        result = solve_window(
                            day, hour, "rolling", mean_price, scenario_price, scenario_net,
                            initial_for_solve, terminal_value, ncur, g0,
                            adjust_purchases=True, fixed_current=False,
                        )
                        frozen_before = geff[:k].copy()
                        geff[k:k_next] = result.z[: k_next - k]
                        if not np.array_equal(geff[:k], frozen_before):
                            raise RuntimeError("4-3滚动更新覆盖了已执行段")
                    else:
                        result = solve_window(
                            day, hour, "fixed_current_value_refresh", mean_price,
                            scenario_price, scenario_net, initial_for_solve,
                            terminal_value, ncur, g0,
                            adjust_purchases=False, fixed_current=True,
                        )
                        if float(np.max(np.abs(result.z[:ncur] - g0[k:]))) > CHECK_TOLERANCE:
                            raise RuntimeError("固定普通购电的日内窗口出现a!=g")
                assert g0 is not None and geff is not None
                solved_count += 1
                solver_rows.append(
                    solver_audit_row(
                        config, result, initial_for_solve, terminal_value,
                        terminal_reference, mean_price,
                    )
                )
                if versions_writer is not None:
                    write_plan_version_rows(
                        versions_writer, config, day, hour, result, g0,
                        point_price, mean_price, terminal_value,
                    )
                span = k_next - k
                responsible_hour[k:k_next] = hour
                point_price_used[k:k_next] = point_price[:span]
                mean_price_used[k:k_next] = mean_price[:span]
                load_forecast_used[k:k_next] = point_load[:span]
                pv_forecast_used[k:k_next] = point_pv[:span]
                if hour == 0 or config.refresh_value:
                    active_values, convexity = average_scenario_value_functions(
                        scenario_net, result.z, scenario_price, terminal_value, grid
                    )
                    active_anchor = k
                    day_convexity = max(day_convexity, convexity)
                previous_kappa = k

            if previous_kappa < T:
                if active_values is None or geff is None:
                    raise RuntimeError("日末执行缺少价值函数或有效计划")
                current_soc = execute_span(
                    actual_price, actual_net, geff, previous_kappa, T,
                    active_values, active_anchor, current_soc, grid,
                    charge, discharge, emergency, unused, soc,
                )
            assert g0 is not None and geff is not None
            dp_result = q2.validate_executor(
                actual_price, actual_net, geff, day_initial_soc,
                charge, discharge, emergency, unused, soc,
                dp_max_convexity_violation=day_convexity,
            )
            # validate_executor会统一清除数值容差内的伪流量；后续结算和导出必须使用
            # 清理后的同一组数组，避免把1e-11量级求解噪声误记为紧急购电。
            charge = dp_result.charge
            discharge = dp_result.discharge
            emergency = dp_result.emergency
            unused = dp_result.unused
            soc = dp_result.soc
            current_soc = dp_result.terminal_soc_kwh
            analytical = q2.execute_analytical(
                actual_price, actual_net, geff, analytical_initial_soc
            )
            delta = geff - g0
            delta_plus = np.maximum(delta, 0.0)
            delta_minus = np.maximum(-delta, 0.0)
            plan_cost = float(np.dot(actual_price, g0))
            adjustment_transaction = float(
                np.dot(
                    actual_price,
                    ADJUST_UP_PRICE_FACTOR * delta_plus
                    - ADJUST_DOWN_PRICE_FACTOR * delta_minus,
                )
            )
            adjustment_surcharge = float(np.dot(actual_price, 0.5 * np.abs(delta)))
            effective_base_cost = float(np.dot(actual_price, geff))
            emergency_cost = float(np.dot(EMERGENCY_PRICE_FACTOR * actual_price, emergency))
            if config.branch == "4-2":
                total_cost = plan_cost + emergency_cost
                total_alt = float(np.dot(actual_price, g0 + EMERGENCY_PRICE_FACTOR * emergency))
            else:
                total_cost = plan_cost + adjustment_transaction + emergency_cost
                total_alt = float(
                    np.dot(
                        actual_price,
                        geff + 0.5 * np.abs(delta) + EMERGENCY_PRICE_FACTOR * emergency,
                    )
                )
            identity = abs(total_cost - total_alt)
            max_billing_identity = max(max_billing_identity, identity)
            if identity > CHECK_TOLERANCE:
                raise RuntimeError(f"{config.name} {day}两种账单公式不一致：{identity}")
            if di > 0:
                max_cross_day_soc = max(max_cross_day_soc, abs(day_initial_soc - dp_initial_soc))
            total_delta_abs += float(np.sum(np.abs(delta)))
            daily_rows.append(
                {
                    "date": day.isoformat(),
                    "config": config.name,
                    "branch": config.branch,
                    "terminal_mode": config.terminal_mode,
                    "updates_solved": solved_count,
                    "dp_initial_soc_kwh": day_initial_soc,
                    "dp_terminal_soc_kwh": dp_result.terminal_soc_kwh,
                    "analytical_initial_soc_kwh": analytical_initial_soc,
                    "analytical_terminal_soc_kwh": analytical.terminal_soc_kwh,
                    "g0_energy_kwh": float(np.sum(g0)),
                    "geff_energy_kwh": float(np.sum(geff)),
                    "adjust_up_energy_kwh": float(np.sum(delta_plus)),
                    "adjust_down_energy_kwh": float(np.sum(delta_minus)),
                    "adjust_abs_energy_kwh": float(np.sum(np.abs(delta))),
                    "dp_charge_kwh": float(np.sum(charge)),
                    "dp_discharge_kwh": float(np.sum(discharge)),
                    "dp_emergency_energy_kwh": float(np.sum(emergency)),
                    "dp_unused_kwh": float(np.sum(unused)),
                    "plan_cost_yuan": plan_cost,
                    "effective_base_cost_yuan": effective_base_cost,
                    "adjustment_transaction_yuan": adjustment_transaction,
                    "adjustment_surcharge_yuan": adjustment_surcharge,
                    "dp_emergency_cost_yuan": emergency_cost,
                    "dp_total_cost_yuan": total_cost,
                    "dp_total_cost_alt_yuan": total_alt,
                    "billing_identity_residual_yuan": identity,
                    "analytical_emergency_energy_kwh": analytical.emergency_energy_kwh,
                    "analytical_emergency_cost_yuan": analytical.emergency_cost_yuan,
                    "analytical_total_cost_yuan": (
                        plan_cost + analytical.emergency_cost_yuan
                        if config.branch == "4-2"
                        else plan_cost + adjustment_transaction + analytical.emergency_cost_yuan
                    ),
                    "dp_max_balance_residual_kwh": dp_result.max_balance_residual_kwh,
                    "dp_max_soc_residual_kwh": dp_result.max_soc_residual_kwh,
                    "dp_max_simultaneous_product_kwh2": dp_result.max_simultaneous_product_kwh2,
                    "dp_max_convexity_violation": day_convexity,
                    "analytical_max_balance_residual_kwh": analytical.max_balance_residual_kwh,
                    "analytical_max_soc_residual_kwh": analytical.max_soc_residual_kwh,
                }
            )
            if config.main_output:
                assert all(
                    value is not None
                    for value in (
                        g0_all, geff_all, charge_all, discharge_all, emergency_all,
                        unused_all, soc_all, initial_soc_all, terminal_soc_all,
                    )
                )
                g0_all[di] = g0
                geff_all[di] = geff
                charge_all[di] = charge
                discharge_all[di] = discharge
                emergency_all[di] = emergency
                unused_all[di] = unused
                soc_all[di] = soc
                initial_soc_all[di] = day_initial_soc
                terminal_soc_all[di] = dp_result.terminal_soc_kwh
                assert interval_writer is not None
                labels = q2.physical_interval_labels()
                for t in range(T):
                    interval_writer.writerow(
                        {
                            "branch": config.branch,
                            "date": day.isoformat(),
                            "t": t + 1,
                            "source_right_end_label": annual.right_end_labels[t],
                            "physical_interval": labels[t],
                            "responsible_issue_hour": int(responsible_hour[t]),
                            "point_price_forecast": float(point_price_used[t]),
                            "scenario_mean_planning_price": float(mean_price_used[t]),
                            "actual_settlement_and_current_price": float(actual_price[t]),
                            "actual_load_kwh": float(actual_load[t]),
                            "actual_pv_kwh": float(actual_pv[t]),
                            "forecast_load_kwh": float(load_forecast_used[t]),
                            "forecast_pv_kwh": float(pv_forecast_used[t]),
                            "g0_kwh": float(g0[t]),
                            "geff_kwh": float(geff[t]),
                            "delta_plus_kwh": float(delta_plus[t]),
                            "delta_minus_kwh": float(delta_minus[t]),
                            "dp_charge_kwh": float(charge[t]),
                            "dp_discharge_kwh": float(discharge[t]),
                            "dp_emergency_kwh": float(emergency[t]),
                            "dp_unused_kwh": float(unused[t]),
                            "dp_end_soc_kwh": float(soc[t]),
                            "interval_plan_cost_yuan": float(actual_price[t] * g0[t]),
                            "interval_adjustment_transaction_yuan": float(
                                actual_price[t]
                                * (
                                    ADJUST_UP_PRICE_FACTOR * delta_plus[t]
                                    - ADJUST_DOWN_PRICE_FACTOR * delta_minus[t]
                                )
                            ),
                            "interval_adjustment_surcharge_yuan": float(
                                0.5 * actual_price[t] * abs(delta[t])
                            ),
                            "interval_emergency_cost_yuan": float(
                                EMERGENCY_PRICE_FACTOR * actual_price[t] * emergency[t]
                            ),
                            "interval_total_bill_yuan": float(
                                actual_price[t]
                                * (
                                    (g0[t] if config.branch == "4-2" else geff[t] + 0.5 * abs(delta[t]))
                                    + EMERGENCY_PRICE_FACTOR * emergency[t]
                                )
                            ),
                        }
                    )
            dp_initial_soc = dp_result.terminal_soc_kwh
            analytical_initial_soc = analytical.terminal_soc_kwh
            max_day_convexity = max(max_day_convexity, day_convexity)
            line = (
                f"{datetime.now().astimezone().isoformat(timespec='seconds')} "
                f"[{di + 1}/{count}] {day} cost={total_cost:.2f} soc={dp_initial_soc:.3f} "
                f"elapsed={time.perf_counter() - started:.0f}s"
            )
            with progress.open("a", encoding="utf-8") as file:
                file.write(line + "\n")
            if args.progress_every > 0 and (
                di == 0 or (di + 1) % args.progress_every == 0 or di + 1 == count
            ):
                print(f"[{config.name}] {line}", flush=True)
    finally:
        if versions_file is not None:
            versions_file.close()
        if interval_file is not None:
            interval_file.close()

    daily_fields = list(daily_rows[0].keys())
    write_csv(paths["daily"], daily_rows, daily_fields)
    solver_fields = list(solver_rows[0].keys())
    write_csv(paths["solver"], solver_rows, solver_fields)

    boundary_plan = np.zeros(T)
    boundary_mean = np.zeros(T)
    if count == FORMAL_DAYS:
        m = int(cache["boundary_count"][bi])
        scenario_price = np.asarray(cache["boundary_price"][bi, :m])
        scenario_net = np.asarray(cache["boundary_net"][bi, :m])
        boundary_mean = np.mean(scenario_price, axis=0)
        terminal_value, _ = terminal_value_coefficient(
            boundary_mean, config.terminal_mode, 0
        )
        boundary_result = solve_window(
            BOUNDARY_PLAN_DATE, 0, "template_boundary", boundary_mean,
            scenario_price, scenario_net, dp_initial_soc, terminal_value, T, None,
            adjust_purchases=False, fixed_current=False,
        )
        boundary_plan = boundary_result.z

    if config.main_output:
        assert all(
            value is not None
            for value in (
                g0_all, geff_all, charge_all, discharge_all, emergency_all,
                unused_all, soc_all, initial_soc_all, terminal_soc_all,
            )
        )
        np.savez(
            paths["arrays"],
            g0=g0_all,
            geff=geff_all,
            charge=charge_all,
            discharge=discharge_all,
            emergency=emergency_all,
            unused=unused_all,
            soc=soc_all,
            initial_soc=initial_soc_all,
            terminal_soc=terminal_soc_all,
            boundary_plan=boundary_plan,
            boundary_price_mean=boundary_mean,
        )

    annual_cost = float(sum(float(row["dp_total_cost_yuan"]) for row in daily_rows))
    summary = {
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "config": asdict(config),
        "formal_days": count,
        "annual": {
            "total_cost_yuan": annual_cost,
            "plan_cost_yuan": float(sum(float(row["plan_cost_yuan"]) for row in daily_rows)),
            "effective_base_cost_yuan": float(sum(float(row["effective_base_cost_yuan"]) for row in daily_rows)),
            "adjustment_transaction_yuan": float(sum(float(row["adjustment_transaction_yuan"]) for row in daily_rows)),
            "adjustment_surcharge_yuan": float(sum(float(row["adjustment_surcharge_yuan"]) for row in daily_rows)),
            "emergency_cost_yuan": float(sum(float(row["dp_emergency_cost_yuan"]) for row in daily_rows)),
            "emergency_energy_kwh": float(sum(float(row["dp_emergency_energy_kwh"]) for row in daily_rows)),
            "adjust_abs_energy_kwh": total_delta_abs,
            "terminal_soc_kwh": dp_initial_soc,
            "analytical_terminal_soc_kwh": analytical_initial_soc,
        },
        "validation": {
            "max_billing_identity_residual_yuan": max_billing_identity,
            "max_cross_day_soc_residual_kwh": max_cross_day_soc,
            "max_dp_balance_residual_kwh": max(float(row["dp_max_balance_residual_kwh"]) for row in daily_rows),
            "max_dp_soc_residual_kwh": max(float(row["dp_max_soc_residual_kwh"]) for row in daily_rows),
            "max_dp_simultaneous_product_kwh2": max(float(row["dp_max_simultaneous_product_kwh2"]) for row in daily_rows),
            "max_dp_value_convexity_violation": max_day_convexity,
            "max_planning_balance_residual_kwh": max(float(row["max_balance_residual_kwh"]) for row in solver_rows),
            "max_planning_soc_residual_kwh": max(float(row["max_soc_residual_kwh"]) for row in solver_rows),
            "max_planning_simultaneous_product_kwh2": max(float(row["max_simultaneous_product_kwh2"]) for row in solver_rows),
            "max_planning_delta_product_kwh2": max(float(row["max_delta_product_kwh2"]) for row in solver_rows),
            "all_planning_statuses_optimal": all("Optimal" in str(row["status"]) for row in solver_rows),
        },
        "boundary": {
            "computed": count == FORMAL_DAYS,
            "initial_soc_kwh": dp_initial_soc if count == FORMAL_DAYS else None,
            "t1_plan_kwh": float(boundary_plan[0]) if count == FORMAL_DAYS else None,
            "t1_predicted_scenario_mean_price": float(boundary_mean[0]) if count == FORMAL_DAYS else None,
        },
        "runtime_seconds": time.perf_counter() - started,
    }
    paths["summary"].write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    print(
        f"[{config.name}] 完成：{count}天，总费用={annual_cost:.2f}，"
        f"年末SOC={dp_initial_soc:.3f}，耗时={summary['runtime_seconds']:.1f}s",
        flush=True,
    )
    return summary


def load_main_arrays(path: Path) -> MainRunArrays:
    with np.load(path) as data:
        return MainRunArrays(
            g0=data["g0"].copy(),
            geff=data["geff"].copy(),
            charge=data["charge"].copy(),
            discharge=data["discharge"].copy(),
            emergency=data["emergency"].copy(),
            unused=data["unused"].copy(),
            soc=data["soc"].copy(),
            initial_soc=data["initial_soc"].copy(),
            terminal_soc=data["terminal_soc"].copy(),
            boundary_plan=data["boundary_plan"].copy(),
            boundary_price_mean=data["boundary_price_mean"].copy(),
        )


def validate_template(path: Path, sheets: tuple[str, ...], kind: str) -> tuple[date, ...]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        if tuple(workbook.sheetnames) != sheets:
            raise ValueError(f"{kind}模板工作表异常：{workbook.sheetnames}")
        expected = formal_dates()
        wide_names = ("计划购电量",) if kind == "result4-2" else ("计划购电量", "调整购电量")
        for name in wide_names:
            sheet = workbook[name]
            if sheet.max_row != 335 or sheet.max_column != 147:
                raise ValueError(f"{kind}{name}尺寸异常：{sheet.max_row}×{sheet.max_column}")
            headers = [sheet.cell(1, column).value for column in range(1, 148)]
            if headers[0] != "日期\\时间" or headers[1] != "0:10-0:20":
                raise ValueError(f"{kind}{name}前部表头异常")
            if headers[143:147] != [
                "23:50-0:00+1", "0:00-0:10+1", "全天购电量", "全天购电费"
            ]:
                raise ValueError(f"{kind}{name}尾部表头异常：{headers[143:147]}")
            dates = tuple(
                q2.as_date(sheet.cell(row, 1).value, f"{kind}{name}第{row}行")
                for row in range(2, 336)
            )
            if dates != expected:
                raise ValueError(f"{kind}{name}日期不连续")
        storage = workbook["充放电量"]
        if storage.max_column != 6:
            raise ValueError(f"{kind}充放电表列数异常")
        if [storage.cell(2 + i, 2).value for i in range(6)] != list(FOUR_HOUR_LABELS):
            raise ValueError(f"{kind}充放电样本时段标签异常")
        emergency = workbook["紧急购电量"]
        if emergency.max_column != 3:
            raise ValueError(f"{kind}紧急购电表列数异常")
        return expected
    finally:
        workbook.close()


def apply_row_style(sheet: Any, target_row: int, source_cells: Sequence[Any], height: float | None) -> None:
    for column, source in enumerate(source_cells, start=1):
        target = sheet.cell(target_row, column)
        target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.protection = copy(source.protection)
    sheet.row_dimensions[target_row].height = height


def displayed_price_path(
    day: date,
    prices: PriceData,
    boundary_mean: np.ndarray,
) -> tuple[np.ndarray, tuple[str, ...]]:
    row = prices.date_to_index[day]
    next_day = day + timedelta(days=1)
    if next_day in prices.date_to_index:
        next_t1 = float(prices.values[prices.date_to_index[next_day], 0])
        roles = ("realized",) * T
    else:
        next_t1 = float(boundary_mean[0])
        roles = ("realized",) * (T - 1) + ("boundary_prediction",)
    return np.concatenate((prices.values[row, 1:], [next_t1])), roles


def export_mapping(
    arrays: MainRunArrays,
    dates: Sequence[date],
    index: int,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    next_t1 = (
        float(arrays.g0[index + 1, 0]) if index + 1 < len(dates) else float(arrays.boundary_plan[0])
    )
    displayed_g0 = np.concatenate((arrays.g0[index, 1:], [next_t1]))
    displayed_geff = np.concatenate((arrays.geff[index, 1:], [next_t1]))
    displayed_delta = displayed_geff - displayed_g0
    return displayed_g0, displayed_geff, displayed_delta


def write_result_workbook(
    kind: str,
    template: Path,
    output: Path,
    prices: PriceData,
    arrays: MainRunArrays,
) -> None:
    sheets = RESULT42_SHEETS if kind == "result4-2" else RESULT43_SHEETS
    dates = validate_template(template, sheets, kind)
    if output.resolve() == template.resolve():
        raise ValueError("Q4输出不得覆盖官方模板")
    source_hash = q2.sha256_file(template)
    shutil.copy2(template, output)
    workbook = load_workbook(output, data_only=False, read_only=False)
    try:
        for i, day in enumerate(dates):
            displayed_g0, displayed_geff, displayed_delta = export_mapping(arrays, dates, i)
            display_prices, _ = displayed_price_path(day, prices, arrays.boundary_price_mean)
            plan_sheet = workbook["计划购电量"]
            row = i + 2
            for column, value in enumerate(displayed_g0, start=2):
                plan_sheet.cell(row, column).value = float(value)
            plan_sheet.cell(row, 146).value = float(np.sum(displayed_g0))
            plan_sheet.cell(row, 147).value = float(np.dot(display_prices, displayed_g0))
            if kind == "result4-3":
                adjusted = workbook["调整购电量"]
                for column, value in enumerate(displayed_geff, start=2):
                    adjusted.cell(row, column).value = float(value)
                adjusted.cell(row, 146).value = float(np.sum(displayed_geff))
                # 用户定案：该列只填0.5附加费，不含原计划费和紧急费。
                adjusted.cell(row, 147).value = float(
                    0.5 * np.dot(display_prices, np.abs(displayed_delta))
                )

        storage = workbook["充放电量"]
        styles = {
            block: (
                [copy(storage.cell(2 + block, column)) for column in range(1, 7)],
                storage.row_dimensions[2 + block].height,
            )
            for block in range(6)
        }
        storage.delete_rows(2, storage.max_row - 1)
        for i, day in enumerate(dates):
            for block in range(6):
                target_row = 2 + i * 6 + block
                source_cells, height = styles[block]
                apply_row_style(storage, target_row, source_cells, height)
                storage.cell(target_row, 1).value = (
                    datetime.combine(day, datetime_time()) if block == 0 else None
                )
                storage.cell(target_row, 2).value = FOUR_HOUR_LABELS[block]
                start = block * 24
                stop = (block + 1) * 24
                storage.cell(target_row, 3).value = float(np.sum(arrays.charge[i, start:stop]))
                storage.cell(target_row, 4).value = float(np.sum(arrays.discharge[i, start:stop]))
                storage.cell(target_row, 5).value = datetime_time(0, 0) if block == 0 else ("24:00" if block == 1 else None)
                storage.cell(target_row, 6).value = float(arrays.initial_soc[i]) if block == 0 else (float(arrays.terminal_soc[i]) if block == 1 else None)

        emergency = workbook["紧急购电量"]
        first_style = [copy(emergency.cell(2, column)) for column in range(1, 4)]
        continuation_style = [copy(emergency.cell(3, column)) for column in range(1, 4)]
        first_height = emergency.row_dimensions[2].height
        continuation_height = emergency.row_dimensions[3].height
        emergency.delete_rows(2, emergency.max_row - 1)
        output_row = 2
        for i, day in enumerate(dates):
            intervals = q2.merge_emergency_intervals(arrays.emergency[i])
            if not intervals:
                intervals = [{"physical_interval": "无", "emergency_energy_kwh": 0.0}]
            for j, interval in enumerate(intervals):
                style = first_style if j == 0 else continuation_style
                height = first_height if j == 0 else continuation_height
                apply_row_style(emergency, output_row, style, height)
                emergency.cell(output_row, 1).value = (
                    datetime.combine(day, datetime_time()) if j == 0 else None
                )
                emergency.cell(output_row, 2).value = interval["physical_interval"]
                emergency.cell(output_row, 3).value = float(interval["emergency_energy_kwh"])
                output_row += 1
        workbook.save(output)
    finally:
        workbook.close()
    if q2.sha256_file(template) != source_hash:
        raise RuntimeError("官方Q4模板在导出过程中被修改")


def verify_result_workbook(
    kind: str,
    path: Path,
    prices: PriceData,
    arrays: MainRunArrays,
) -> dict[str, Any]:
    sheets = RESULT42_SHEETS if kind == "result4-2" else RESULT43_SHEETS
    dates = formal_dates()
    workbook = load_workbook(path, data_only=True, read_only=True)
    max_difference = 0.0
    try:
        if tuple(workbook.sheetnames) != sheets:
            raise RuntimeError(f"{kind}输出工作表名称或顺序改变")
        plan_rows = list(
            workbook["计划购电量"].iter_rows(
                min_row=2, max_row=FORMAL_DAYS + 1, values_only=True
            )
        )
        adjusted_rows = (
            list(
                workbook["调整购电量"].iter_rows(
                    min_row=2, max_row=FORMAL_DAYS + 1, values_only=True
                )
            )
            if kind == "result4-3"
            else None
        )
        for i, day in enumerate(dates):
            displayed_g0, displayed_geff, displayed_delta = export_mapping(arrays, dates, i)
            display_prices, _ = displayed_price_path(day, prices, arrays.boundary_price_mean)
            plan_values = plan_rows[i]
            actual = np.asarray(plan_values[1:145], dtype=np.float64)
            max_difference = max(max_difference, float(np.max(np.abs(actual - displayed_g0))))
            if abs(float(plan_values[145]) - float(np.sum(displayed_g0))) > CHECK_TOLERANCE:
                raise RuntimeError(f"{kind} {day}计划表全天购电量错误")
            if abs(float(plan_values[146]) - float(np.dot(display_prices, displayed_g0))) > CHECK_TOLERANCE:
                raise RuntimeError(f"{kind} {day}计划表展示窗口费用错误")
            if adjusted_rows is not None:
                values = adjusted_rows[i]
                actual_eff = np.asarray(values[1:145], dtype=np.float64)
                max_difference = max(max_difference, float(np.max(np.abs(actual_eff - displayed_geff))))
                expected_fee = float(0.5 * np.dot(display_prices, np.abs(displayed_delta)))
                if abs(float(values[145]) - float(np.sum(displayed_geff))) > CHECK_TOLERANCE:
                    raise RuntimeError(f"{kind} {day}调整表全天购电量错误")
                if abs(float(values[146]) - expected_fee) > CHECK_TOLERANCE:
                    raise RuntimeError(f"{kind} {day}调整表0.5附加费错误")

        storage = workbook["充放电量"]
        if storage.max_row != 1 + FORMAL_DAYS * 6 or storage.max_column != 6:
            raise RuntimeError(f"{kind}充放电量未按334天×6段展开")
        storage_rows = list(storage.iter_rows(min_row=2, values_only=True))
        for i, day in enumerate(dates):
            base = i * 6
            if q2.as_date(storage_rows[base][0], f"{kind}充放电日期") != day:
                raise RuntimeError(f"{kind} {day}充放电首行日期错误")
            for block in range(6):
                row = storage_rows[base + block]
                if row[1] != FOUR_HOUR_LABELS[block]:
                    raise RuntimeError(f"{kind} {day}充放电物理时段错误")
                start = block * 24
                stop = (block + 1) * 24
                expected_c = float(np.sum(arrays.charge[i, start:stop]))
                expected_d = float(np.sum(arrays.discharge[i, start:stop]))
                max_difference = max(
                    max_difference, abs(float(row[2]) - expected_c), abs(float(row[3]) - expected_d)
                )
                if block > 0 and row[0] is not None:
                    raise RuntimeError(f"{kind} {day}充放电续行不应重复日期")
            if storage_rows[base][4] != datetime_time(0, 0) or storage_rows[base + 1][4] != "24:00":
                raise RuntimeError(f"{kind} {day}SOC时刻标签错误")
            if abs(float(storage_rows[base][5]) - float(arrays.initial_soc[i])) > CHECK_TOLERANCE:
                raise RuntimeError(f"{kind} {day}期初SOC错误")
            if abs(float(storage_rows[base + 1][5]) - float(arrays.terminal_soc[i])) > CHECK_TOLERANCE:
                raise RuntimeError(f"{kind} {day}期末SOC错误")

        emergency = workbook["紧急购电量"]
        rows = list(emergency.iter_rows(min_row=2, values_only=True))
        cursor = 0
        days_with_emergency = 0
        for i, day in enumerate(dates):
            intervals = q2.merge_emergency_intervals(arrays.emergency[i])
            if not intervals:
                intervals = [{"physical_interval": "无", "emergency_energy_kwh": 0.0}]
            else:
                days_with_emergency += 1
            for j, expected in enumerate(intervals):
                if cursor >= len(rows):
                    raise RuntimeError(f"{kind}紧急购电输出行数不足")
                row = rows[cursor]
                if j == 0 and q2.as_date(row[0], f"{kind}紧急购电日期") != day:
                    raise RuntimeError(f"{kind} {day}紧急购电首行日期错误")
                if j > 0 and row[0] is not None:
                    raise RuntimeError(f"{kind} {day}紧急购电续行日期应为空")
                if row[1] != expected["physical_interval"]:
                    raise RuntimeError(f"{kind} {day}紧急购电物理区间错误")
                if abs(float(row[2]) - float(expected["emergency_energy_kwh"])) > CHECK_TOLERANCE:
                    raise RuntimeError(f"{kind} {day}紧急购电量错误")
                cursor += 1
        if cursor != len(rows):
            raise RuntimeError(f"{kind}紧急购电输出有多余行")
    finally:
        workbook.close()
    if max_difference > CHECK_TOLERANCE:
        raise RuntimeError(f"{kind}回读最大差异超限：{max_difference}")
    return {
        "sheet_names_preserved": True,
        "wide_rows": FORMAL_DAYS,
        "wide_physical_mapping": "date d physical t=2..144 plus date d+1 physical t=1",
        "storage_rows": FORMAL_DAYS * 6,
        "storage_all_days_filled": True,
        "emergency_all_days_filled": True,
        "emergency_days_with_purchase": days_with_emergency,
        "december_31_boundary_prediction_used": True,
        "max_roundtrip_difference": max_difference,
        "adjusted_fee_column": (
            "0.5*sum(display_price*abs(display_geff-display_g0)); excludes original and emergency"
            if kind == "result4-3" else None
        ),
    }


def combine_main_csvs(out: Path, filename: str, output_name: str) -> None:
    sources = [out / f"q4-2_{filename}.csv", out / f"q4-3_{filename}.csv"]
    rows: list[dict[str, str]] = []
    fields: list[str] | None = None
    for source in sources:
        source_rows = read_csv_rows(source)
        if source_rows:
            if fields is None:
                fields = list(source_rows[0].keys())
            elif fields != list(source_rows[0].keys()):
                raise RuntimeError(f"合并{filename}时列不一致")
            rows.extend(source_rows)
    if fields is None:
        raise RuntimeError(f"没有可合并的{filename}")
    write_csv(out / output_name, rows, fields)


def monthly_rows(out: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for branch in BRANCHES:
        metrics = read_csv_rows(out / f"q4-{branch[-1]}_daily_metrics.csv")
        for month in range(2, 13):
            selected = [row for row in metrics if date.fromisoformat(row["date"]).month == month]
            rows.append(
                {
                    "branch": branch,
                    "month": month,
                    "days": len(selected),
                    "total_cost_yuan": sum(float(row["dp_total_cost_yuan"]) for row in selected),
                    "plan_cost_yuan": sum(float(row["plan_cost_yuan"]) for row in selected),
                    "adjustment_transaction_yuan": sum(float(row["adjustment_transaction_yuan"]) for row in selected),
                    "adjustment_surcharge_yuan": sum(float(row["adjustment_surcharge_yuan"]) for row in selected),
                    "emergency_cost_yuan": sum(float(row["dp_emergency_cost_yuan"]) for row in selected),
                    "emergency_energy_kwh": sum(float(row["dp_emergency_energy_kwh"]) for row in selected),
                    "terminal_soc_kwh": float(selected[-1]["dp_terminal_soc_kwh"]),
                }
            )
    return rows


def write_paper_tables(
    out: Path,
    prices: PriceData,
    arrays42: MainRunArrays,
    arrays43: MainRunArrays,
) -> None:
    dates = formal_dates()
    six_rows: list[dict[str, Any]] = []
    for branch, arrays in (("4-2", arrays42), ("4-3", arrays43)):
        for i, day in enumerate(dates):
            display_g0, display_geff, display_delta = export_mapping(arrays, dates, i)
            display_prices, _ = displayed_price_path(day, prices, arrays.boundary_price_mean)
            for plan_type, displayed in (
                ("original", display_g0),
                *(((("effective", display_geff),)) if branch == "4-3" else ()),
            ):
                row: dict[str, Any] = {
                    "branch": branch,
                    "date": day.isoformat(),
                    "plan_type": plan_type,
                }
                for model_t, label in PAPER_INTERVALS:
                    # 模板列号=物理模型t，因为第2列从t=2开始；按标签再次断言。
                    column = model_t
                    expected_header_position = model_t - 2
                    row[label] = float(displayed[expected_header_position])
                    row[f"{label}_template_column"] = column
                row["display_window_energy_kwh"] = float(np.sum(displayed))
                row["display_window_plan_cost_yuan"] = float(np.dot(display_prices, display_g0))
                row["display_window_adjustment_surcharge_yuan"] = float(
                    0.5 * np.dot(display_prices, np.abs(display_delta))
                ) if branch == "4-3" else 0.0
                six_rows.append(row)
    six_fields = list(six_rows[0].keys())
    write_csv(out / "q4_paper_six_intervals.csv", six_rows, six_fields)

    emergency_rows: list[dict[str, Any]] = []
    date_index = {day: i for i, day in enumerate(dates)}
    for branch, arrays in (("4-2", arrays42), ("4-3", arrays43)):
        for day in PAPER_EMERGENCY_DATES:
            intervals = q2.merge_emergency_intervals(arrays.emergency[date_index[day]])
            if not intervals:
                intervals = [{"physical_interval": "无", "emergency_energy_kwh": 0.0}]
            for rank, interval in enumerate(intervals, start=1):
                emergency_rows.append(
                    {
                        "branch": branch,
                        "date": day.isoformat(),
                        "interval_rank": rank,
                        "physical_interval": interval["physical_interval"],
                        "emergency_energy_kwh": interval["emergency_energy_kwh"],
                    }
                )
    write_csv(
        out / "q4_paper_emergency_dates.csv",
        emergency_rows,
        ["branch", "date", "interval_rank", "physical_interval", "emergency_energy_kwh"],
    )


def write_template_mapping_audit(
    out: Path,
    prices: PriceData,
    arrays42: MainRunArrays,
    arrays43: MainRunArrays,
) -> None:
    rows: list[dict[str, Any]] = []
    dates = formal_dates()
    labels = q2.physical_interval_labels()
    for branch, arrays in (("4-2", arrays42), ("4-3", arrays43)):
        for di, day in enumerate(dates):
            display_prices, roles = displayed_price_path(day, prices, arrays.boundary_price_mean)
            g0, geff, delta = export_mapping(arrays, dates, di)
            for j in range(T):
                if j < T - 1:
                    physical_day = day
                    physical_t = j + 2
                else:
                    physical_day = day + timedelta(days=1)
                    physical_t = 1
                rows.append(
                    {
                        "branch": branch,
                        "template_row_date": day.isoformat(),
                        "template_column": j + 2,
                        "template_header_physical_interval": (
                            f"{(j + 1) * 10 // 60}:{(j + 1) * 10 % 60:02d}-"
                            f"{(j + 2) * 10 // 60}:{(j + 2) * 10 % 60:02d}"
                            if j < T - 2
                            else ("23:50-0:00+1" if j == T - 2 else "0:00-0:10+1")
                        ),
                        "physical_date": physical_day.isoformat(),
                        "physical_t": physical_t,
                        "physical_interval": labels[physical_t - 1],
                        "price": float(display_prices[j]),
                        "price_role": roles[j],
                        "g0_kwh": float(g0[j]),
                        "geff_kwh": float(geff[j]),
                        "adjustment_surcharge_yuan": float(0.5 * display_prices[j] * abs(delta[j])),
                    }
                )
    write_csv(out / "q4_template_mapping_audit.csv", rows, list(rows[0].keys()))


def reprice_existing_baselines(args: argparse.Namespace, prices: PriceData) -> list[dict[str, Any]]:
    # 优先使用终值系数修复后（v=凌晨均价/eta_c=0.481548）的Q2/Q3重跑动作；
    # 旧目录仅作回退并在note中标记。
    candidates = [
        (
            "Q2_existing_actions_terminal_fix",
            args.root / "output/q2_terminal_fix/q2_interval_details.csv",
            args.root / "output/q2/q2_interval_details.csv",
            "4-2",
        ),
        (
            "Q3_existing_actions_terminal_fix",
            args.root / "output/q3_terminal_fix/q3_all_interval_details.csv",
            args.root / "output/q3/q3_all_interval_details.csv",
            "4-3",
        ),
    ]
    rows: list[dict[str, Any]] = []
    for label, primary_path, fallback_path, branch in candidates:
        path = primary_path if primary_path.is_file() else fallback_path
        note = "baseline only; not a bound"
        if not path.is_file():
            rows.append(
                {
                    "baseline": label,
                    "branch_contract": branch,
                    "status": "skipped_missing_file",
                    "source_path": q2.relative_path(primary_path, args.root),
                    "repriced_total_cost_yuan": "",
                    "note": note,
                }
            )
            continue
        if path.resolve() != primary_path.resolve():
            note = "baseline only; not a bound; fallback to pre-fix actions"
        source = read_csv_rows(path)
        total = 0.0
        for row in source:
            day = date.fromisoformat(row["date"])
            t = int(row["t"])
            price = float(prices.values[prices.date_to_index[day], t - 1])
            if branch == "4-2":
                total += price * (
                    float(row["plan_grid_kwh"]) + EMERGENCY_PRICE_FACTOR * float(row["dp_emergency_kwh"])
                )
            else:
                g0 = float(row["g0_kwh"])
                geff = float(row["geff_kwh"])
                emergency = float(row["dp_emergency_kwh"])
                total += price * (geff + 0.5 * abs(geff - g0) + EMERGENCY_PRICE_FACTOR * emergency)
        rows.append(
            {
                "baseline": label,
                "branch_contract": branch,
                "status": "repriced",
                "source_path": q2.relative_path(path, args.root),
                "repriced_total_cost_yuan": total,
                "note": note,
            }
        )
    return rows


def aggregate_config_comparison(out: Path, config_names: Sequence[str]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for name in config_names:
        config = RUN_CONFIGS[name]
        path = config_paths(out, config)["summary"]
        if not path.is_file():
            continue
        summary = json.loads(path.read_text(encoding="utf-8"))
        rows.append(
            {
                "config": name,
                "branch": config.branch,
                "terminal_mode": config.terminal_mode,
                "adjust_purchases": config.adjust_purchases,
                "refresh_value": config.refresh_value,
                "update_hours": ",".join(map(str, config.update_hours)),
                "formal_days": summary["formal_days"],
                "total_cost_yuan": summary["annual"]["total_cost_yuan"],
                "terminal_soc_kwh": summary["annual"]["terminal_soc_kwh"],
                "emergency_energy_kwh": summary["annual"]["emergency_energy_kwh"],
                "adjust_abs_energy_kwh": summary["annual"]["adjust_abs_energy_kwh"],
                "official_output": config.main_output,
                "description": config.description,
            }
        )
    return rows


def assemble_outputs(args: argparse.Namespace, count: int, config_names: Sequence[str]) -> dict[str, Any]:
    out = get_output_dir(args)
    if count != FORMAL_DAYS:
        comparison = aggregate_config_comparison(out, config_names)
        if comparison:
            write_csv(out / "q4_configuration_comparison.csv", comparison, list(comparison[0].keys()))
        report = {
            "run_type": "partial_smoke_test",
            "formal_days": count,
            "official_workbooks_generated": False,
            "reason": "正式Excel只在334日完整运行后生成",
            "configurations": comparison,
        }
        (out / "q4_run_summary.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
            encoding="utf-8",
        )
        return report

    for main_name in ("42-main", "43-main"):
        if not configuration_complete(out, RUN_CONFIGS[main_name], count):
            raise RuntimeError(f"缺少正式配置{main_name}，无法组装结果")
    prices = read_price_data(input_path(args, "prices"))
    annual = q2.read_annual_data(input_path(args, "annual"))
    validate_inputs(annual, prices)
    arrays42 = load_main_arrays(config_paths(out, RUN_CONFIGS["42-main"])["arrays"])
    arrays43 = load_main_arrays(config_paths(out, RUN_CONFIGS["43-main"])["arrays"])
    template42 = input_path(args, "template42")
    template43 = input_path(args, "template43")
    result42 = out / "result4-2.xlsx"
    result43 = out / "result4-3.xlsx"
    write_result_workbook("result4-2", template42, result42, prices, arrays42)
    write_result_workbook("result4-3", template43, result43, prices, arrays43)
    validation42 = verify_result_workbook("result4-2", result42, prices, arrays42)
    validation43 = verify_result_workbook("result4-3", result43, prices, arrays43)

    combine_main_csvs(out, "plan_versions", "q4_plan_versions.csv")
    combine_main_csvs(out, "interval_details", "q4_interval_details.csv")
    monthly = monthly_rows(out)
    write_csv(out / "q4_monthly_metrics.csv", monthly, list(monthly[0].keys()))
    comparison = aggregate_config_comparison(out, config_names)
    write_csv(out / "q4_configuration_comparison.csv", comparison, list(comparison[0].keys()))
    baselines = reprice_existing_baselines(args, prices)
    write_csv(
        out / "q4_repriced_baselines.csv",
        baselines,
        ["baseline", "branch_contract", "status", "source_path", "repriced_total_cost_yuan", "note"],
    )
    write_paper_tables(out, prices, arrays42, arrays43)
    write_template_mapping_audit(out, prices, arrays42, arrays43)

    prep = json.loads((out / CACHE_FILES["metadata"]).read_text(encoding="utf-8"))
    main42 = json.loads(config_paths(out, RUN_CONFIGS["42-main"])["summary"].read_text(encoding="utf-8"))
    main43 = json.loads(config_paths(out, RUN_CONFIGS["43-main"])["summary"].read_text(encoding="utf-8"))
    output_files = [
        result42, result43,
        out / "q4_price_forecast_audit.csv",
        out / "q4_scenario_audit.csv",
        out / "q4_plan_versions.csv",
        out / "q4_interval_details.csv",
        out / "q4-2_daily_metrics.csv",
        out / "q4-3_daily_metrics.csv",
        out / "q4_monthly_metrics.csv",
        out / "q4_configuration_comparison.csv",
        out / "q4_repriced_baselines.csv",
        out / "q4_paper_six_intervals.csv",
        out / "q4_paper_emergency_dates.csv",
        out / "q4_template_mapping_audit.csv",
        out / "q4_price_model_selection.csv",
        out / "q4_price_forecast_metrics.csv",
    ]
    report = {
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "run_type": "full",
        "formal_days": FORMAL_DAYS,
        "status": "success",
        "errors": [],
        "parameters": {
            "periods_per_day": T,
            "delta_t_hours": q2.DELTA_T_HOURS,
            "eta_charge": ETA_CHARGE,
            "eta_discharge": ETA_DISCHARGE,
            "soc_min_kwh": SOC_MIN_KWH,
            "soc_max_kwh": SOC_MAX_KWH,
            "initial_soc_2025_02_01_kwh": INITIAL_SOC_KWH,
            "max_action_kwh_per_period": MAX_ACTION_KWH,
            "price_history_days": PRICE_HISTORY_DAYS,
            "max_scenarios": MAX_SCENARIOS,
            "selected_price_model_family_shared": prep["selected_price_model_family"],
            "selected_short_rho_shared": prep["selected_short_rho"],
            "formal_terminal_value": "next 00:00-05:00 window scenario mean / eta_c",
            "terminal_efficiency_basis": "charge",
            "terminal_price_window": "next physical 00:00-05:00 after issue time",
            "terminal_price_slot_count": OVERNIGHT_WINDOW_SLOTS,
            "terminal_sensitivity": ["zero", "eta_d*minimum", "eta_d*median"],
            "q2_q3_fix_alignment": (
                "q2.terminal_value_coefficient shared implementation checked in self-test; "
                "Q4 formal terminal proxy uses the same charge-margin overnight-window structure "
                "with Q4 scenario prices instead of attachment-1 prices"
            ),
            "result_wide_mapping": "date d physical t=2..144 plus date d+1 physical t=1",
            "result43_adjusted_fee_column": "0.5*sum(price*abs(geff-g0)); excludes original and emergency",
        },
        "information_sets": {
            "4-2_attachment3_read": False,
            "4-3_attachment3_read": True,
            "price_reveal": "each interval price revealed immediately before execution",
            "future_actual_price_in_planning": False,
        },
        "price_forecast": prep,
        "main_results": {
            "4-2": main42["annual"],
            "4-3": main43["annual"],
            "difference_4-2_minus_4-3_yuan": main42["annual"]["total_cost_yuan"] - main43["annual"]["total_cost_yuan"],
        },
        "configuration_comparison": comparison,
        "repriced_existing_action_baselines": baselines,
        "validation": {
            "4-2_simulation": main42["validation"],
            "4-3_simulation": main43["validation"],
            "result4-2": validation42,
            "result4-3": validation43,
            "daily_bill_sum_matches_annual": True,
            "scenario_close_times_checked": True,
            "all_checks_passed": True,
        },
        "inputs": {
            key: {
                "path": q2.relative_path(input_path(args, key), args.root),
                "sha256": q2.sha256_file(input_path(args, key)),
            }
            for key in DATA_PATHS
        },
        "versions": q2.package_versions(),
        "outputs": {
            path.name: {
                "path": q2.relative_path(path, args.root),
                "sha256": q2.sha256_file(path),
            }
            for path in output_files
        },
        "warnings": [
            "模板宽表按其文字的物理时间展示跨日24小时窗口；其金额不等于自然日账单。",
            "result4-3调整购电量表费用列按用户指定仅为0.5调整附加费；完整账单见daily metrics。",
            "终值敏感性（zero/eta_d*min/eta_d*median）不用于事后切换正式凌晨窗口方案。",
            "正式终值代理已按Q2/Q3修复对齐为充电边际（凌晨窗口场景均价/eta_c），需在论文说明与orderForQ4暂设eta_d*min的差异。",
        ],
    }
    (out / "q4_run_summary.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    print("Q4正式结果已生成并回读验收通过", flush=True)
    print(f"  result4-2: {result42}", flush=True)
    print(f"  result4-3: {result43}", flush=True)
    print(
        f"  4-2总费用={main42['annual']['total_cost_yuan']:.2f}, "
        f"4-3总费用={main43['annual']['total_cost_yuan']:.2f}",
        flush=True,
    )
    return report


def verify_existing(args: argparse.Namespace) -> dict[str, Any]:
    out = get_output_dir(args)
    prices = read_price_data(input_path(args, "prices"))
    arrays42 = load_main_arrays(config_paths(out, RUN_CONFIGS["42-main"])["arrays"])
    arrays43 = load_main_arrays(config_paths(out, RUN_CONFIGS["43-main"])["arrays"])
    result = {
        "result4-2": verify_result_workbook("result4-2", out / "result4-2.xlsx", prices, arrays42),
        "result4-3": verify_result_workbook("result4-3", out / "result4-3.xlsx", prices, arrays43),
    }
    print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
    return result


def self_test(args: argparse.Namespace) -> None:
    labels = q2.physical_interval_labels()
    assert labels[0] == "0:00-0:10"
    assert labels[60] == "10:00-10:10"
    assert labels[-1] == "23:50-00:00+1"
    day_positions, t_abs = q3.window_positions(40, 18)
    assert np.all(day_positions[:36] == 40) and np.all(day_positions[36:] == 41)
    np.testing.assert_array_equal(t_abs[:36], np.arange(109, 145))
    np.testing.assert_array_equal(t_abs[36:], np.arange(1, 109))

    synthetic_dates = tuple(date(2025, 1, 1) + timedelta(days=i) for i in range(20))
    x = np.linspace(0.4, 0.8, 20)
    values = np.stack([np.full(T, value) for value in x])
    synthetic = PriceData(
        synthetic_dates,
        tuple("" for _ in range(T)),
        values,
        {day: i for i, day in enumerate(synthetic_dates)},
        np.mean(values, axis=1),
    )
    fit = fit_ar_level(date(2025, 1, 21), synthetic, synthetic_dates[-10:])
    assert 0.0 <= float(fit["coefficient"]) <= 0.99
    shape = normalized_shape(synthetic, synthetic_dates[-10:])
    np.testing.assert_allclose(shape, 1.0, atol=1.0e-12)

    # Q2/Q3终值修复的共享实现与验收断言（orderForQ2rivise/orderForQ3rivise §6）
    attachment1_prices, attachment1_labels = q2.read_prices(args.root / "data/附件/附件1.xlsx")
    assert len(attachment1_prices) == T
    assert attachment1_labels[0] in ("00:10", "0:10")
    assert np.isclose(np.mean(attachment1_prices[:30]), 0.43339333333333335, atol=1e-12)
    fixed_coefficient = q2.terminal_value_coefficient(attachment1_prices)
    assert np.isclose(fixed_coefficient, 0.4815481481481482, atol=1e-12)
    assert not np.isclose(fixed_coefficient, 0.33417, atol=1e-8)
    assert not hasattr(q3, "TERMINAL_VALUE_COEFFICIENT")

    mean = np.arange(1.0, T + 1.0) * 0.01  # 严格递增，便于定位窗口
    value_zero, ref_zero = terminal_value_coefficient(mean, "zero", 0)
    assert value_zero == 0.0 and ref_zero == 0.0
    value_min, ref_min = terminal_value_coefficient(mean, "minimum", 0)
    assert math.isclose(ref_min, 0.01) and math.isclose(value_min, ETA_DISCHARGE * 0.01)
    value_med, _ = terminal_value_coefficient(mean, "median", 0)
    assert math.isclose(value_med, ETA_DISCHARGE * float(np.median(mean)))
    # 正式口径：凌晨窗口场景均价/eta_c。h=0取窗口0..29，h>0取次日t=1..30。
    for hour, expected_slice in (
        (0, slice(0, 30)),
        (6, slice(108, 138)),
        (12, slice(72, 102)),
        (18, slice(36, 66)),
    ):
        expected_reference = float(np.mean(mean[expected_slice]))
        value, reference = terminal_value_coefficient(mean, "overnight", hour)
        assert math.isclose(reference, expected_reference)
        assert math.isclose(value, expected_reference / ETA_CHARGE)
    np.testing.assert_array_equal(
        overnight_window_positions(0), np.arange(0, 30)
    )
    np.testing.assert_array_equal(
        overnight_window_positions(18), np.arange(36, 66)
    )

    rng = np.random.default_rng(9)
    p = rng.uniform(0.1, 1.5, T)
    g0 = rng.uniform(0.0, 100.0, T)
    geff = rng.uniform(0.0, 100.0, T)
    emergency = rng.uniform(0.0, 10.0, T)
    delta = geff - g0
    left = float(
        np.dot(p, g0)
        + np.dot(p, 1.5 * np.maximum(delta, 0.0) - 0.5 * np.maximum(-delta, 0.0))
        + np.dot(5.0 * p, emergency)
    )
    right = float(np.dot(p, geff + 0.5 * np.abs(delta) + 5.0 * emergency))
    assert abs(left - right) <= 1.0e-9
    surcharge_only = float(0.5 * np.dot(p, np.abs(delta)))
    assert surcharge_only >= 0.0 and surcharge_only < left

    residual = ResidualEvent(
        "4-2", date(2025, 1, 1), 6,
        issue_datetime(date(2025, 1, 1), 6) + timedelta(hours=24),
        np.zeros(T), np.zeros(T), np.zeros(T),
    )
    assert not recent_closed_residuals([residual], issue_datetime(date(2025, 1, 2), 5))
    assert recent_closed_residuals([residual], issue_datetime(date(2025, 1, 2), 6))

    scenario_net = np.vstack((np.full(T, 100.0), np.full(T, 110.0)))
    plan = np.full(T, 105.0)
    scenario_price = np.vstack((np.full(T, 0.2), np.full(T, 1.2)))
    grid = q2.state_grid(100.0)
    values_dp, violation = average_scenario_value_functions(
        scenario_net, plan, scenario_price, 0.1, grid
    )
    assert values_dp.shape == (T + 1, grid.size)
    assert violation <= 1.0e-7

    # 安全候选集：无约束最小值远低于可行区间时仍必须全部位于[lower, upper]内
    grid = q2.state_grid(100.0)
    transformed = (grid - 1200.0) ** 2  # 最小值在1200
    for lower, upper in ((5000.0, 6000.0), (7000.0, 7500.0)):
        candidates = safe_candidate_minimizers(
            transformed, grid, float(lower), float(upper), 6200.0
        )
        assert candidates
        assert all(lower - 1.0e-9 <= value <= upper + 1.0e-9 for value in candidates)
    assert min(safe_candidate_minimizers(transformed, grid, 8000.0, 8500.0, 8200.0)) >= 8000.0 - 1.0e-9
    assert max(safe_candidate_minimizers(transformed, grid, 8000.0, 8500.0, 8200.0)) <= 8500.0 + 1.0e-9

    template42 = input_path(args, "template42")
    template43 = input_path(args, "template43")
    validate_template(template42, RESULT42_SHEETS, "result4-2")
    validate_template(template43, RESULT43_SHEETS, "result4-3")
    print("Q4 self-test全部通过", flush=True)


def selected_config_names(main_only: bool) -> tuple[str, ...]:
    if main_only:
        return ("42-main", "43-main")
    return tuple(RUN_CONFIGS)


def main() -> None:
    args = parse_args()
    if args.max_days is None:
        count = FORMAL_DAYS
    else:
        if not 1 <= args.max_days <= FORMAL_DAYS:
            raise SystemExit("--max-days必须位于1..334")
        count = args.max_days
    if args.self_test:
        self_test(args)
        return
    if args.verify:
        verify_existing(args)
        return
    if args.prep:
        run_prep(args, count)
        return
    if args.run_config is not None:
        run_configuration(args, RUN_CONFIGS[args.run_config], count)
        return
    config_names = selected_config_names(args.main_only)
    if args.assemble:
        assemble_outputs(args, count, config_names)
        return
    # 默认行为与--all相同。
    run_prep(args, count)
    for name in config_names:
        run_configuration(args, RUN_CONFIGS[name], count)
    assemble_outputs(args, count, config_names)


if __name__ == "__main__":
    main()
