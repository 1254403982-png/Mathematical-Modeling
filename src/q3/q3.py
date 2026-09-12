"""问题3：日内滚动调度主程序（自包含，依据 order/orderForQ3.md v4）。

三阶段：
  A  python q3.py --prep                     一次性生成预测/场景缓存与Bates-Granger校准
  B  python q3.py --scheme all|-6|-12|-18|0  单方案全年仿真（五个方案可并行）
  C  python q3.py --assemble                 由S_all结果回填result3.xlsx并出汇总

时间口径：附件1/2/3按位置映射（附件第i个零基数据位置对应模型t=i+1），时间标签为
区间右端点；result3表格按物理时间填写；充放电量与紧急购电量334天全量展开
（每天行数不限、无省略号）。

更新时刻 h∈{0,6,12,18}，窗口144段；h=0 解原始计划g⁰；h>0 相对g⁰调整未执行时段
（a−g⁰=δ⁺−δ⁻），跨日部分g̃仅作前瞻；g^eff=t≤36←g⁰、37..72←a⁶、73..108←a¹²、
109..144←a¹⁸。执行器：DP（10kWh网格价值函数，主方案）与解析（基线）。
计费：C=Σp·g⁰ + Σp(1.5Δ⁺−0.5Δ⁻) + Σ5p·b，与等价式 Σp[g^eff+0.5|Δ|+5b] 对账。
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import shutil
import sys
import time
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, time as datetime_time, timedelta
from pathlib import Path
from typing import Any, Sequence

import numpy as np
import pulp
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import q2  # noqa: E402

T = q2.T
DELTA_T_HOURS = q2.DELTA_T_HOURS
ETA_CHARGE = q2.ETA_CHARGE
ETA_DISCHARGE = q2.ETA_DISCHARGE
SOC_MIN_KWH = q2.SOC_MIN_KWH
SOC_MAX_KWH = q2.SOC_MAX_KWH
INITIAL_SOC_KWH = q2.INITIAL_SOC_KWH
MAX_ACTION_KWH = q2.MAX_ACTION_KWH
EMERGENCY_PRICE_FACTOR = q2.EMERGENCY_PRICE_FACTOR
MAX_SCENARIOS = q2.MAX_SCENARIOS
CHECK_TOLERANCE = q2.CHECK_TOLERANCE
SIMULTANEOUS_PRODUCT_TOLERANCE = q2.SIMULTANEOUS_PRODUCT_TOLERANCE
ACTION_ZERO_TOLERANCE = q2.ACTION_ZERO_TOLERANCE
FORECAST_START_DATE = q2.FORECAST_START_DATE
FORECAST_END_DATE = q2.FORECAST_END_DATE
BOUNDARY_PLAN_DATE = q2.BOUNDARY_PLAN_DATE
FOUR_HOUR_LABELS = q2.FOUR_HOUR_LABELS

ISSUE_HOURS = (0, 6, 12, 18)
HOUR_INDEX = {0: 0, 6: 1, 12: 2, 18: 3}
KAPPA = {0: 0, 6: 36, 12: 72, 18: 108}
DP_GRID_STEP_KWH = 10.0
ADJUST_UP_PRICE_FACTOR = 1.5
ADJUST_DOWN_PRICE_FACTOR = 0.5
CALIBRATION_SCOPE_END = date(2025, 1, 31)
FORMAL_DAYS = 334
RESULT3_SHEETS = ("计划购电量", "调整购电量", "充放电量", "紧急购电量")
SCHEME_UPDATE_HOURS = {
    "0": (0,),
    "all": (0, 6, 12, 18),
    "-6": (0, 12, 18),
    "-12": (0, 6, 18),
    "-18": (0, 6, 12),
}
SCHEME_NAMES = ("0", "all", "-6", "-12", "-18")
SCHEME_LABELS = {"0": "S0", "all": "S_all", "-6": "S_-6", "-12": "S_-12", "-18": "S_-18"}
SCHEME_MISSING_HOUR = {"0": (6, 12, 18), "all": (), "-6": (6,), "-12": (12,), "-18": (18,)}

DATA_PATHS = {
    "附件1": "data/附件/附件1.xlsx",
    "附件2": "data/附件/附件2.xlsx",
    "附件3": "data/附件/附件3.xlsx",
    "result3模板": "data/附件/附件5/result3.xlsx",
}


@dataclass(frozen=True)
class Attachment3:
    forecasts: dict[tuple[date, int], np.ndarray]
    row_numbers: dict[tuple[date, int], int]


@dataclass
class PendingEvent:
    """一次发布事件(d,h)的窗口预测；窗口完整实现后转为残差进入场景池。"""

    day_idx: int
    hour: int
    day: np.ndarray  # (144,) 各窗口段的绝对日期索引（annual内）
    t_abs: np.ndarray  # (144,) 1-based绝对时段
    load_forecast: np.ndarray  # (144,) 负荷窗口预测
    pv_formal: np.ndarray  # (144,) 附件3正式预报路径
    pv_historical: np.ndarray  # (144,) 历史光伏窗口预测
    pv_combined: np.ndarray | None = None  # Bates-Granger组合（w定后填）


@dataclass(frozen=True)
class UpdateResult:
    issue_date: date
    issue_hour: int
    kind: str  # 'original' | 'rolling'
    z: np.ndarray  # (144,) 窗口共享购电计划（0:00为g⁰；h>0为[a; g̃]）
    delta_plus: np.ndarray  # (144,) 相对g⁰上调（跨日部分补0）
    delta_minus: np.ndarray  # (144,) 相对g⁰下调（跨日部分补0）
    objective_value: float
    scenario_count: int
    status: str
    solve_seconds: float
    tiebreak_used: bool
    binary_fallback_used: bool
    max_balance_residual_kwh: float
    max_soc_residual_kwh: float
    max_delta_product_kwh2: float
    max_simultaneous_product_kwh2: float
    highs_max_primal_infeasibility: float
    simplex_iterations: int
    ipm_iterations: int
    mip_gap: float | None


@dataclass(frozen=True)
class ExportDay3:
    """回填result3所需的一整天数据（与求解结构解耦，便于从CSV重建）。"""

    plan_grid: np.ndarray  # g⁰ (144,)
    eff_grid: np.ndarray  # g^eff (144,)
    dp_charge: np.ndarray
    dp_discharge: np.ndarray
    dp_emergency: np.ndarray
    dp_initial_soc_kwh: float
    dp_terminal_soc_kwh: float


@dataclass
class WindowModel:
    problem: pulp.LpProblem
    original_objective: pulp.LpAffineExpression
    z: list[pulp.LpVariable]
    charge: list[list[pulp.LpVariable]]
    discharge: list[list[pulp.LpVariable]]
    emergency: list[list[pulp.LpVariable]]
    unused: list[list[pulp.LpVariable]]
    soc: list[list[pulp.LpVariable]]
    delta_plus: list[pulp.LpVariable]
    delta_minus: list[pulp.LpVariable]
    ncur: int
    mode: list[list[pulp.LpVariable]] | None


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Q3日内滚动调度")
    parser.add_argument("--prep", action="store_true", help="阶段A：预测/场景缓存与校准")
    parser.add_argument("--scheme", choices=SCHEME_NAMES, default=None, help="阶段B：单方案仿真")
    parser.add_argument("--assemble", action="store_true", help="阶段C：回填result3并出汇总")
    parser.add_argument("--limit-days", type=int, default=None, help="冒烟测试：仅处理前N个正式日")
    parser.add_argument(
        "--root", type=Path, default=Path(__file__).resolve().parent.parent
    )
    parser.add_argument("--out", type=Path, default=None)
    return parser.parse_args()


def input_path(root: Path, key: str) -> Path:
    return root / DATA_PATHS[key]


def output_dir(args: argparse.Namespace) -> Path:
    out = args.out if args.out is not None else args.root / "output" / "q3"
    out.mkdir(parents=True, exist_ok=True)
    return out


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as file:
        return list(csv.DictReader(file))


# ---------------------------------------------------------------------------
# 附件3读取与窗口映射
# ---------------------------------------------------------------------------


def read_attachment3(path: Path) -> Attachment3:
    if not path.is_file():
        raise FileNotFoundError(f"找不到附件3：{path}")
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if tuple(workbook.sheetnames) != ("Sheet1",):
            raise ValueError(f"附件3工作表应为Sheet1，实际为{workbook.sheetnames}")
        worksheet = workbook["Sheet1"]
        rows = list(worksheet.iter_rows(min_row=2, values_only=True))
        if len(rows) != 365 * 4:
            raise ValueError(f"附件3数据行数应为1460，实际为{len(rows)}")
        forecasts: dict[tuple[date, int], np.ndarray] = {}
        row_numbers: dict[tuple[date, int], int] = {}
        for i, row in enumerate(rows):
            excel_row = i + 2
            quarter, hour_slot = divmod(i, 4)
            expected_date = date(2025, 1, 1) + timedelta(days=quarter)
            if hour_slot == 0:
                if not isinstance(row[0], str):
                    raise ValueError(f"附件3第{excel_row}行日期缺失或非文本")
                try:
                    parsed = datetime.strptime(row[0].strip(), "%Y-%m-%d").date()
                except ValueError as exc:
                    raise ValueError(f"附件3第{excel_row}行日期无法解析：{row[0]!r}") from exc
                if parsed != expected_date:
                    raise ValueError(
                        f"附件3第{excel_row}行日期{parsed}应为{expected_date}"
                    )
            else:
                if row[0] not in (None, ""):
                    raise ValueError(f"附件3第{excel_row}行不应有日期：{row[0]!r}")
            hour_label = row[1]
            if hour_label != f"{ISSUE_HOURS[hour_slot]}:00":
                raise ValueError(
                    f"附件3第{excel_row}行预报时刻{hour_label!r}应为"
                    f"{ISSUE_HOURS[hour_slot]}:00"
                )
            values = np.asarray(
                [
                    q2.checked_float(
                        value,
                        context=f"附件3第{excel_row}行第{column}列",
                        nonnegative=True,
                    )
                    for column, value in enumerate(row[2:26], start=3)
                ],
                dtype=np.float64,
            )
            forecasts[(expected_date, ISSUE_HOURS[hour_slot])] = values
            row_numbers[(expected_date, ISSUE_HOURS[hour_slot])] = excel_row
    finally:
        workbook.close()
    if len(forecasts) != 365 * 4:
        raise ValueError(f"附件3有效事件数{len(forecasts)}不足1460")
    return Attachment3(forecasts, row_numbers)


def window_positions(
    day_idx: int, issue_hour: int
) -> tuple[np.ndarray, np.ndarray]:
    """更新(d,h)窗口u=1..144的绝对位置：返回(day, t_abs)，均为1-based。"""
    kappa = KAPPA[issue_hour]
    t = np.arange(kappa + 1, kappa + T + 1, dtype=np.int64)
    overflow = t > T
    day = np.where(overflow, day_idx + 1, day_idx).astype(np.int64)
    t_abs = np.where(overflow, t - T, t).astype(np.int64)
    return day, t_abs


def window_prices(t_abs: np.ndarray, prices: np.ndarray) -> np.ndarray:
    return prices[t_abs - 1]


def formal_pv_window(
    day_idx: int,
    issue_hour: int,
    att3: Attachment3,
    annual: q2.AnnualData,
) -> np.ndarray:
    """附件3(d,h)行24个整点预测展开为窗口144段正式预报电量（kWh）。

    P_0 = 发布时刻已观测光伏功率（kW，取自附件2已完成区间；h=0取前一日t=144）；
    P_k = F_{d,h,k}；小时k内6段电量 = (P_{k-1}+P_k)/2/6。
    """
    forecast = att3.forecasts[(annual.dates[day_idx], issue_hour)]
    if issue_hour == 0:
        p0 = annual.pv_kwh[day_idx - 1, T - 1] * 6.0 if day_idx >= 1 else 0.0
    else:
        p0 = annual.pv_kwh[day_idx, KAPPA[issue_hour] - 1] * 6.0
    points = np.concatenate(([p0], forecast))
    hourly_mean = (points[:-1] + points[1:]) / 2.0
    return np.repeat(hourly_mean / 6.0, 6)  # (144,) kWh


def bates_granger_weight(
    error_formal: np.ndarray, error_historical: np.ndarray
) -> dict[str, Any]:
    """Bates-Granger组合权重：w* = Σ eH(eH−eF) / Σ(eF−eH)²，截断[0,1]，分母0→0.5。"""
    eF = np.asarray(error_formal, dtype=np.float64)
    eH = np.asarray(error_historical, dtype=np.float64)
    if eF.shape != eH.shape or eF.ndim != 1:
        raise ValueError("Bates-Granger误差向量维度异常")
    denominator = float(np.sum((eF - eH) ** 2))
    if denominator <= 0.0:
        return {
            "w": 0.5,
            "w_raw": 0.5,
            "denominator": denominator,
            "sample_count": int(eF.size),
            "clipped": False,
        }
    w_raw = float(np.sum(eH * (eH - eF)) / denominator)
    w = float(min(1.0, max(0.0, w_raw)))
    return {
        "w": w,
        "w_raw": w_raw,
        "denominator": denominator,
        "sample_count": int(eF.size),
        "clipped": w_raw != w,
    }


def combine_pv(
    w: float, formal: np.ndarray, historical: np.ndarray
) -> np.ndarray:
    """组合光伏预测：w·V̂^F+(1−w)·V̂^H，截断非负，夜间双零段保持0。"""
    combined = w * formal + (1.0 - w) * historical
    combined = np.maximum(combined, 0.0)
    both_zero = (formal == 0.0) & (historical == 0.0)
    combined[both_zero] = 0.0
    return combined


def error_triplet(error: np.ndarray) -> dict[str, float]:
    return {
        "mae_kwh": float(np.mean(np.abs(error))),
        "rmse_kwh": float(np.sqrt(np.mean(np.square(error)))),
        "bias_forecast_minus_actual_kwh": float(np.mean(error)),
    }


# ---------------------------------------------------------------------------
# 阶段A：预测、校准与场景缓存
# ---------------------------------------------------------------------------


def build_event(
    day_idx: int,
    hour: int,
    draft_d: q2.ForecastDraft,
    next_draft: q2.ForecastDraft | None,
    att3: Attachment3,
    annual: q2.AnnualData,
) -> PendingEvent | None:
    """为事件(d,h)构建窗口144段预测。历史不足返回None。

    负荷/历史光伏：当前日部分=当天日预测切片；跨日部分=次日日预测
    （用发布当时完整历史≤d-1生成，次日0:00预测当时尚不存在）。
    """
    kappa = KAPPA[hour]
    cur_load = draft_d.load_forecast
    cur_pv = draft_d.pv_forecast
    if cur_load is None or cur_pv is None:
        return None
    if kappa == 0:
        load_win = cur_load
        pv_hist = cur_pv
    else:
        if (
            next_draft is None
            or next_draft.load_forecast is None
            or next_draft.pv_forecast is None
        ):
            return None
        load_win = np.concatenate((cur_load[kappa:], next_draft.load_forecast[:kappa]))
        pv_hist = np.concatenate((cur_pv[kappa:], next_draft.pv_forecast[:kappa]))
    pv_formal = formal_pv_window(day_idx, hour, att3, annual)
    day, t_abs = window_positions(day_idx, hour)
    return PendingEvent(day_idx, hour, day, t_abs, load_win, pv_formal, pv_hist)


def realize_event(
    ev: PendingEvent, annual: q2.AnnualData, w: float
) -> tuple[date, int, np.ndarray, np.ndarray]:
    """事件窗口完整实现后计算配对残差eL、eV（组合光伏）。"""
    actual_load = annual.load_kwh[ev.day, ev.t_abs - 1]
    actual_pv = annual.pv_kwh[ev.day, ev.t_abs - 1]
    combined = combine_pv(w, ev.pv_formal, ev.pv_historical)
    e_load = actual_load - ev.load_forecast
    e_pv = actual_pv - combined
    return (annual.dates[ev.day_idx], ev.hour, e_load, e_pv)


def realize_events_ending_on(
    pending: list[PendingEvent], day_idx: int, annual: q2.AnnualData, w: float
) -> list[tuple[date, int, np.ndarray, np.ndarray]]:
    """实现窗口终点落在day_idx当天的全部pending事件，返回残差并移除。"""
    realized: list[tuple[date, int, np.ndarray, np.ndarray]] = []
    remaining: list[PendingEvent] = []
    for ev in pending:
        if ev.day_idx == day_idx - 1:
            realized.append(realize_event(ev, annual, w))
        else:
            remaining.append(ev)
    pending[:] = remaining
    return realized


def run_prep(args: argparse.Namespace) -> dict[str, Any]:
    out = output_dir(args)
    prices, price_labels = q2.read_prices(input_path(args.root, "附件1"))
    annual = q2.read_annual_data(input_path(args.root, "附件2"))
    q2.validate_time_inputs(price_labels, annual)
    att3 = read_attachment3(input_path(args.root, "附件3"))
    limit = args.limit_days
    started = time.perf_counter()

    class_rows, class_summary = q2.validate_january_load_classes(
        annual.dates[:31], annual.load_kwh[:31]
    )
    q2.write_csv(out / "q3_load_class_validation.csv", class_rows)

    history = q2.ForecastHistory.empty()
    realized: dict[int, list[tuple[date, int, np.ndarray, np.ndarray]]] = {
        h: [] for h in ISSUE_HOURS
    }

    # ---- Pass 1：2025年1月（仅用于Bates-Granger校准与残差池，不进入正式事件） ----
    calib_events: list[PendingEvent] = []
    for day_idx in range(31):
        d = annual.dates[day_idx]
        draft = q2.forecast_one_day(d, history)
        next_draft = q2.forecast_one_day(d + timedelta(days=1), history)
        for h in ISSUE_HOURS:
            ev = build_event(day_idx, h, draft, next_draft, att3, annual)
            if ev is not None:
                calib_events.append(ev)
        q2.append_observation(
            history, d, annual.load_kwh[day_idx], annual.pv_kwh[day_idx], draft
        )
    if not calib_events:
        raise RuntimeError("1月没有任何可校准事件（历史不足）")

    error_formal_parts: list[np.ndarray] = []
    error_historical_parts: list[np.ndarray] = []
    calibration_sample_rows: list[dict[str, Any]] = []
    for ev in calib_events:
        actual_pv = annual.pv_kwh[ev.day, ev.t_abs - 1]
        eF = actual_pv - ev.pv_formal
        eH = actual_pv - ev.pv_historical
        error_formal_parts.append(eF)
        error_historical_parts.append(eH)
        issue_date = annual.dates[ev.day_idx].isoformat()
        for u in range(T):
            calibration_sample_rows.append(
                {
                    "issue_date": issue_date,
                    "issue_hour": ev.hour,
                    "window_position": u + 1,
                    "absolute_date": annual.dates[int(ev.day[u])].isoformat(),
                    "absolute_t": int(ev.t_abs[u]),
                    "pv_formal_kwh": float(ev.pv_formal[u]),
                    "pv_historical_kwh": float(ev.pv_historical[u]),
                    "actual_pv_kwh": float(actual_pv[u]),
                    "error_formal_kwh": float(eF[u]),
                    "error_historical_kwh": float(eH[u]),
                }
            )
    bg = bates_granger_weight(
        np.concatenate(error_formal_parts), np.concatenate(error_historical_parts)
    )
    w = float(bg["w"])
    # 1月事件全部实现（窗口最晚终点为2月1日18:00，附件2均有实际值）
    for ev in calib_events:
        realized[ev.hour].append(realize_event(ev, annual, w))

    # 校准指标（F/H/组合在1月样本上的MAE/RMSE/偏差）
    eF_all = np.concatenate(error_formal_parts)
    eH_all = np.concatenate(error_historical_parts)
    combined_parts = [
        combine_pv(w, ev.pv_formal, ev.pv_historical) for ev in calib_events
    ]
    actual_parts = [annual.pv_kwh[ev.day, ev.t_abs - 1] for ev in calib_events]
    eC_all = np.concatenate(
        [actual - combined for actual, combined in zip(actual_parts, combined_parts)]
    )
    per_hour: dict[str, Any] = {}
    for h in ISSUE_HOURS:
        mask = np.asarray(
            [ev.hour == h for ev in calib_events for _ in range(T)], dtype=bool
        )
        per_hour[str(h)] = {
            "events": sum(1 for ev in calib_events if ev.hour == h),
            "formal": error_triplet(eF_all[mask]),
            "historical": error_triplet(eH_all[mask]),
            "combined": error_triplet(eC_all[mask]),
        }
    calibration = {
        "scope": "2025-01-only",
        "uses_february_to_december": False,
        "sample_events": len(calib_events),
        "sample_positions": int(eF_all.size),
        "w": w,
        "w_raw": float(bg["w_raw"]),
        "denominator": float(bg["denominator"]),
        "clipped": bool(bg["clipped"]),
        "formal": error_triplet(eF_all),
        "historical": error_triplet(eH_all),
        "combined": error_triplet(eC_all),
        "per_issue_hour": per_hour,
    }
    q2.write_csv(out / "q3_bg_calibration_samples.csv", calibration_sample_rows)
    (out / "q3_bg_calibration.json").write_text(
        json.dumps(calibration, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    # ---- Pass 2：2025-02-01..2025-12-31 正式事件 ----
    formal_count = FORMAL_DAYS if limit is None else min(FORMAL_DAYS, limit)
    total_events = formal_count * 4
    L_all = np.zeros((total_events, MAX_SCENARIOS, T), dtype=np.float64)
    V_all = np.zeros((total_events, MAX_SCENARIOS, T), dtype=np.float64)
    M_arr = np.zeros(total_events, dtype=np.int64)
    index_rows: list[dict[str, Any]] = []
    scenario_audit_rows: list[dict[str, Any]] = []
    pending: list[PendingEvent] = []
    labels = q2.physical_interval_labels()

    audit_path = out / "q3_forecast_audit.csv"
    audit_file = audit_path.open("w", encoding="utf-8-sig", newline="")
    audit_writer = csv.writer(audit_file)
    audit_writer.writerow(
        [
            "issue_date",
            "issue_hour",
            "attachment_row",
            "forecast_hour_column",
            "absolute_date",
            "absolute_t",
            "physical_interval",
            "load_forecast_kwh",
            "pv_historical_kwh",
            "pv_formal_kwh",
            "pv_combined_kwh",
            "actual_load_kwh",
            "actual_pv_kwh",
        ]
    )
    try:
        for di in range(formal_count):
            day_idx = 31 + di
            d = annual.dates[day_idx]
            for item in realize_events_ending_on(pending, day_idx, annual, w):
                realized[item[1]].append(item)
            draft = q2.forecast_one_day(d, history)
            next_draft = q2.forecast_one_day(d + timedelta(days=1), history)
            for hi, h in enumerate(ISSUE_HOURS):
                ev = build_event(day_idx, h, draft, next_draft, att3, annual)
                if ev is None:
                    raise RuntimeError(f"{d} {h}:00历史不足，无法构建窗口预测")
                ev.pv_combined = combine_pv(w, ev.pv_formal, ev.pv_historical)
                pool = realized[h][-MAX_SCENARIOS:]
                if not pool:
                    raise RuntimeError(f"{d} {h}:00没有可用残差场景")
                m = len(pool)
                scenario_load = np.stack(
                    [np.maximum(ev.load_forecast + item[2], 0.0) for item in pool]
                )
                scenario_pv = np.stack(
                    [np.maximum(ev.pv_combined + item[3], 0.0) for item in pool]
                )
                idx = di * 4 + hi
                M_arr[idx] = m
                L_all[idx, :m] = scenario_load
                V_all[idx, :m] = scenario_pv
                index_rows.append(
                    {
                        "date": d.isoformat(),
                        "issue_hour": h,
                        "event_index": idx,
                        "scenario_count": m,
                        "pool_min_date": pool[0][0].isoformat(),
                        "pool_max_date": pool[-1][0].isoformat(),
                    }
                )
                weight = 1.0 / m
                for rank, item in enumerate(pool, start=1):
                    scenario_audit_rows.append(
                        {
                            "issue_date": d.isoformat(),
                            "issue_hour": h,
                            "rank": rank,
                            "residual_date": item[0].isoformat(),
                            "residual_hour": item[1],
                            "weight": weight,
                        }
                    )
                attachment_row = att3.row_numbers[(d, h)]
                for u in range(T):
                    actual_day = int(ev.day[u])
                    audit_writer.writerow(
                        [
                            d.isoformat(),
                            h,
                            attachment_row,
                            u // 6 + 1,
                            annual.dates[actual_day].isoformat() if actual_day < 365 else "",
                            int(ev.t_abs[u]),
                            labels[int(ev.t_abs[u]) - 1],
                            float(ev.load_forecast[u]),
                            float(ev.pv_historical[u]),
                            float(ev.pv_formal[u]),
                            float(ev.pv_combined[u]),
                            float(annual.load_kwh[actual_day, int(ev.t_abs[u]) - 1])
                            if actual_day < 365
                            else "",
                            float(annual.pv_kwh[actual_day, int(ev.t_abs[u]) - 1])
                            if actual_day < 365
                            else "",
                        ]
                    )
                pending.append(ev)
            q2.append_observation(
                history, d, annual.load_kwh[day_idx], annual.pv_kwh[day_idx], draft
            )
            if (di + 1) % 30 == 0 or di + 1 == formal_count:
                print(
                    f"[prep] {d} done, {di + 1}/{formal_count} 天, "
                    f"elapsed={time.perf_counter() - started:.0f}s",
                    flush=True,
                )
    finally:
        audit_file.close()

    # ---- 边界日2026-01-01（0:00）：附件3无2026行 → 光伏仅历史预测 ----
    boundary_draft = q2.forecast_one_day(BOUNDARY_PLAN_DATE, history)
    if boundary_draft.load_forecast is None or boundary_draft.pv_forecast is None:
        raise RuntimeError("边界日历史不足")
    boundary_pool = realized[0][-MAX_SCENARIOS:]
    if not boundary_pool:
        raise RuntimeError("边界日没有可用残差场景")
    boundary_load = np.stack(
        [
            np.maximum(boundary_draft.load_forecast + item[2], 0.0)
            for item in boundary_pool
        ]
    )
    boundary_pv = np.stack(
        [
            np.maximum(boundary_draft.pv_forecast + item[3], 0.0)
            for item in boundary_pool
        ]
    )
    np.savez(
        out / "q3_boundary.npz",
        L=boundary_load,
        V=boundary_pv,
        M=np.asarray([boundary_load.shape[0]], dtype=np.int64),
    )
    boundary_meta = {
        "target_date": BOUNDARY_PLAN_DATE.isoformat(),
        "issue_hour": 0,
        "formal_forecast_available": False,
        "note": "附件3无2026年数据，边界日光伏仅用历史预测（Q2口径），场景沿用(i,0)事件残差。",
        "w": w,
        "scenario_count": int(boundary_load.shape[0]),
        "residual_dates": [item[0].isoformat() for item in boundary_pool],
        "load_forecast_kwh": boundary_draft.load_forecast.tolist(),
        "pv_historical_kwh": boundary_draft.pv_forecast.tolist(),
    }
    (out / "q3_boundary_meta.json").write_text(
        json.dumps(boundary_meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    np.savez(out / "q3_scenarios.npz", L=L_all, V=V_all, M=M_arr)
    q2.write_csv(out / "q3_event_index.csv", index_rows)
    q2.write_csv(out / "q3_scenario_audit.csv", scenario_audit_rows)

    summary = {
        "bates_granger": calibration,
        "load_class_validation": class_summary,
        "formal_events": total_events,
        "boundary_scenario_count": int(boundary_load.shape[0]),
        "elapsed_seconds": time.perf_counter() - started,
    }
    print(
        json.dumps(
            {
                "prep_done": True,
                "w": w,
                "formal_events": total_events,
                "calibration_sample_events": len(calib_events),
                "elapsed_seconds": round(summary["elapsed_seconds"], 1),
            },
            ensure_ascii=False,
        ),
        flush=True,
    )
    return summary


# ---------------------------------------------------------------------------
# 窗口LP（0:00原始计划 / h>0滚动调整）
# ---------------------------------------------------------------------------


def build_window_lp(
    issue_label: str,
    prices_window: np.ndarray,
    scenario_net: np.ndarray,
    initial_soc: float,
    terminal_value_coefficient: float,
    ncur: int,
    g0: np.ndarray | None,
    *,
    binary_mutex: bool,
) -> WindowModel:
    """g0为None→0:00原始计划模型（z=g，无δ）；否则→滚动模型（z=[a;g̃]，a−g0=δ⁺−δ⁻）。"""
    scenarios = scenario_net.shape[0]
    problem = pulp.LpProblem(f"Q3_{issue_label}", pulp.LpMinimize)
    z = [pulp.LpVariable(f"z_{u + 1:03d}", lowBound=0.0) for u in range(T)]
    charge = [
        [
            pulp.LpVariable(
                f"C_{sc + 1:02d}_{u + 1:03d}", lowBound=0.0, upBound=MAX_ACTION_KWH
            )
            for u in range(T)
        ]
        for sc in range(scenarios)
    ]
    discharge = [
        [
            pulp.LpVariable(
                f"D_{sc + 1:02d}_{u + 1:03d}", lowBound=0.0, upBound=MAX_ACTION_KWH
            )
            for u in range(T)
        ]
        for sc in range(scenarios)
    ]
    emergency = [
        [pulp.LpVariable(f"b_{sc + 1:02d}_{u + 1:03d}", lowBound=0.0) for u in range(T)]
        for sc in range(scenarios)
    ]
    unused = [
        [pulp.LpVariable(f"U_{sc + 1:02d}_{u + 1:03d}", lowBound=0.0) for u in range(T)]
        for sc in range(scenarios)
    ]
    soc = [
        [
            pulp.LpVariable(
                f"E_{sc + 1:02d}_{u + 1:03d}",
                lowBound=SOC_MIN_KWH,
                upBound=SOC_MAX_KWH,
            )
            for u in range(T)
        ]
        for sc in range(scenarios)
    ]
    delta_plus: list[pulp.LpVariable] = []
    delta_minus: list[pulp.LpVariable] = []
    mode: list[list[pulp.LpVariable]] | None = None
    if g0 is not None:
        delta_plus = [
            pulp.LpVariable(f"dp_{u + 1:03d}", lowBound=0.0) for u in range(ncur)
        ]
        delta_minus = [
            pulp.LpVariable(f"dm_{u + 1:03d}", lowBound=0.0) for u in range(ncur)
        ]
    if binary_mutex:
        mode = [
            [
                pulp.LpVariable(f"m_{sc + 1:02d}_{u + 1:03d}", cat=pulp.LpBinary)
                for u in range(T)
            ]
            for sc in range(scenarios)
        ]

    for sc in range(scenarios):
        for u in range(T):
            problem += (
                z[u] + emergency[sc][u] + discharge[sc][u]
                == float(scenario_net[sc, u]) + charge[sc][u] + unused[sc][u],
                f"balance_{sc + 1:02d}_{u + 1:03d}",
            )
            previous_soc: float | pulp.LpVariable = (
                initial_soc if u == 0 else soc[sc][u - 1]
            )
            problem += (
                soc[sc][u]
                == previous_soc
                + ETA_CHARGE * charge[sc][u]
                - discharge[sc][u] / ETA_DISCHARGE,
                f"soc_{sc + 1:02d}_{u + 1:03d}",
            )
            if mode is not None:
                problem += (
                    charge[sc][u] <= MAX_ACTION_KWH * mode[sc][u],
                    f"charge_mutex_{sc + 1:02d}_{u + 1:03d}",
                )
                problem += (
                    discharge[sc][u] <= MAX_ACTION_KWH * (1.0 - mode[sc][u]),
                    f"discharge_mutex_{sc + 1:02d}_{u + 1:03d}",
                )

    if g0 is not None:
        kappa = T - ncur
        for u in range(ncur):
            problem += (
                z[u] - float(g0[kappa + u])
                == delta_plus[u] - delta_minus[u],
                f"delta_{u + 1:03d}",
            )

    probability = 1.0 / scenarios
    if g0 is None:
        plan_cost = pulp.lpSum(float(prices_window[u]) * z[u] for u in range(T))
    else:
        plan_cost = pulp.lpSum(
            float(prices_window[u])
            * (z[u] + 0.5 * (delta_plus[u] + delta_minus[u]))
            for u in range(ncur)
        ) + pulp.lpSum(float(prices_window[u]) * z[u] for u in range(ncur, T))
    expected_recourse = probability * pulp.lpSum(
        pulp.lpSum(
            EMERGENCY_PRICE_FACTOR * float(prices_window[u]) * emergency[sc][u]
            for u in range(T)
        )
        - terminal_value_coefficient * soc[sc][-1]
        for sc in range(scenarios)
    )
    original_objective = plan_cost + expected_recourse
    problem += original_objective, "expected_total_cost_with_terminal_value"
    return WindowModel(
        problem,
        original_objective,
        z,
        charge,
        discharge,
        emergency,
        unused,
        soc,
        delta_plus,
        delta_minus,
        ncur,
        mode,
    )


def solve_current_model(model: WindowModel, *, mip: bool) -> tuple[str, float, Any]:
    started = time.perf_counter()
    status_code = model.problem.solve(q2.make_highs_solver(mip=mip))
    elapsed = time.perf_counter() - started
    status = pulp.LpStatus.get(status_code, f"Unknown({status_code})")
    if status != "Optimal":
        raise RuntimeError(f"窗口LP未达到Optimal：{status}")
    return status, elapsed, model.problem.solverModel.getInfo()


def values_1d(variables: Sequence[pulp.LpVariable]) -> np.ndarray:
    values = [variable.value() for variable in variables]
    if any(value is None or not math.isfinite(float(value)) for value in values):
        raise RuntimeError("求解器返回了空值或非有限变量")
    return np.asarray(values, dtype=np.float64)


def values_2d(variables: Sequence[Sequence[pulp.LpVariable]]) -> np.ndarray:
    return np.asarray([values_1d(row) for row in variables], dtype=np.float64)


def window_arrays(model: WindowModel) -> dict[str, np.ndarray]:
    delta_plus = np.zeros(T, dtype=np.float64)
    delta_minus = np.zeros(T, dtype=np.float64)
    if model.delta_plus:
        delta_plus[: model.ncur] = values_1d(model.delta_plus)
        delta_minus[: model.ncur] = values_1d(model.delta_minus)
    return {
        "z": values_1d(model.z),
        "c": values_2d(model.charge),
        "d": values_2d(model.discharge),
        "b": values_2d(model.emergency),
        "u": values_2d(model.unused),
        "e": values_2d(model.soc),
        "dp": delta_plus,
        "dm": delta_minus,
    }


def solve_window_lp(
    issue_date: date,
    issue_hour: int,
    kind: str,
    prices_window: np.ndarray,
    scenario_net: np.ndarray,
    initial_soc: float,
    g0: np.ndarray | None,
    ncur: int,
    *,
    terminal_value_coefficient: float,
) -> UpdateResult:
    if scenario_net.ndim != 2 or scenario_net.shape[1] != T or scenario_net.shape[0] < 1:
        raise ValueError(f"窗口场景净负荷维度异常：{scenario_net.shape}")
    if not (SOC_MIN_KWH - CHECK_TOLERANCE <= initial_soc <= SOC_MAX_KWH + CHECK_TOLERANCE):
        raise ValueError(f"窗口LP初始SOC越界：{initial_soc}")
    if (g0 is None) != (ncur == T):
        raise ValueError("原始计划模型的ncur必须为144，滚动模型必须提供g0")
    if g0 is not None and (g0.shape != (T,) or ncur <= 0 or ncur >= T):
        raise ValueError("滚动模型g0/当前日部分长度异常")

    issue_label = f"{issue_date.isoformat()}_{issue_hour:02d}"
    model = build_window_lp(
        issue_label,
        prices_window,
        scenario_net,
        initial_soc,
        terminal_value_coefficient,
        ncur,
        g0,
        binary_mutex=False,
    )
    status, solve_seconds, info = solve_current_model(model, mip=False)
    best_objective = float(pulp.value(model.original_objective))
    arrays = window_arrays(model)
    max_product = float(np.max(arrays["c"] * arrays["d"]))
    tiebreak_used = False
    binary_fallback = False

    if max_product > SIMULTANEOUS_PRODUCT_TOLERANCE:
        tiebreak_used = True
        cost_tolerance = max(1.0e-7, abs(best_objective) * 1.0e-10)
        model.problem += (
            model.original_objective <= best_objective + cost_tolerance,
            "preserve_primary_optimal_cost",
        )
        model.problem.setObjective(
            pulp.lpSum(
                variable
                for rows in (model.charge, model.discharge)
                for row in rows
                for variable in row
            )
        )
        tie_status, tie_seconds, info = solve_current_model(model, mip=False)
        status = f"{status}; tiebreak={tie_status}"
        solve_seconds += tie_seconds
        arrays = window_arrays(model)
        max_product = float(np.max(arrays["c"] * arrays["d"]))

    if max_product > SIMULTANEOUS_PRODUCT_TOLERANCE:
        binary_fallback = True
        model = build_window_lp(
            issue_label,
            prices_window,
            scenario_net,
            initial_soc,
            terminal_value_coefficient,
            ncur,
            g0,
            binary_mutex=True,
        )
        status, binary_seconds, info = solve_current_model(model, mip=True)
        solve_seconds += binary_seconds
        status = f"binary_fallback={status}"
        arrays = window_arrays(model)
        max_product = float(np.max(arrays["c"] * arrays["d"]))

    objective_value = float(pulp.value(model.original_objective))
    z = arrays["z"]
    c = arrays["c"]
    d = arrays["d"]
    b = arrays["b"]
    u = arrays["u"]
    e = arrays["e"]
    previous_e = np.concatenate(
        [np.full((e.shape[0], 1), initial_soc), e[:, :-1]], axis=1
    )
    balance_residual = z[None, :] + b + d - scenario_net - c - u
    soc_residual = e - previous_e - ETA_CHARGE * c + d / ETA_DISCHARGE
    max_balance = float(np.max(np.abs(balance_residual)))
    max_soc_residual = float(np.max(np.abs(soc_residual)))
    max_delta_product = float(np.max(arrays["dp"] * arrays["dm"]))
    min_flow = float(min(np.min(z), np.min(c), np.min(d), np.min(b), np.min(u)))
    min_soc = float(np.min(e))
    max_soc = float(np.max(e))
    max_charge = float(np.max(c))
    max_discharge = float(np.max(d))

    checks = {
        "balance": max_balance <= CHECK_TOLERANCE,
        "soc_transition": max_soc_residual <= CHECK_TOLERANCE,
        "soc_bounds": min_soc >= SOC_MIN_KWH - CHECK_TOLERANCE
        and max_soc <= SOC_MAX_KWH + CHECK_TOLERANCE,
        "action_bounds": max_charge <= MAX_ACTION_KWH + CHECK_TOLERANCE
        and max_discharge <= MAX_ACTION_KWH + CHECK_TOLERANCE,
        "nonnegative": min_flow >= -CHECK_TOLERANCE,
        "no_simultaneous": max_product <= SIMULTANEOUS_PRODUCT_TOLERANCE,
        "delta_mutex": max_delta_product <= SIMULTANEOUS_PRODUCT_TOLERANCE,
    }
    if not all(checks.values()):
        raise RuntimeError(f"{issue_label}窗口LP验收失败：{checks}")

    cleaned_z = z.copy()
    cleaned_z[np.abs(cleaned_z) <= ACTION_ZERO_TOLERANCE] = 0.0
    highs_mip_gap = getattr(info, "mip_gap", None) if binary_fallback else None
    return UpdateResult(
        issue_date=issue_date,
        issue_hour=issue_hour,
        kind=kind,
        z=cleaned_z,
        delta_plus=arrays["dp"],
        delta_minus=arrays["dm"],
        objective_value=objective_value,
        scenario_count=scenario_net.shape[0],
        status=status,
        solve_seconds=solve_seconds,
        tiebreak_used=tiebreak_used,
        binary_fallback_used=binary_fallback,
        max_balance_residual_kwh=max_balance,
        max_soc_residual_kwh=max_soc_residual,
        max_delta_product_kwh2=max_delta_product,
        max_simultaneous_product_kwh2=max_product,
        highs_max_primal_infeasibility=float(
            getattr(info, "max_primal_infeasibility", float("nan"))
        ),
        simplex_iterations=int(getattr(info, "simplex_iteration_count", -1)),
        ipm_iterations=int(getattr(info, "ipm_iteration_count", -1)),
        mip_gap=None if highs_mip_gap is None else float(highs_mip_gap),
    )


# ---------------------------------------------------------------------------
# DP执行器与方案仿真
# ---------------------------------------------------------------------------


def execute_span(
    prices: np.ndarray,
    actual_net: np.ndarray,
    plan: np.ndarray,
    k_from: int,
    k_to: int,
    future_values: np.ndarray,
    current_soc: float,
    grid: np.ndarray,
    charge: np.ndarray,
    discharge: np.ndarray,
    emergency: np.ndarray,
    unused: np.ndarray,
    soc_arr: np.ndarray,
) -> float:
    """用更新时刻k_from生成的价值函数执行时段(k_from, k_to]（0-based半开）。"""
    for t_idx in range(k_from, k_to):
        u_idx = t_idx - k_from
        future = future_values[u_idx + 1]
        r = float(actual_net[t_idx] - plan[t_idx])
        if r > 0.0:
            lower = max(
                SOC_MIN_KWH,
                current_soc - min(MAX_ACTION_KWH, r) / ETA_DISCHARGE,
            )
            upper = current_soc
            slope = EMERGENCY_PRICE_FACTOR * float(prices[t_idx]) * ETA_DISCHARGE
            transformed = future + slope * grid
        else:
            lower = current_soc
            upper = min(
                SOC_MAX_KWH,
                current_soc + ETA_CHARGE * min(MAX_ACTION_KWH, -r),
            )
            transformed = future
        candidates = q2.candidate_minimizers(
            transformed, grid, lower, upper, current_soc
        )
        evaluated: list[tuple[float, float, float]] = []
        for next_soc in candidates:
            x = next_soc - current_soc
            psi = x / ETA_CHARGE if x >= 0.0 else ETA_DISCHARGE * x
            stage_cost = (
                EMERGENCY_PRICE_FACTOR
                * float(prices[t_idx])
                * max(r + psi, 0.0)
            )
            objective = stage_cost + float(np.interp(next_soc, grid, future))
            evaluated.append((objective, abs(x), next_soc))
        best_objective = min(item[0] for item in evaluated)
        objective_tolerance = max(1.0e-9, abs(best_objective) * 1.0e-12)
        best_candidates = [
            item for item in evaluated if item[0] <= best_objective + objective_tolerance
        ]
        _, _, next_soc = min(best_candidates, key=lambda item: (item[1], item[2]))
        x = next_soc - current_soc
        if abs(x) <= ACTION_ZERO_TOLERANCE:
            x = 0.0
            next_soc = current_soc
        if x >= 0.0:
            charge[t_idx] = x / ETA_CHARGE
        else:
            discharge[t_idx] = -ETA_DISCHARGE * x
        psi = x / ETA_CHARGE if x >= 0.0 else ETA_DISCHARGE * x
        emergency[t_idx] = max(r + psi, 0.0)
        unused[t_idx] = max(-r - psi, 0.0)
        current_soc = min(max(next_soc, SOC_MIN_KWH), SOC_MAX_KWH)
        soc_arr[t_idx] = current_soc
    return current_soc


def compose_effective_plan(
    updates: Sequence[tuple[int, np.ndarray]]
) -> np.ndarray:
    """按§8.5由各更新版本合成最终生效计划g^eff。

    updates: [(hour, z)]，hour升序且必须含0:00；每个z的当前日部分
    z[0:ncur]覆盖t=κ_h+1..144，生效区间为(κ_h, κ_next]。
    """
    if not updates or updates[0][0] != 0:
        raise ValueError("生效计划合成必须包含0:00更新")
    geff = np.zeros(T, dtype=np.float64)
    ordered = sorted(updates)
    for i, (h, z) in enumerate(ordered):
        k = KAPPA[h]
        k_next = KAPPA[ordered[i + 1][0]] if i + 1 < len(ordered) else T
        if z.shape != (T,):
            raise ValueError(f"{h}:00窗口计划维度异常：{z.shape}")
        geff[k:k_next] = z[: k_next - k]
    return geff


def run_scheme(args: argparse.Namespace) -> None:
    out = output_dir(args)
    scheme = args.scheme
    if scheme is None:
        raise ValueError("--scheme 必须给出")
    prices, _ = q2.read_prices(input_path(args.root, "附件1"))
    terminal_value_coefficient = q2.terminal_value_coefficient(prices)
    annual = q2.read_annual_data(input_path(args.root, "附件2"))
    index = read_csv_rows(out / "q3_event_index.csv")
    if len(index) % 4 != 0 or len(index) < 4:
        raise RuntimeError("q3_event_index.csv缺失或异常，请先运行 --prep")
    formal_count = len(index) // 4
    if args.limit_days is not None:
        formal_count = min(formal_count, args.limit_days)
    data = np.load(out / "q3_scenarios.npz")
    L_all = data["L"]
    V_all = data["V"]
    M_arr = data["M"]
    if L_all.shape[0] < formal_count * 4:
        raise RuntimeError(
            "q3_scenarios.npz事件数不足，prep与scheme的--limit-days口径需一致"
        )

    update_hours = SCHEME_UPDATE_HOURS[scheme]
    grid = q2.state_grid(DP_GRID_STEP_KWH)
    dp_soc = INITIAL_SOC_KWH
    ana_soc = INITIAL_SOC_KWH
    daily_rows: list[dict[str, Any]] = []
    solver_rows: list[dict[str, Any]] = []
    executor_rows: list[dict[str, Any]] = []
    interval_rows: list[dict[str, Any]] = []
    started = time.perf_counter()
    progress = out / f"q3_{scheme}_progress.log"
    versions_path = out / "q3_all_plan_versions.csv"
    versions_file: Any = None
    versions_writer: Any = None
    if scheme == "all":
        versions_file = versions_path.open("w", encoding="utf-8-sig", newline="")
        versions_writer = csv.writer(versions_file)
        versions_writer.writerow(
            [
                "date",
                "issue_hour",
                "window_position",
                "absolute_date_offset",
                "absolute_t",
                "variable_type",
                "plan_kwh",
                "delta_plus_kwh",
                "delta_minus_kwh",
                "g0_reference_kwh",
            ]
        )
    try:
        with progress.open("w", encoding="utf-8") as progress_log:
            for di in range(formal_count):
                day_idx = 31 + di
                d = annual.dates[day_idx]
                actual_net = annual.load_kwh[day_idx] - annual.pv_kwh[day_idx]
                ordered = sorted(update_hours)
                kappas = [KAPPA[h] for h in ordered] + [T]

                dp_charge = np.zeros(T, dtype=np.float64)
                dp_discharge = np.zeros(T, dtype=np.float64)
                dp_emergency = np.zeros(T, dtype=np.float64)
                dp_unused = np.zeros(T, dtype=np.float64)
                dp_soc_arr = np.zeros(T, dtype=np.float64)
                current_soc = dp_soc
                previous_kappa = 0
                future_values = None
                day_conv = 0.0
                solved: list[tuple[int, UpdateResult]] = []

                for ui, h in enumerate(ordered):
                    k = KAPPA[h]
                    k_next = kappas[ui + 1]
                    if previous_kappa < k:
                        current_soc = execute_span(
                            prices,
                            actual_net,
                            geff,
                            previous_kappa,
                            k,
                            future_values,
                            current_soc,
                            grid,
                            dp_charge,
                            dp_discharge,
                            dp_emergency,
                            dp_unused,
                            dp_soc_arr,
                        )
                    ev_idx = di * 4 + HOUR_INDEX[h]
                    scenario_count = int(M_arr[ev_idx])
                    if scenario_count <= 0:
                        raise RuntimeError(f"{d} {h}:00场景缺失")
                    scenario_net = (
                        L_all[ev_idx, :scenario_count] - V_all[ev_idx, :scenario_count]
                    )
                    _, t_abs = window_positions(day_idx, h)
                    pw = window_prices(t_abs, prices)
                    if h == 0:
                        result = solve_window_lp(
                            d, 0, "original", pw, scenario_net, current_soc, None, T,
                            terminal_value_coefficient=terminal_value_coefficient,
                        )
                        g0 = result.z
                        geff = g0.copy()
                    else:
                        result = solve_window_lp(
                            d, h, "rolling", pw, scenario_net, current_soc, g0, T - k,
                            terminal_value_coefficient=terminal_value_coefficient,
                        )
                        geff[k:k_next] = result.z[: k_next - k]
                    future_values, conv = q2.average_scenario_value_functions(
                        scenario_net,
                        result.z,
                        pw,
                        terminal_value_coefficient,
                        grid,
                    )
                    day_conv = max(day_conv, conv)
                    previous_kappa = k
                    solved.append((h, result))
                    solver_rows.append(
                        {
                            "date": d.isoformat(),
                            "issue_hour": h,
                            "kind": result.kind,
                            "initial_soc_kwh": float(current_soc),
                            "objective_value": result.objective_value,
                            "scenario_count": result.scenario_count,
                            "status": result.status,
                            "solve_seconds": result.solve_seconds,
                            "tiebreak_used": result.tiebreak_used,
                            "binary_fallback_used": result.binary_fallback_used,
                            "max_balance_residual_kwh": result.max_balance_residual_kwh,
                            "max_soc_residual_kwh": result.max_soc_residual_kwh,
                            "max_delta_product_kwh2": result.max_delta_product_kwh2,
                            "max_simultaneous_product_kwh2": result.max_simultaneous_product_kwh2,
                            "highs_max_primal_infeasibility": result.highs_max_primal_infeasibility,
                            "simplex_iterations": result.simplex_iterations,
                            "ipm_iterations": result.ipm_iterations,
                            "mip_gap": result.mip_gap,
                        }
                    )
                    if scheme == "all":
                        ncur = T if h == 0 else T - k
                        for u in range(T):
                            versions_writer.writerow(
                                [
                                    d.isoformat(),
                                    h,
                                    u + 1,
                                    1 if u + 1 + k > T else 0,
                                    int(t_abs[u]),
                                    "g0" if h == 0 else (f"a{h}" if u < ncur else "gtilde"),
                                    float(result.z[u]),
                                    float(result.delta_plus[u]) if u < ncur and h > 0 else "",
                                    float(result.delta_minus[u]) if u < ncur and h > 0 else "",
                                    float(g0[int(t_abs[u]) - 1]) if h > 0 and u < ncur else "",
                                ]
                            )

                if previous_kappa < T:
                    current_soc = execute_span(
                        prices,
                        actual_net,
                        geff,
                        previous_kappa,
                        T,
                        future_values,
                        current_soc,
                        grid,
                        dp_charge,
                        dp_discharge,
                        dp_emergency,
                        dp_unused,
                        dp_soc_arr,
                    )
                dp_result = q2.validate_executor(
                    prices,
                    actual_net,
                    geff,
                    dp_soc,
                    dp_charge,
                    dp_discharge,
                    dp_emergency,
                    dp_unused,
                    dp_soc_arr,
                    dp_max_convexity_violation=day_conv,
                )
                ana_result = q2.execute_analytical(prices, actual_net, geff, ana_soc)

                plan_cost = float(np.dot(prices, g0))
                delta = geff - g0
                adj_cost = float(
                    np.dot(
                        prices,
                        ADJUST_UP_PRICE_FACTOR * np.maximum(delta, 0.0)
                        - ADJUST_DOWN_PRICE_FACTOR * np.maximum(-delta, 0.0),
                    )
                )
                emg_cost = float(np.dot(EMERGENCY_PRICE_FACTOR * prices, dp_emergency))
                total_cost = plan_cost + adj_cost + emg_cost
                total_cost_alt = float(
                    np.dot(
                        prices,
                        geff + 0.5 * np.abs(delta) + EMERGENCY_PRICE_FACTOR * dp_emergency,
                    )
                )
                ana_total = plan_cost + adj_cost + ana_result.emergency_cost_yuan
                daily_rows.append(
                    {
                        "date": d.isoformat(),
                        "dp_initial_soc_kwh": dp_soc,
                        "dp_terminal_soc_kwh": dp_result.terminal_soc_kwh,
                        "ana_initial_soc_kwh": ana_soc,
                        "ana_terminal_soc_kwh": ana_result.terminal_soc_kwh,
                        "plan_cost_yuan": plan_cost,
                        "adj_cost_yuan": adj_cost,
                        "adj_up_energy_kwh": float(np.sum(np.maximum(delta, 0.0))),
                        "adj_down_energy_kwh": float(np.sum(np.maximum(-delta, 0.0))),
                        "dp_emergency_energy_kwh": dp_result.emergency_energy_kwh,
                        "dp_emergency_cost_yuan": emg_cost,
                        "dp_total_cost_yuan": total_cost,
                        "total_cost_yuan_alt": total_cost_alt,
                        "ana_emergency_energy_kwh": ana_result.emergency_energy_kwh,
                        "ana_emergency_cost_yuan": ana_result.emergency_cost_yuan,
                        "ana_total_cost_yuan": ana_total,
                        "g0_energy_kwh": float(np.sum(g0)),
                        "geff_energy_kwh": float(np.sum(geff)),
                        "updates_solved": len(ordered),
                        "dp_max_balance_residual_kwh": dp_result.max_balance_residual_kwh,
                        "dp_max_soc_residual_kwh": dp_result.max_soc_residual_kwh,
                        "ana_max_balance_residual_kwh": ana_result.max_balance_residual_kwh,
                        "ana_max_soc_residual_kwh": ana_result.max_soc_residual_kwh,
                        "dp_max_convexity_violation": day_conv,
                    }
                )
                executor_rows.append(
                    {
                        "date": d.isoformat(),
                        "dp_emergency_cost_yuan": emg_cost,
                        "ana_emergency_cost_yuan": ana_result.emergency_cost_yuan,
                        "dp_emergency_energy_kwh": dp_result.emergency_energy_kwh,
                        "ana_emergency_energy_kwh": ana_result.emergency_energy_kwh,
                        "dp_max_balance_residual_kwh": dp_result.max_balance_residual_kwh,
                        "ana_max_balance_residual_kwh": ana_result.max_balance_residual_kwh,
                        "dp_max_soc_residual_kwh": dp_result.max_soc_residual_kwh,
                        "ana_max_soc_residual_kwh": ana_result.max_soc_residual_kwh,
                        "dp_total_cost_yuan": total_cost,
                        "ana_total_cost_yuan": ana_total,
                    }
                )
                if scheme == "all":
                    for t in range(T):
                        interval_rows.append(
                            {
                                "date": d.isoformat(),
                                "t": t + 1,
                                "g0_kwh": float(g0[t]),
                                "geff_kwh": float(geff[t]),
                                "dp_charge_kwh": float(dp_charge[t]),
                                "dp_discharge_kwh": float(dp_discharge[t]),
                                "dp_emergency_kwh": float(dp_emergency[t]),
                                "dp_unused_kwh": float(dp_unused[t]),
                                "dp_soc_kwh": float(dp_soc_arr[t]),
                                "ana_charge_kwh": float(ana_result.charge[t]),
                                "ana_discharge_kwh": float(ana_result.discharge[t]),
                                "ana_emergency_kwh": float(ana_result.emergency[t]),
                                "ana_unused_kwh": float(ana_result.unused[t]),
                                "ana_soc_kwh": float(ana_result.soc[t]),
                            }
                        )

                dp_soc = dp_result.terminal_soc_kwh
                ana_soc = ana_result.terminal_soc_kwh
                progress_log.write(
                    f"{d} updates={len(ordered)} dp_soc={dp_soc:.3f} "
                    f"total={total_cost:.2f}\n"
                )
                progress_log.flush()
                print(
                    f"[{scheme}] {d} total={total_cost:.2f} "
                    f"elapsed={time.perf_counter() - started:.0f}s",
                    flush=True,
                )
    finally:
        if versions_file is not None:
            versions_file.close()

    q2.write_csv(
        out / f"q3_{scheme}_daily_metrics.csv",
        daily_rows,
        fieldnames=list(daily_rows[0]) if daily_rows else None,
    )
    q2.write_csv(
        out / f"q3_{scheme}_solver_audit.csv",
        solver_rows,
        fieldnames=list(solver_rows[0]) if solver_rows else None,
    )
    q2.write_csv(
        out / f"q3_{scheme}_executor_comparison.csv",
        executor_rows,
        fieldnames=list(executor_rows[0]) if executor_rows else None,
    )
    if scheme == "all":
        q2.write_csv(
            out / "q3_all_interval_details.csv",
            interval_rows,
            fieldnames=list(interval_rows[0]) if interval_rows else None,
        )
    print(
        f"[{scheme}] 完成：{len(daily_rows)}天，总耗时{time.perf_counter() - started:.0f}s",
        flush=True,
    )


# ---------------------------------------------------------------------------
# 阶段C：result3回填与汇总
# ---------------------------------------------------------------------------


def validate_result3_template(path: Path) -> tuple[date, ...]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        if tuple(workbook.sheetnames) != RESULT3_SHEETS:
            raise ValueError(f"result3工作表结构异常：{workbook.sheetnames}")
        expected = tuple(
            FORECAST_START_DATE + timedelta(days=i) for i in range(FORMAL_DAYS)
        )
        for name in ("计划购电量", "调整购电量"):
            sheet = workbook[name]
            if sheet.max_row != 335 or sheet.max_column != 147:
                raise ValueError(f"result3{name}尺寸异常：{sheet.max_row}×{sheet.max_column}")
            headers = [sheet.cell(1, column).value for column in range(1, 148)]
            if headers[0] != "日期\\时间" or headers[1] != "0:10-0:20":
                raise ValueError(f"result3{name}前部表头异常")
            if headers[143:147] != [
                "23:50-0:00+1",
                "0:00-0:10+1",
                "全天购电量",
                "全天购电费",
            ]:
                raise ValueError(f"result3{name}尾部表头异常：{headers[143:147]}")
            dates = tuple(
                q2.as_date(sheet.cell(row, 1).value, f"result3{name}第{row}行")
                for row in range(2, 336)
            )
            if dates != expected:
                raise ValueError(f"result3{name}日期不连续或不完整")
        storage = workbook["充放电量"]
        if storage.max_row != 26 or storage.max_column != 6:
            raise ValueError(f"result3充放电量表尺寸异常：{storage.max_row}×{storage.max_column}")
        storage_labels = [
            storage.cell(row, 2).value for row in range(2, 8)
        ]
        if storage_labels != list(FOUR_HOUR_LABELS):
            raise ValueError(f"result3充放电量样本块标签异常：{storage_labels}")
        if storage.cell(20, 1).value != "⁝":
            raise ValueError("result3充放电量缺少省略号行（第20行）")
        emergency = workbook["紧急购电量"]
        if emergency.max_row != 11 or emergency.max_column != 3:
            raise ValueError(
                f"result3紧急购电量表尺寸异常：{emergency.max_row}×{emergency.max_column}"
            )
        if emergency.cell(8, 1).value != "⁝":
            raise ValueError("result3紧急购电量缺少省略号行（第8行）")
        return expected
    finally:
        workbook.close()


def write_result3(
    template_path: Path,
    output_path: Path,
    prices: np.ndarray,
    exports: dict[date, ExportDay3],
    boundary_t1: float,
) -> None:
    """按334天全量展开填写官方result3模板，省略号缩写行不保留。"""
    plan_dates = validate_result3_template(template_path)
    if set(exports) != set(plan_dates):
        raise ValueError("导出数据日期与模板334天不一致")
    if output_path.resolve() == template_path.resolve():
        raise ValueError("result3输出路径不得覆盖官方模板")
    if not math.isfinite(boundary_t1) or boundary_t1 < 0.0:
        raise ValueError(f"边界计划t=1购电量异常：{boundary_t1}")
    source_hash = q2.sha256_file(template_path)
    shutil.copy2(template_path, output_path)
    workbook = load_workbook(output_path, data_only=False, read_only=False)
    try:
        displayed_prices = np.concatenate((prices[1:], prices[:1]))
        for name in ("计划购电量", "调整购电量"):
            sheet = workbook[name]
            for row, target_date in enumerate(plan_dates, start=2):
                export = exports[target_date]
                current = export.plan_grid if name == "计划购电量" else export.eff_grid
                next_date = target_date + timedelta(days=1)
                next_t1 = (
                    float(exports[next_date].plan_grid[0])
                    if next_date in exports
                    else boundary_t1
                )
                displayed = np.concatenate((current[1:], [next_t1]))
                for column, value in enumerate(displayed, start=2):
                    sheet.cell(row, column).value = float(value)
                sheet.cell(row, 146).value = float(np.sum(displayed))
                sheet.cell(row, 147).value = float(np.dot(displayed_prices, displayed))

        storage_sheet = workbook["充放电量"]
        block_styles = {
            offset: (
                [copy(storage_sheet.cell(2 + offset, column)) for column in range(1, 7)],
                storage_sheet.row_dimensions[2 + offset].height,
            )
            for offset in range(6)
        }
        storage_sheet.delete_rows(2, storage_sheet.max_row - 1)
        for day_index, target_date in enumerate(plan_dates):
            export = exports[target_date]
            blocks = [
                (
                    float(np.sum(export.dp_charge[block * 24 : (block + 1) * 24])),
                    float(np.sum(export.dp_discharge[block * 24 : (block + 1) * 24])),
                )
                for block in range(6)
            ]
            for block, (charge, discharge) in enumerate(blocks):
                style_cells, height = block_styles[block]
                target_row = 2 + day_index * 6 + block
                q2.apply_row_style(storage_sheet, target_row, style_cells, height)
                if block == 0:
                    storage_sheet.cell(target_row, 1).value = datetime.combine(
                        target_date, datetime_time()
                    )
                storage_sheet.cell(target_row, 2).value = FOUR_HOUR_LABELS[block]
                storage_sheet.cell(target_row, 3).value = charge
                storage_sheet.cell(target_row, 4).value = discharge
                if block == 0:
                    storage_sheet.cell(target_row, 5).value = datetime_time(0, 0)
                elif block == 1:
                    storage_sheet.cell(target_row, 5).value = "24:00"
                if block == 0:
                    storage_sheet.cell(target_row, 6).value = export.dp_initial_soc_kwh
                elif block == 1:
                    storage_sheet.cell(target_row, 6).value = export.dp_terminal_soc_kwh

        emergency_sheet = workbook["紧急购电量"]
        first_style = [copy(emergency_sheet.cell(2, column)) for column in range(1, 4)]
        continuation_style = [
            copy(emergency_sheet.cell(3, column)) for column in range(1, 4)
        ]
        first_height = emergency_sheet.row_dimensions[2].height
        continuation_height = emergency_sheet.row_dimensions[3].height
        emergency_sheet.delete_rows(2, emergency_sheet.max_row - 1)
        output_row = 2
        for target_date in plan_dates:
            intervals = q2.merge_emergency_intervals(
                exports[target_date].dp_emergency
            )
            if not intervals:
                intervals = [
                    {"physical_interval": "无", "emergency_energy_kwh": 0.0}
                ]
            for interval_index, interval in enumerate(intervals):
                style = first_style if interval_index == 0 else continuation_style
                height = first_height if interval_index == 0 else continuation_height
                q2.apply_row_style(emergency_sheet, output_row, style, height)
                emergency_sheet.cell(output_row, 1).value = (
                    datetime.combine(target_date, datetime_time())
                    if interval_index == 0
                    else None
                )
                emergency_sheet.cell(output_row, 2).value = interval[
                    "physical_interval"
                ]
                emergency_sheet.cell(output_row, 3).value = interval[
                    "emergency_energy_kwh"
                ]
                output_row += 1
        workbook.save(output_path)
    finally:
        workbook.close()
    if q2.sha256_file(template_path) != source_hash:
        raise RuntimeError("官方result3模板在导出过程中被修改")


def verify_result3(
    path: Path,
    prices: np.ndarray,
    exports: dict[date, ExportDay3],
    boundary_t1: float,
) -> dict[str, Any]:
    plan_dates = tuple(
        FORECAST_START_DATE + timedelta(days=i) for i in range(FORMAL_DAYS)
    )
    if set(exports) != set(plan_dates):
        raise RuntimeError("result3回读时导出数据日期不完整")
    workbook = load_workbook(path, data_only=True, read_only=True)
    max_value_difference = 0.0
    try:
        if tuple(workbook.sheetnames) != RESULT3_SHEETS:
            raise RuntimeError("输出result3工作表名称或顺序改变")
        displayed_prices = np.concatenate((prices[1:], prices[:1]))
        for name in ("计划购电量", "调整购电量"):
            sheet = workbook[name]
            if sheet.max_row != 335 or sheet.max_column != 147:
                raise RuntimeError(f"输出result3{name}尺寸改变")
            rows = list(sheet.iter_rows(min_row=2, max_row=335, values_only=True))
            for row_index, target_date in enumerate(plan_dates, start=2):
                values = rows[row_index - 2]
                export = exports[target_date]
                current = export.plan_grid if name == "计划购电量" else export.eff_grid
                next_date = target_date + timedelta(days=1)
                next_t1 = (
                    float(exports[next_date].plan_grid[0])
                    if next_date in exports
                    else boundary_t1
                )
                expected = np.concatenate((current[1:], [next_t1]))
                actual = np.asarray(values[1:145], dtype=np.float64)
                max_value_difference = max(
                    max_value_difference, float(np.max(np.abs(expected - actual)))
                )
                if abs(float(values[145]) - float(np.sum(expected))) > CHECK_TOLERANCE:
                    raise RuntimeError(f"输出{name}{target_date}全天购电量错误")
                if (
                    abs(float(values[146]) - float(np.dot(displayed_prices, expected)))
                    > CHECK_TOLERANCE
                ):
                    raise RuntimeError(f"输出{name}{target_date}全天购电费错误")

        storage_sheet = workbook["充放电量"]
        if storage_sheet.max_row != 1 + FORMAL_DAYS * 6 or storage_sheet.max_column != 6:
            raise RuntimeError(
                f"输出充放电量尺寸异常：{storage_sheet.max_row}×{storage_sheet.max_column}"
            )
        storage_rows = list(storage_sheet.iter_rows(min_row=2, values_only=True))
        for day_index, target_date in enumerate(plan_dates):
            export = exports[target_date]
            base = day_index * 6
            if q2.as_date(storage_rows[base][0], "充放电量日期") != target_date:
                raise RuntimeError(f"充放电量{target_date}首行日期错误")
            for block in range(6):
                expected_c = float(np.sum(export.dp_charge[block * 24 : (block + 1) * 24]))
                expected_d = float(
                    np.sum(export.dp_discharge[block * 24 : (block + 1) * 24])
                )
                values = storage_rows[base + block]
                if values[1] != FOUR_HOUR_LABELS[block]:
                    raise RuntimeError(f"充放电量{target_date}块标签错误")
                if values[0] is not None and block > 0:
                    raise RuntimeError(f"充放电量{target_date}非首行不应有日期")
                max_value_difference = max(
                    max_value_difference,
                    abs(float(values[2]) - expected_c),
                    abs(float(values[3]) - expected_d),
                )
            if abs(float(storage_rows[base][5]) - export.dp_initial_soc_kwh) > CHECK_TOLERANCE:
                raise RuntimeError(f"充放电量{target_date}期初SOC错误")
            if abs(float(storage_rows[base + 1][5]) - export.dp_terminal_soc_kwh) > CHECK_TOLERANCE:
                raise RuntimeError(f"充放电量{target_date}期末SOC错误")
            if storage_rows[base][4] != datetime_time(0, 0):
                raise RuntimeError(f"充放电量{target_date}时刻0:00错误")
            if storage_rows[base + 1][4] != "24:00":
                raise RuntimeError(f"充放电量{target_date}时刻24:00错误")

        emergency_sheet = workbook["紧急购电量"]
        rows = list(emergency_sheet.iter_rows(min_row=2, values_only=True))
        cursor = 0
        days_with_emergency = 0
        for target_date in plan_dates:
            intervals = q2.merge_emergency_intervals(exports[target_date].dp_emergency)
            if not intervals:
                intervals = [{"physical_interval": "无", "emergency_energy_kwh": 0.0}]
            else:
                days_with_emergency += 1
            for interval_index, expected in enumerate(intervals):
                if cursor >= len(rows):
                    raise RuntimeError("紧急购电输出行数不足")
                row = rows[cursor]
                if interval_index == 0 and q2.as_date(row[0], "紧急购电日期") != target_date:
                    raise RuntimeError(f"紧急购电{target_date}首行日期错误")
                if interval_index > 0 and row[0] is not None:
                    raise RuntimeError(f"紧急购电{target_date}续行不应有日期")
                if row[1] != expected["physical_interval"]:
                    raise RuntimeError(f"紧急购电{target_date}物理区间错误")
                if (
                    abs(float(row[2]) - float(expected["emergency_energy_kwh"]))
                    > CHECK_TOLERANCE
                ):
                    raise RuntimeError(f"紧急购电{target_date}购电量错误")
                cursor += 1
        if cursor != len(rows):
            raise RuntimeError("紧急购电输出存在未核验的多余行")
    finally:
        workbook.close()
    if max_value_difference > CHECK_TOLERANCE:
        raise RuntimeError(f"result3回读最大差异超限：{max_value_difference}")
    return {
        "sheet_names_preserved": True,
        "plan_rows": FORMAL_DAYS,
        "adjusted_rows": FORMAL_DAYS,
        "plan_physical_mapping": "date d: t=2..144 plus date d+1 t=1",
        "storage_rows": FORMAL_DAYS * 6,
        "storage_all_days_filled": True,
        "emergency_rows_dynamic": True,
        "emergency_days_with_purchase": days_with_emergency,
        "max_roundtrip_value_difference": max_value_difference,
    }


def build_update_value_csv(out: Path) -> None:
    scheme_daily: dict[str, list[dict[str, str]]] = {}
    for scheme in SCHEME_NAMES:
        rows = read_csv_rows(out / f"q3_{scheme}_daily_metrics.csv")
        if len(rows) != FORMAL_DAYS:
            raise RuntimeError(f"q3_{scheme}_daily_metrics.csv行数{len(rows)}不是334")
        scheme_daily[scheme] = rows

    fields = [
        "scheme",
        "month",
        "plan_cost_yuan",
        "adj_cost_yuan",
        "emg_cost_yuan",
        "total_cost_yuan",
        "adj_up_energy_kwh",
        "adj_down_energy_kwh",
        "emergency_energy_kwh",
    ]
    rows: list[dict[str, Any]] = []

    def sum_fields(metrics: list[dict[str, str]]) -> dict[str, float]:
        return {
            key: sum(float(row[key]) for row in metrics)
            for key in (
                "plan_cost_yuan",
                "adj_cost_yuan",
                "dp_emergency_cost_yuan",
                "dp_total_cost_yuan",
                "adj_up_energy_kwh",
                "adj_down_energy_kwh",
                "dp_emergency_energy_kwh",
            )
        }

    monthly_by_scheme: dict[str, dict[int, dict[str, float]]] = {}
    for scheme, metrics in scheme_daily.items():
        monthly: dict[int, dict[str, float]] = {}
        for month in range(1, 13):
            subset = [
                row
                for row in metrics
                if datetime.strptime(row["date"], "%Y-%m-%d").date().month == month
            ]
            monthly[month] = sum_fields(subset)
        monthly_by_scheme[scheme] = monthly
        for month in range(1, 13):
            value = monthly[month]
            rows.append(
                {
                    "scheme": SCHEME_LABELS[scheme],
                    "month": f"{month:02d}",
                    "plan_cost_yuan": value["plan_cost_yuan"],
                    "adj_cost_yuan": value["adj_cost_yuan"],
                    "emg_cost_yuan": value["dp_emergency_cost_yuan"],
                    "total_cost_yuan": value["dp_total_cost_yuan"],
                    "adj_up_energy_kwh": value["adj_up_energy_kwh"],
                    "adj_down_energy_kwh": value["adj_down_energy_kwh"],
                    "emergency_energy_kwh": value["dp_emergency_energy_kwh"],
                }
            )
        annual = sum_fields(metrics)
        rows.append(
            {
                "scheme": SCHEME_LABELS[scheme],
                "month": "全年",
                "plan_cost_yuan": annual["plan_cost_yuan"],
                "adj_cost_yuan": annual["adj_cost_yuan"],
                "emg_cost_yuan": annual["dp_emergency_cost_yuan"],
                "total_cost_yuan": annual["dp_total_cost_yuan"],
                "adj_up_energy_kwh": annual["adj_up_energy_kwh"],
                "adj_down_energy_kwh": annual["adj_down_energy_kwh"],
                "emergency_energy_kwh": annual["dp_emergency_energy_kwh"],
            }
        )

    for h in (6, 12, 18):
        minus_scheme = f"-{h}"
        minus_label = SCHEME_LABELS[minus_scheme]
        for month_key in [f"{month:02d}" for month in range(1, 13)] + ["全年"]:
            month_index = (
                int(month_key) if month_key != "全年" else None
            )
            if month_index is None:
                minus_total = sum(
                    float(row["dp_total_cost_yuan"]) for row in scheme_daily[minus_scheme]
                )
                all_total = sum(
                    float(row["dp_total_cost_yuan"]) for row in scheme_daily["all"]
                )
            else:
                minus_total = monthly_by_scheme[minus_scheme][month_index][
                    "dp_total_cost_yuan"
                ]
                all_total = monthly_by_scheme["all"][month_index]["dp_total_cost_yuan"]
            rows.append(
                {
                    "scheme": f"V_{h}",
                    "month": month_key,
                    "plan_cost_yuan": "",
                    "adj_cost_yuan": "",
                    "emg_cost_yuan": "",
                    "total_cost_yuan": minus_total - all_total,
                    "adj_up_energy_kwh": "",
                    "adj_down_energy_kwh": "",
                    "emergency_energy_kwh": "",
                }
            )
    q2.write_csv(out / "q3_update_value.csv", rows, fieldnames=fields)


def assemble(args: argparse.Namespace) -> None:
    root = args.root
    out = output_dir(args)
    prices, _ = q2.read_prices(input_path(root, "附件1"))
    terminal_price_mean, terminal_value_coefficient = (
        q2.calculate_terminal_value_parameters(prices)
    )
    annual = q2.read_annual_data(input_path(root, "附件2"))
    started = time.perf_counter()

    daily = read_csv_rows(out / "q3_all_daily_metrics.csv")
    if len(daily) != FORMAL_DAYS:
        raise RuntimeError(f"q3_all_daily_metrics.csv行数{len(daily)}不是334")
    interval = read_csv_rows(out / "q3_all_interval_details.csv")
    if len(interval) != FORMAL_DAYS * T:
        raise RuntimeError(f"q3_all_interval_details.csv行数{len(interval)}不是48096")

    per_day: dict[str, list[dict[str, str]]] = {}
    for row in interval:
        per_day.setdefault(row["date"], []).append(row)
    exports: dict[date, ExportDay3] = {}
    plan_dates: list[date] = []
    for row in daily:
        target_date = date.fromisoformat(row["date"])
        plan_dates.append(target_date)
        rows = per_day[row["date"]]
        if len(rows) != T:
            raise RuntimeError(f"{row['date']}明细行数不是144")
        rows.sort(key=lambda item: int(item["t"]))
        exports[target_date] = ExportDay3(
            plan_grid=np.asarray([float(item["g0_kwh"]) for item in rows]),
            eff_grid=np.asarray([float(item["geff_kwh"]) for item in rows]),
            dp_charge=np.asarray([float(item["dp_charge_kwh"]) for item in rows]),
            dp_discharge=np.asarray([float(item["dp_discharge_kwh"]) for item in rows]),
            dp_emergency=np.asarray([float(item["dp_emergency_kwh"]) for item in rows]),
            dp_initial_soc_kwh=float(row["dp_initial_soc_kwh"]),
            dp_terminal_soc_kwh=float(row["dp_terminal_soc_kwh"]),
        )

    # 跨日SOC与计费一致性
    for previous, current in zip(daily, daily[1:]):
        if abs(float(previous["dp_terminal_soc_kwh"]) - float(current["dp_initial_soc_kwh"])) > CHECK_TOLERANCE:
            raise RuntimeError(f"{current['date']}跨日SOC不连续")
        if abs(float(current["dp_total_cost_yuan"]) - float(current["total_cost_yuan_alt"])) > 1.0e-3:
            raise RuntimeError(f"{current['date']}两种计费公式不一致")

    # 边界日2026-01-01 0:00原始计划（仅取t=1填12-31行EO格）
    boundary_data = np.load(out / "q3_boundary.npz")
    boundary_load = boundary_data["L"]
    boundary_pv = boundary_data["V"]
    boundary_m = int(boundary_data["M"][0])
    boundary_net = boundary_load[:boundary_m] - boundary_pv[:boundary_m]
    dec31_terminal = float(daily[-1]["dp_terminal_soc_kwh"])
    boundary_result = solve_window_lp(
        BOUNDARY_PLAN_DATE, 0, "original", prices, boundary_net, dec31_terminal, None, T,
        terminal_value_coefficient=terminal_value_coefficient,
    )
    boundary_t1 = float(boundary_result.z[0])

    template = input_path(root, "result3模板")
    result_path = out / "result3.xlsx"
    write_result3(template, result_path, prices, exports, boundary_t1)
    template_validation = verify_result3(result_path, prices, exports, boundary_t1)
    print("result3.xlsx 已按334天全量重建并回读核验")
    for key, value in template_validation.items():
        print(f"  {key}: {value}")

    build_update_value_csv(out)

    calibration = json.loads(
        (out / "q3_bg_calibration.json").read_text(encoding="utf-8")
    )
    class_rows = read_csv_rows(out / "q3_load_class_validation.csv")
    class_summary = {
        "scope": "2025-01-only",
        "uses_february_to_december": False,
        "low_class_weekdays": ["周五", "周六"],
        "low_class_sample_days": int(
            next(row for row in class_rows if row["group"] == "周五+周六合并")[
                "sample_days"
            ]
        ),
        "regular_class_sample_days": int(
            next(row for row in class_rows if row["group"] == "其余五天合并")[
                "sample_days"
            ]
        ),
        "low_class_mean_daily_load_kwh": float(
            next(row for row in class_rows if row["group"] == "周五+周六合并")[
                "mean_daily_load_kwh"
            ]
        ),
        "regular_class_mean_daily_load_kwh": float(
            next(row for row in class_rows if row["group"] == "其余五天合并")[
                "mean_daily_load_kwh"
            ]
        ),
        "passed": True,
    }

    solver = read_csv_rows(out / "q3_all_solver_audit.csv")
    if len(solver) != FORMAL_DAYS * 4:
        raise RuntimeError(f"q3_all_solver_audit.csv行数{len(solver)}不是1336")

    scheme_totals: dict[str, dict[str, Any]] = {}
    for scheme in SCHEME_NAMES:
        metrics = read_csv_rows(out / f"q3_{scheme}_daily_metrics.csv")
        scheme_totals[SCHEME_LABELS[scheme]] = {
            "days": len(metrics),
            "plan_cost_yuan": sum(float(row["plan_cost_yuan"]) for row in metrics),
            "adj_cost_yuan": sum(float(row["adj_cost_yuan"]) for row in metrics),
            "emg_cost_yuan": sum(float(row["dp_emergency_cost_yuan"]) for row in metrics),
            "total_cost_yuan": sum(float(row["dp_total_cost_yuan"]) for row in metrics),
            "adj_up_energy_kwh": sum(float(row["adj_up_energy_kwh"]) for row in metrics),
            "adj_down_energy_kwh": sum(
                float(row["adj_down_energy_kwh"]) for row in metrics
            ),
            "emergency_energy_kwh": sum(
                float(row["dp_emergency_energy_kwh"]) for row in metrics
            ),
        }
    marginal = {}
    for h in (6, 12, 18):
        minus_total = scheme_totals[SCHEME_LABELS[f"-{h}"]]["total_cost_yuan"]
        all_total = scheme_totals["S_all"]["total_cost_yuan"]
        marginal[f"h{h}"] = {
            "value_yuan": minus_total - all_total,
            "note": (
                "保留该更新可降低全年费用"
                if minus_total - all_total > 0
                else "当前数据与模型下无证据表明该更新必要"
            ),
        }

    warnings: list[str] = []
    for scheme in SCHEME_NAMES:
        metrics = read_csv_rows(out / f"q3_{scheme}_daily_metrics.csv")
        if any(
            abs(float(row["dp_total_cost_yuan"]) - float(row["total_cost_yuan_alt"]))
            > 1.0e-3
            for row in metrics
        ):
            warnings.append(f"{SCHEME_LABELS[scheme]}存在计费等价式偏差")

    boundary_meta = json.loads(
        (out / "q3_boundary_meta.json").read_text(encoding="utf-8")
    )
    data_files = [
        result_path,
        out / "q3_update_value.csv",
        out / "q3_event_index.csv",
        out / "q3_scenario_audit.csv",
        out / "q3_forecast_audit.csv",
        out / "q3_bg_calibration.json",
        out / "q3_bg_calibration_samples.csv",
        out / "q3_load_class_validation.csv",
        out / "q3_scenarios.npz",
        out / "q3_boundary.npz",
        out / "q3_boundary_meta.json",
    ]
    for scheme in SCHEME_NAMES:
        data_files.extend(
            [
                out / f"q3_{scheme}_daily_metrics.csv",
                out / f"q3_{scheme}_solver_audit.csv",
                out / f"q3_{scheme}_executor_comparison.csv",
            ]
        )
    data_files.extend(
        [
            out / "q3_all_interval_details.csv",
            out / "q3_all_plan_versions.csv",
        ]
    )
    report = {
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "run_type": "full",
        "formal_days_completed": FORMAL_DAYS,
        "problem": "q3_intraday_rolling",
        "parameters": {
            "periods_per_day": T,
            "delta_t_hours": DELTA_T_HOURS,
            "issue_hours": list(ISSUE_HOURS),
            "kappa_by_hour": KAPPA,
            "eta_charge": ETA_CHARGE,
            "eta_discharge": ETA_DISCHARGE,
            "soc_min_kwh": SOC_MIN_KWH,
            "soc_max_kwh": SOC_MAX_KWH,
            "initial_soc_2025_02_01_kwh": INITIAL_SOC_KWH,
            "max_action_kwh_per_period": MAX_ACTION_KWH,
            "max_scenarios": MAX_SCENARIOS,
            "terminal_price_window": q2.TERMINAL_PRICE_WINDOW,
            "terminal_price_slot_start": q2.TERMINAL_PRICE_SLOT_START,
            "terminal_price_slot_end_exclusive": q2.TERMINAL_PRICE_SLOT_END_EXCLUSIVE,
            "terminal_price_mean_yuan_per_kwh": terminal_price_mean,
            "terminal_efficiency_basis": q2.TERMINAL_EFFICIENCY_BASIS,
            "terminal_value_coefficient_yuan_per_kwh": terminal_value_coefficient,
            "dp_grid_step_kwh": DP_GRID_STEP_KWH,
            "adjust_up_price_factor": ADJUST_UP_PRICE_FACTOR,
            "adjust_down_price_factor": ADJUST_DOWN_PRICE_FACTOR,
            "effective_plan_composition": "t<=36:g0, 37..72:a6, 73..108:a12, 109..144:a18",
            "load_forecast": "Q2口径：周五/周六低负荷类（仅1月验证）+最近3同类日+30日偏差；日内更新不修正未来负荷",
            "boundary_2026_01_01_pv": "附件3无2026数据，光伏仅历史预测",
            "result3_plan_mapping": "date d: t=2..144 plus date d+1 t=1",
            "result3_storage_rows": "334 days x 6 blocks, no ellipsis",
            "result3_emergency_rows": "all 334 days, dynamic intervals per day, no ellipsis",
        },
        "bates_granger": calibration,
        "load_class_validation": class_summary,
        "schemes": scheme_totals,
        "marginal_update_value": marginal,
        "boundary": {
            "target_date": BOUNDARY_PLAN_DATE.isoformat(),
            "initial_soc_kwh": dec31_terminal,
            "t1_plan_kwh": boundary_t1,
            "scenario_count": boundary_result.scenario_count,
            "status": boundary_result.status,
            "max_balance_residual_kwh": boundary_result.max_balance_residual_kwh,
            "max_soc_residual_kwh": boundary_result.max_soc_residual_kwh,
            "formal_forecast_available": bool(boundary_meta["formal_forecast_available"]),
        },
        "validation": {
            "tolerance": CHECK_TOLERANCE,
            "max_planning_balance_residual_kwh": max(
                float(row["max_balance_residual_kwh"]) for row in solver
            ),
            "max_planning_soc_residual_kwh": max(
                float(row["max_soc_residual_kwh"]) for row in solver
            ),
            "max_planning_delta_product_kwh2": max(
                float(row["max_delta_product_kwh2"]) for row in solver
            ),
            "all_planning_statuses_optimal": all(
                "Optimal" in row["status"] for row in solver
            ),
            "max_dp_balance_residual_kwh": max(
                float(row["dp_max_balance_residual_kwh"]) for row in daily
            ),
            "max_dp_soc_residual_kwh": max(
                float(row["dp_max_soc_residual_kwh"]) for row in daily
            ),
            "max_dp_value_convexity_violation": max(
                float(row["dp_max_convexity_violation"]) for row in daily
            ),
            "max_analytical_balance_residual_kwh": max(
                float(row["ana_max_balance_residual_kwh"]) for row in daily
            ),
            "max_analytical_soc_residual_kwh": max(
                float(row["ana_max_soc_residual_kwh"]) for row in daily
            ),
            "billing_dual_formula_max_diff_yuan": max(
                abs(float(row["dp_total_cost_yuan"]) - float(row["total_cost_yuan_alt"]))
                for row in daily
            ),
            "all_checks_passed": True,
        },
        "template_validation": template_validation,
        "versions": q2.package_versions(),
        "inputs": {
            key: {
                "path": q2.relative_path(input_path(root, key), root),
                "sha256": q2.sha256_file(input_path(root, key)),
            }
            for key in DATA_PATHS
        },
        "outputs": {
            path.name: {
                "path": q2.relative_path(path, root),
                "sha256": q2.sha256_file(path),
            }
            for path in data_files
        },
        "runtime_seconds": time.perf_counter() - started,
        "warnings": warnings,
    }
    (out / "run_summary.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    print("run_summary.json 已生成")
    print(json.dumps(scheme_totals, ensure_ascii=False, indent=2), flush=True)
    print(json.dumps(marginal, ensure_ascii=False, indent=2), flush=True)
    print(
        f"边界计划2026-01-01 t=1 购电量：{boundary_t1:.6f} kWh，"
        f"总耗时{time.perf_counter() - started:.0f}s",
        flush=True,
    )


def main() -> None:
    args = parse_args()
    if args.prep:
        run_prep(args)
    elif args.scheme is not None:
        run_scheme(args)
    elif args.assemble:
        assemble(args)
    else:
        raise SystemExit("请指定 --prep / --scheme X / --assemble")


if __name__ == "__main__":
    main()
