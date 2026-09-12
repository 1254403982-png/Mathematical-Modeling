"""Q3产物独立复核（正式运行后运行，只读不改产物）。

覆盖：
 1. 附件/模板哈希未变
 2. prep产物：事件索引公式与覆盖、场景池因果实现规则与等权、BG权重复算
 3. 每方案：更新集合正确、LP全部Optimal且无回退、残差≤1e-6、计费双公式对账、
    SOC跨日连续且在[1200,10800]、DP执行残差
 4. 主方案冻结与组合：geff = g0(t≤36)/a6(37..72)/a12(73..108)/a18(109..144)、
    δ链接 a−g0=δ⁺−δ⁻、interval与plan_versions一致
 5. result3.xlsx：独立重建exports回读核验 + 边界日EO格复算 + 充放电/紧急表结构
 6. q3_update_value.csv：全年费用与V_h从方案CSV复算
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import q2  # noqa: E402
import q3  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "output" / "q3"
ATTACH1 = ROOT / "data" / "附件" / "附件1.xlsx"
ATTACH2 = ROOT / "data" / "附件" / "附件2.xlsx"
ATTACH3 = ROOT / "data" / "附件" / "附件3.xlsx"
TEMPLATE3 = ROOT / "data" / "附件" / "附件5" / "result3.xlsx"
RESULT3 = OUT / "result3.xlsx"
FORMAL_DAYS = q3.FORMAL_DAYS
T = q3.T

HASHES = {
    "附件1": ATTACH1,
    "附件2": ATTACH2,
    "附件3": ATTACH3,
    "result3模板": TEMPLATE3,
}
SCHEMES = ("0", "all", "-6", "-12", "-18")
FAILED: list[str] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f"  ({detail})" if detail else ""))
    if not ok:
        FAILED.append(name)


# --------------------------------------------------------------------- #
# 1. 哈希与prep产物
# --------------------------------------------------------------------- #
def verify_hashes(summary: dict) -> None:
    inputs = summary.get("inputs", {})
    for label, path in HASHES.items():
        actual = q2.sha256_file(path)
        recorded = inputs.get(label, {}).get("sha256")
        check(f"哈希[{label}]", actual == recorded,
              f"实际={actual[:16]}… 记录={recorded[:16] if recorded else '缺失'}")
    for name in ("q3_scenarios.npz", "q3_event_index.csv", "q3_boundary.npz"):
        check(f"prep产物存在[{name}]", (OUT / name).is_file())


def verify_prep(summary: dict) -> None:
    events = q3.read_csv_rows(OUT / "q3_event_index.csv")
    check("事件索引行数=1336", len(events) == 1336, f"实际{len(events)}")
    idx_ok = all(
        int(row["event_index"]) == di
        for di, row in enumerate(events)
    )
    check("事件索引与行序一致0..1335", idx_ok)
    check("HOUR_INDEX={0:0,6:1,12:2,18:3}",
          q3.HOUR_INDEX == {0: 0, 6: 1, 12: 2, 18: 3})
    check("事件覆盖2-1..12-31",
          events[0]["date"] == "2025-02-01" and events[-1]["date"] == "2025-12-31")
    hours = [int(row["issue_hour"]) for row in events]
    check("每日4事件按0/6/12/18",
          all(hours[i * 4:(i + 1) * 4] == [0, 6, 12, 18] for i in range(FORMAL_DAYS)))
    check("场景数0..30",
          all(0 <= int(row["scenario_count"]) <= 30 for row in events))

    # 场景池审计：每行=(事件,场景rank,残差来源,权重)；因果实现规则与等权
    audit = q3.read_csv_rows(OUT / "q3_scenario_audit.csv")
    causal_ok, equal_ok, size_ok = True, True, True
    by_event: dict[tuple[str, int], list[dict]] = defaultdict(list)
    for row in audit:
        by_event[(row["issue_date"], int(row["issue_hour"]))].append(row)
    for (issue_date, issue_hour), rows in by_event.items():
        size = len(rows)
        for row in rows:
            d_issue = date.fromisoformat(issue_date)
            d_res = date.fromisoformat(row["residual_date"])
            gap_h = ((d_issue - d_res).days * 24 + issue_hour
                     - int(row["residual_hour"]))
            if gap_h < 24:
                causal_ok = False
            if abs(float(row["weight"]) - 1.0 / size) > 1e-9:
                equal_ok = False
    for event_row in events:
        key = (event_row["date"], int(event_row["issue_hour"]))
        if len(by_event.get(key, [])) != int(event_row["scenario_count"]):
            size_ok = False
    check("场景池因果(历史发布时间+24h≤当前)", causal_ok)
    check("场景审计行数与索引一致", size_ok)
    check("场景等权1/n", equal_ok)

    # 场景npz形状
    npz = np.load(OUT / "q3_scenarios.npz")
    L, V = npz["L"], npz["V"]
    check("场景形状(1336,30,144)",
          L.shape == (1336, 30, 144) and V.shape == (1336, 30, 144),
          f"实际L={L.shape} V={V.shape}")
    check("场景负荷/光伏非负", L.min() >= 0 and V.min() >= -1e-9)
    npz.close()

    # BG权重复算
    samples = q3.read_csv_rows(OUT / "q3_bg_calibration_samples.csv")
    ef = np.array([float(r["error_formal_kwh"]) for r in samples])
    eh = np.array([float(r["error_historical_kwh"]) for r in samples])
    denom = float(np.sum((ef - eh) ** 2))
    w = 0.5 if denom == 0 else float(np.sum(eh * (eh - ef)) / denom)
    w = min(max(w, 0.0), 1.0)
    if "bates_granger" in summary:
        recorded = float(summary["bates_granger"]["w"])
        check("BG权重复算一致", abs(w - recorded) < 1e-9,
              f"复算={w:.9f} 记录={recorded:.9f}")
    else:
        check("BG权重复算一致", False, f"复算={w:.9f} 但run_summary缺失")

    # 负荷分类验证（1月only）
    classes = q3.read_csv_rows(OUT / "q3_load_class_validation.csv")
    groups = {r["group"] for r in classes}
    check("1月负荷分类覆盖周五周六与其余",
          "周五+周六合并" in groups and "其余五天合并" in groups)


# --------------------------------------------------------------------- #
# 2. 各方案
# --------------------------------------------------------------------- #
def verify_scheme(scheme: str) -> None:
    daily_path = OUT / f"q3_{scheme}_daily_metrics.csv"
    solver_path = OUT / f"q3_{scheme}_solver_audit.csv"
    exec_path = OUT / f"q3_{scheme}_executor_comparison.csv"
    missing = [p.name for p in (daily_path, solver_path, exec_path) if not p.is_file()]
    if missing:
        check(f"[{scheme}]产物齐全", False, f"缺{missing}")
        return
    daily = q3.read_csv_rows(daily_path)
    solver = q3.read_csv_rows(solver_path)
    execr = q3.read_csv_rows(exec_path)
    expect_hours = {"0": (0,), "all": (0, 6, 12, 18),
                    "-6": (0, 12, 18), "-12": (0, 6, 18), "-18": (0, 6, 12)}[scheme]
    check(f"[{scheme}]daily行数=334", len(daily) == 334, f"实际{len(daily)}")
    check(f"[{scheme}]solver行数=334×{len(expect_hours)}",
          len(solver) == 334 * len(expect_hours), f"实际{len(solver)}")
    check(f"[{scheme}]executor行数=334", len(execr) == 334, f"实际{len(execr)}")
    got = [int(r["issue_hour"]) for r in solver]
    check(f"[{scheme}]更新集合正确", got == list(expect_hours) * 334,
          f"前4个={got[:4]} 期望={expect_hours}")
    check(f"[{scheme}]updates_solved一致",
          all(int(r["updates_solved"]) == len(expect_hours) for r in daily))

    # 计费双公式对账与三分量
    alt_ok = all(
        abs(float(r["dp_total_cost_yuan"]) - float(r["total_cost_yuan_alt"])) < 1e-4
        for r in daily
    )
    check(f"[{scheme}]计费双公式对账(1e-4)", alt_ok)
    comp_ok = all(
        abs(float(r["plan_cost_yuan"]) + float(r["adj_cost_yuan"])
            + float(r["dp_emergency_cost_yuan"])
            - float(r["dp_total_cost_yuan"])) < 1e-3
        for r in daily
    )
    check(f"[{scheme}]plan+adj+emg=total", comp_ok)

    # SOC
    check(f"[{scheme}]DP SOC跨日连续",
          all(abs(float(daily[i]["dp_terminal_soc_kwh"])
                  - float(daily[i + 1]["dp_initial_soc_kwh"])) < 1e-6
              for i in range(len(daily) - 1)))
    check(f"[{scheme}]DP期初SOC=6000",
          abs(float(daily[0]["dp_initial_soc_kwh"]) - 6000.0) < 1e-6)
    check(f"[{scheme}]SOC∈[1200,10800]",
          all(1200.0 - 1e-6 <= float(r["dp_terminal_soc_kwh"]) <= 10800.0 + 1e-6
              and 1200.0 - 1e-6 <= float(r["dp_initial_soc_kwh"]) <= 10800.0 + 1e-6
              for r in daily))

    # solver：全部Optimal（含tiebreak二次求解标记）、无回退、残差
    check(f"[{scheme}]LP全部Optimal", all("Optimal" in r["status"] for r in solver))
    check(f"[{scheme}]无二元互斥回退",
          all(r["binary_fallback_used"] == "False" for r in solver))
    bres = max(float(r["max_balance_residual_kwh"]) for r in solver)
    sres = max(float(r["max_soc_residual_kwh"]) for r in solver)
    dres = max(float(r["max_delta_product_kwh2"] or 0) for r in solver)
    check(f"[{scheme}]LP残差≤1e-6",
          bres <= 1e-6 and sres <= 1e-6 and dres <= 1e-8,
          f"bal={bres:.2e} soc={sres:.2e} δδ={dres:.2e}")

    # DP执行残差
    dp_bal = max(float(r["dp_max_balance_residual_kwh"]) for r in daily)
    dp_soc = max(float(r["dp_max_soc_residual_kwh"]) for r in daily)
    check(f"[{scheme}]DP执行残差≤1e-6",
          dp_bal <= 1e-6 and dp_soc <= 1e-6,
          f"bal={dp_bal:.2e} soc={dp_soc:.2e}")


# --------------------------------------------------------------------- #
# 3. 主方案冻结/组合/δ链接
# --------------------------------------------------------------------- #
def verify_freeze_and_composition() -> None:
    pv = q3.read_csv_rows(OUT / "q3_all_plan_versions.csv")
    check("plan_versions行数=192384", len(pv) == 192384, f"实际{len(pv)}")
    interval = q3.read_csv_rows(OUT / "q3_all_interval_details.csv")
    check("interval_details行数=48096", len(interval) == 48096, f"实际{len(interval)}")

    by_day: dict[str, dict[str, dict[int, float]]] = defaultdict(
        lambda: defaultdict(dict))
    for r in pv:
        by_day[r["date"]][r["variable_type"]][int(r["absolute_t"])] = float(r["plan_kwh"])

    delta_ok = True
    max_delta = 0.0
    for r in pv:
        if r["variable_type"] in ("a6", "a12", "a18"):
            diff = abs(float(r["plan_kwh"]) - float(r["g0_reference_kwh"])
                       - (float(r["delta_plus_kwh"]) - float(r["delta_minus_kwh"])))
            max_delta = max(max_delta, diff)
            if diff > 1e-5:
                delta_ok = False
    check("δ链接a−g0=δ⁺−δ⁻", delta_ok, f"max={max_delta:.2e}")

    interval_by_day = defaultdict(list)
    for r in interval:
        interval_by_day[r["date"]].append(r)
    freeze_ok = True
    comp_ok = True
    g0_consistent = True
    max_freeze = 0.0
    for day, rows in interval_by_day.items():
        rows.sort(key=lambda r: int(r["t"]))
        g0_i = np.array([float(r["g0_kwh"]) for r in rows])
        geff_i = np.array([float(r["geff_kwh"]) for r in rows])
        expect = np.empty(T)
        expect[:36] = g0_i[:36]                       # 0:00计划执行段
        expect[36:72] = [by_day[day]["a6"][t] for t in range(37, 73)]
        expect[72:108] = [by_day[day]["a12"][t] for t in range(73, 109)]
        expect[108:] = [by_day[day]["a18"][t] for t in range(109, 145)]
        dmax = float(np.abs(geff_i - expect).max())
        max_freeze = max(max_freeze, dmax)
        if dmax > 1e-6:
            freeze_ok = False
        if float(np.abs(g0_i - np.array(
                [by_day[day]["g0"][t] for t in range(1, 145)])).max()) > 1e-9:
            g0_consistent = False
    check("geff组合=t≤36g0/37..72a6/73..108a12/109..144a18",
          freeze_ok, f"max_diff={max_freeze:.2e}")
    check("interval的g0与plan_versions一致", g0_consistent)


# --------------------------------------------------------------------- #
# 4. result3.xlsx
# --------------------------------------------------------------------- #
def verify_result3(summary: dict, prices: np.ndarray) -> None:
    if not RESULT3.is_file():
        check("result3.xlsx存在", False)
        return
    daily = q3.read_csv_rows(OUT / "q3_all_daily_metrics.csv")
    if len(daily) != FORMAL_DAYS:
        check("result3回读核验（需334天完整数据）", False,
              f"daily仅{len(daily)}天")
        return
    interval = q3.read_csv_rows(OUT / "q3_all_interval_details.csv")
    per_day: dict[str, list] = defaultdict(list)
    for row in interval:
        per_day[row["date"]].append(row)
    exports = {}
    for row in daily:
        target = date.fromisoformat(row["date"])
        rows = sorted(per_day[row["date"]], key=lambda r: int(r["t"]))
        exports[target] = q3.ExportDay3(
            plan_grid=np.asarray([float(r["g0_kwh"]) for r in rows]),
            eff_grid=np.asarray([float(r["geff_kwh"]) for r in rows]),
            dp_charge=np.asarray([float(r["dp_charge_kwh"]) for r in rows]),
            dp_discharge=np.asarray([float(r["dp_discharge_kwh"]) for r in rows]),
            dp_emergency=np.asarray([float(r["dp_emergency_kwh"]) for r in rows]),
            dp_initial_soc_kwh=float(row["dp_initial_soc_kwh"]),
            dp_terminal_soc_kwh=float(row["dp_terminal_soc_kwh"]),
        )

    # 独立复算边界日2026-01-01的t=1计划
    boundary = np.load(OUT / "q3_boundary.npz")
    bL, bV, bM = boundary["L"], boundary["V"], int(boundary["M"][0])
    boundary_net = bL[:bM] - bV[:bM]
    dec31_terminal = float(daily[-1]["dp_terminal_soc_kwh"])
    boundary_res = q3.solve_window_lp(
        q3.BOUNDARY_PLAN_DATE, 0, "original", prices, boundary_net,
        dec31_terminal, None, T,
        terminal_value_coefficient=q2.terminal_value_coefficient(prices),
    )
    boundary_t1 = float(boundary_res.z[0])
    check("边界日LP复算Optimal", boundary_res.status == "Optimal",
          f"status={boundary_res.status}")
    if "boundary" in summary:
        check("边界日t1与run_summary一致",
              abs(boundary_t1 - float(summary["boundary"]["t1_plan_kwh"])) < 1e-6,
              f"复算={boundary_t1:.4f} "
              f"记录={summary['boundary']['t1_plan_kwh']:.4f}")

    # 用独立重建的exports调用q3.verify_result3回读核验
    validation = q3.verify_result3(RESULT3, prices, exports, boundary_t1)
    check("result3四表名与顺序", validation["sheet_names_preserved"])
    check("result3计划/调整表144列与CSV一致",
          validation["max_roundtrip_value_difference"] <= 1e-6,
          f"max={validation['max_roundtrip_value_difference']:.2e}")
    check("result3充放电行数=2004", validation["storage_rows"] == 334 * 6,
          f"实际{validation['storage_rows']}")

    # 紧急购电日期集与CSV一致
    emg_csv = {r["date"] for r in daily
               if float(r["dp_emergency_energy_kwh"]) > 1e-6}
    wb = load_workbook(RESULT3, data_only=True, read_only=True)
    read_emg = set()
    for row in wb["紧急购电量"].iter_rows(min_row=2, values_only=True):
        if row[0] is not None and row[1] not in (None, "无", ""):
            read_emg.add(str(row[0])[:10])
    wb.close()
    check("result3紧急购电日期集与CSV一致",
          emg_csv == read_emg,
          f"CSV={len(emg_csv)}天 表={len(read_emg)}天")

    # 边界日EO格直接核验：12-31行EO列(145) == boundary_t1
    wb = load_workbook(RESULT3, data_only=True, read_only=True)
    ws = wb["计划购电量"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    eo_plan = float(rows[333][144])  # 12-31行 EO格
    check("12-31行EO格=边界日t1计划", abs(eo_plan - boundary_t1) < 1e-6,
          f"表={eo_plan:.4f} 复算={boundary_t1:.4f}")
    # 其余日EO格=次日g0的t=1
    eo_ok = True
    for i in range(FORMAL_DAYS - 1):
        next_plan_t1 = exports[date.fromisoformat(daily[i + 1]["date"])].plan_grid[0]
        if abs(float(rows[i][144]) - next_plan_t1) > 1e-6:
            eo_ok = False
            break
    check("EO格=次日g0[t=1]", eo_ok)


# --------------------------------------------------------------------- #
# 5. update_value
# --------------------------------------------------------------------- #
# --------------------------------------------------------------------- #
# 6. 终值系数
# --------------------------------------------------------------------- #
def verify_terminal(summary: dict, prices: np.ndarray) -> None:
    params = summary.get("parameters", {})
    mean = float(np.mean(prices[:30]))
    coefficient = mean / q2.ETA_CHARGE
    check("凌晨00:00-05:00均价=0.433393333333",
          np.isclose(mean, 0.43339333333333335, atol=1e-12))
    recorded = params.get("terminal_value_coefficient_yuan_per_kwh")
    if recorded is None:
        check("终值系数符合定案", False, "run_summary缺少该字段")
        check("run_summary终值窗口与效率口径", False, "run_summary缺少该字段")
    else:
        recorded = float(recorded)
        check("终值系数符合定案",
              abs(recorded - 0.4815481481481482) < 1e-9
              and not np.isclose(recorded, 0.33417, atol=1e-8),
              f"记录={recorded:.12f}")
        check("run_summary终值窗口与效率口径",
              params.get("terminal_price_window") == "physical 00:00-05:00"
              and params.get("terminal_price_slot_start") == 0
              and params.get("terminal_price_slot_end_exclusive") == 30
              and params.get("terminal_efficiency_basis") == "charge"
              and np.isclose(
                  float(params.get("terminal_price_mean_yuan_per_kwh", float("nan"))),
                  mean,
                  atol=1e-12,
              ))
    check("run_summary all_checks_passed=true",
          summary.get("validation", {}).get("all_checks_passed") is True)
    annual = sum(
        float(row["dp_total_cost_yuan"])
        for row in q3.read_csv_rows(OUT / "q3_all_daily_metrics.csv")
    )
    check("Q3 全年费用量级对照完整稿(-5%内)",
          abs(annual / 13024699.25 - 1) < 0.05, f"复算={annual:,.2f}")


def verify_update_value() -> None:
    path = OUT / "q3_update_value.csv"
    if not path.is_file():
        check("q3_update_value.csv存在", False)
        return
    rows = q3.read_csv_rows(path)
    totals = {}
    for s in SCHEMES:
        d = q3.read_csv_rows(OUT / f"q3_{s}_daily_metrics.csv")
        totals[s] = sum(float(r["dp_total_cost_yuan"]) for r in d)
    labels = {"0": "S0", "all": "S_all", "-6": "S_-6", "-12": "S_-12", "-18": "S_-18"}
    annual = {r["scheme"]: r for r in rows if r["month"] == "全年" and r["scheme"] in labels.values()}
    ok = True
    for s, label in labels.items():
        if label not in annual:
            ok = False
            continue
        if abs(float(annual[label]["total_cost_yuan"]) - totals[s]) > 0.5:
            ok = False
    check("update_value全年费用与daily一致", ok,
          f"表S_all={float(annual['S_all']['total_cost_yuan']):.2f} "
          f"复算={totals['all']:.2f}")
    for h in (6, 12, 18):
        v_row = next((r for r in rows
                      if r["scheme"] == f"V_{h}" and r["month"] == "全年"), None)
        if v_row is None:
            check(f"V_{h}全年行存在", False)
            continue
        expect = totals[f"-{h}"] - totals["all"]
        check(f"V_{h}=C(S_-{h})−C(S_all)复算", abs(float(v_row["total_cost_yuan"]) - expect) < 0.5,
              f"表={float(v_row['total_cost_yuan']):.2f} 复算={expect:.2f}")


def main() -> None:
    global OUT, RESULT3
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, default=None,
                        help="输出目录，默认output/q3")
    args = parser.parse_args()
    if args.out is not None:
        OUT = args.out.resolve()
        RESULT3 = OUT / "result3.xlsx"
    print("=" * 70)
    print("Q3 产物独立复核")
    print(f"输出目录：{OUT}")
    print("=" * 70)
    summary_path = OUT / "run_summary.json"
    if not summary_path.is_file():
        print("警告：未找到 run_summary.json（哈希与BG权重对比将按缺失处理）")
        summary: dict = {}
    else:
        with open(summary_path, encoding="utf-8") as fh:
            summary = json.load(fh)
    prices, _ = q2.read_prices(ATTACH1)

    verify_hashes(summary)
    verify_prep(summary)
    verify_terminal(summary, prices)
    for scheme in SCHEMES:
        verify_scheme(scheme)
    verify_freeze_and_composition()
    verify_result3(summary, prices)
    verify_update_value()

    print("=" * 70)
    if FAILED:
        print(f"复核未通过：{len(FAILED)} 项")
        for name in FAILED:
            print(f"  - {name}")
        sys.exit(1)
    print("全部复核通过")


if __name__ == "__main__":
    main()
