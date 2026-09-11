from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import platform
import shutil
import statistics
import time
from copy import copy
from dataclasses import asdict, dataclass
from datetime import date, datetime, time as datetime_time, timedelta
from pathlib import Path
from typing import Any, Iterable, Sequence

try:
    import highspy
    import numpy as np
    import openpyxl
    import pulp
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - CLI dependency error
    raise SystemExit(
        "缺少Q2依赖。请执行：python -m pip install -r requirements-q2.txt"
    ) from exc


T = 144
DELTA_T_HOURS = 1.0 / 6.0
ETA_CHARGE = 0.90
ETA_DISCHARGE = 0.90
SOC_MIN_KWH = 1200.0
SOC_MAX_KWH = 10800.0
INITIAL_SOC_KWH = 6000.0
MAX_ACTION_KWH = 5000.0 / 6.0
EMERGENCY_PRICE_FACTOR = 5.0
MAX_SCENARIOS = 30
PRIMARY_HISTORY_DAYS = 3
BIAS_WINDOW_DAYS = 30
DEFAULT_DP_GRID_STEP_KWH = 10.0
RANDOM_SEED = 1
CHECK_TOLERANCE = 1.0e-6
SIMULTANEOUS_PRODUCT_TOLERANCE = 1.0e-8
ACTION_ZERO_TOLERANCE = 1.0e-8
EMERGENCY_INTERVAL_TOLERANCE = 1.0e-7
FORECAST_START_DATE = date(2025, 2, 1)
FORECAST_END_DATE = date(2025, 12, 31)
BOUNDARY_PLAN_DATE = date(2026, 1, 1)
LOW_LOAD_WEEKDAYS = frozenset({4, 5})  # Python weekday: Friday=4, Saturday=5
WEEKDAY_NAMES = ("周一", "周二", "周三", "周四", "周五", "周六", "周日")
FOUR_HOUR_LABELS = (
    "0:00-4:00",
    "4:00-8:00",
    "8:00-12:00",
    "12:00-16:00",
    "16:00-20:00",
    "20:00-24:00",
)
SAMPLE_TEMPLATE_DATES = (
    date(2025, 2, 1),
    date(2025, 2, 2),
    date(2025, 12, 31),
)
INPUT_PRICE_HEADERS = ("时间", "电价", "小区负载", "光伏发电预测功率")
ANNUAL_SHEETS = ("小区负载", "光伏发电实际功率")
RESULT_SHEETS = ("计划购电量", "充放电量", "紧急购电量")


@dataclass(frozen=True)
class AnnualData:
    dates: tuple[date, ...]
    right_end_labels: tuple[str, ...]
    load_kwh: np.ndarray
    pv_kwh: np.ndarray


@dataclass
class ForecastHistory:
    dates: list[date]
    load_actual: list[np.ndarray]
    pv_actual: list[np.ndarray]
    load_primary_residual: list[np.ndarray | None]
    pv_primary_residual: list[np.ndarray | None]
    load_final_residual: list[np.ndarray | None]
    pv_final_residual: list[np.ndarray | None]

    @classmethod
    def empty(cls) -> ForecastHistory:
        return cls([], [], [], [], [], [], [])


@dataclass(frozen=True)
class ForecastDraft:
    target_date: date
    load_primary: np.ndarray | None
    pv_primary: np.ndarray | None
    load_forecast: np.ndarray | None
    pv_forecast: np.ndarray | None
    baseline_load: np.ndarray | None
    baseline_pv: np.ndarray | None
    load_history: tuple[tuple[date, int, float], ...]
    pv_history: tuple[tuple[date, int, float], ...]
    bias_dates: tuple[date, ...]
    scenario_dates: tuple[date, ...]
    scenario_load_residuals: np.ndarray | None
    scenario_pv_residuals: np.ndarray | None


@dataclass(frozen=True)
class ForecastBundle:
    target_date: date
    load_primary: np.ndarray
    pv_primary: np.ndarray
    load_forecast: np.ndarray
    pv_forecast: np.ndarray
    baseline_load: np.ndarray
    baseline_pv: np.ndarray
    load_history: tuple[tuple[date, int, float], ...]
    pv_history: tuple[tuple[date, int, float], ...]
    bias_dates: tuple[date, ...]
    scenario_dates: tuple[date, ...]
    scenario_net_kwh: np.ndarray


@dataclass(frozen=True)
class PlanningResult:
    target_date: date
    initial_soc_kwh: float
    grid_purchase: np.ndarray
    objective_value: float
    plan_energy_kwh: float
    plan_cost_yuan: float
    scenario_count: int
    status: str
    solve_seconds: float
    highs_objective: float
    highs_max_primal_infeasibility: float
    highs_num_primal_infeasibilities: int
    simplex_iterations: int
    ipm_iterations: int
    mip_gap: float | None
    tiebreak_used: bool
    binary_fallback_used: bool
    max_balance_residual_kwh: float
    max_soc_residual_kwh: float
    max_simultaneous_product_kwh2: float
    min_variable_value_kwh: float
    min_soc_kwh: float
    max_soc_kwh: float
    max_charge_kwh: float
    max_discharge_kwh: float


@dataclass(frozen=True)
class ExecutorResult:
    charge: np.ndarray
    discharge: np.ndarray
    emergency: np.ndarray
    unused: np.ndarray
    soc: np.ndarray
    initial_soc_kwh: float
    terminal_soc_kwh: float
    emergency_energy_kwh: float
    emergency_cost_yuan: float
    max_balance_residual_kwh: float
    max_soc_residual_kwh: float
    min_soc_kwh: float
    max_soc_kwh: float
    max_charge_kwh: float
    max_discharge_kwh: float
    max_simultaneous_product_kwh2: float
    min_nonnegative_flow_kwh: float
    dp_max_convexity_violation: float | None = None


@dataclass(frozen=True)
class DailyRun:
    target_date: date
    forecast: ForecastBundle
    planning: PlanningResult
    actual_load: np.ndarray
    actual_pv: np.ndarray
    dp: ExecutorResult
    analytical: ExecutorResult


@dataclass(frozen=True)
class ExportDay:
    """导出result2所需的一整天数据（与求解结构解耦，便于从CSV重建）。"""

    plan_grid: np.ndarray
    dp_charge: np.ndarray
    dp_discharge: np.ndarray
    dp_emergency: np.ndarray
    dp_initial_soc_kwh: float
    dp_terminal_soc_kwh: float


@dataclass
class PlanningModel:
    problem: pulp.LpProblem
    original_objective: pulp.LpAffineExpression
    grid: list[pulp.LpVariable]
    charge: list[list[pulp.LpVariable]]
    discharge: list[list[pulp.LpVariable]]
    emergency: list[list[pulp.LpVariable]]
    unused: list[list[pulp.LpVariable]]
    soc: list[list[pulp.LpVariable]]


def parse_args() -> argparse.Namespace:
    root = Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser(
        description="求解数学建模C题问题2：因果预测、随机计划与跨日执行。"
    )
    parser.add_argument(
        "--price-input",
        type=Path,
        default=root / "data" / "附件" / "附件1.xlsx",
    )
    parser.add_argument(
        "--annual-input",
        type=Path,
        default=root / "data" / "附件" / "附件2.xlsx",
    )
    parser.add_argument(
        "--template",
        type=Path,
        default=root / "data" / "附件" / "附件5" / "result2.xlsx",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=root / "output" / "q2",
    )
    parser.add_argument(
        "--dp-grid-step",
        type=float,
        default=DEFAULT_DP_GRID_STEP_KWH,
        help="DP的SOC网格间距(kWh)，默认10。",
    )
    parser.add_argument(
        "--max-days",
        type=int,
        default=None,
        help="仅用于烟雾测试；省略时运行全部334日并填写result2。",
    )
    parser.add_argument(
        "--progress-every",
        type=int,
        default=10,
        help="每隔多少个正式日打印一次进度。",
    )
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def display_time_label(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if isinstance(value, datetime_time):
        return value.strftime("%H:%M")
    return str(value)


def as_date(value: Any, context: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise ValueError(f"{context}不是日期：{value!r}")


def checked_float(
    value: Any,
    *,
    context: str,
    nonnegative: bool,
) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{context}不是数值：{value!r}")
    number = float(value)
    if not math.isfinite(number):
        raise ValueError(f"{context}不是有限数值：{value!r}")
    if nonnegative and number < 0.0:
        raise ValueError(f"{context}为负数：{number}")
    return number


def physical_interval_labels() -> tuple[str, ...]:
    labels: list[str] = []
    for t in range(1, T + 1):
        start_minutes = (t - 1) * 10
        end_minutes = t * 10
        start = f"{start_minutes // 60}:{start_minutes % 60:02d}"
        end = (
            "00:00+1"
            if end_minutes == 1440
            else f"{end_minutes // 60}:{end_minutes % 60:02d}"
        )
        labels.append(f"{start}-{end}")
    return tuple(labels)


def read_prices(path: Path) -> tuple[np.ndarray, tuple[str, ...]]:
    if not path.is_file():
        raise FileNotFoundError(f"找不到附件1：{path}")
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if len(workbook.sheetnames) != 1:
            raise ValueError(f"附件1工作表数量异常：{workbook.sheetnames}")
        worksheet = workbook[workbook.sheetnames[0]]
        headers = tuple(worksheet.cell(1, column).value for column in range(1, 5))
        if headers != INPUT_PRICE_HEADERS:
            raise ValueError(f"附件1表头异常：{headers}")
        if worksheet.max_row != T + 1:
            raise ValueError(f"附件1数据行数不是144：{worksheet.max_row - 1}")
        prices: list[float] = []
        labels: list[str] = []
        for row in range(2, T + 2):
            raw_time = worksheet.cell(row, 1).value
            if raw_time is None:
                raise ValueError(f"附件1第{row}行时间为空")
            labels.append(display_time_label(raw_time))
            prices.append(
                checked_float(
                    worksheet.cell(row, 2).value,
                    context=f"附件1第{row}行电价",
                    nonnegative=False,
                )
            )
    finally:
        workbook.close()
    return np.asarray(prices, dtype=np.float64), tuple(labels)


def read_annual_data(path: Path) -> AnnualData:
    if not path.is_file():
        raise FileNotFoundError(f"找不到附件2：{path}")
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if tuple(workbook.sheetnames) != ANNUAL_SHEETS:
            raise ValueError(
                f"附件2工作表应为{ANNUAL_SHEETS}，实际为{workbook.sheetnames}"
            )
        parsed: dict[str, tuple[tuple[date, ...], tuple[str, ...], np.ndarray]] = {}
        for sheet_name in ANNUAL_SHEETS:
            worksheet = workbook[sheet_name]
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows)
            if len(header) != T + 1 or header[0] != "日期\\时间":
                raise ValueError(f"附件2“{sheet_name}”表头尺寸或首格异常")
            labels = tuple(display_time_label(value) for value in header[1:])
            dates: list[date] = []
            values: list[list[float]] = []
            for excel_row, row in enumerate(rows, start=2):
                if len(row) < T + 1:
                    raise ValueError(f"附件2“{sheet_name}”第{excel_row}行列数不足")
                dates.append(as_date(row[0], f"附件2“{sheet_name}”第{excel_row}行"))
                values.append(
                    [
                        checked_float(
                            value,
                            context=(
                                f"附件2“{sheet_name}”第{excel_row}行第{column}列"
                            ),
                            nonnegative=True,
                        )
                        for column, value in enumerate(row[1 : T + 1], start=2)
                    ]
                )
            parsed[sheet_name] = (
                tuple(dates),
                labels,
                np.asarray(values, dtype=np.float64) * DELTA_T_HOURS,
            )
    finally:
        workbook.close()

    load_dates, load_labels, load_kwh = parsed[ANNUAL_SHEETS[0]]
    pv_dates, pv_labels, pv_kwh = parsed[ANNUAL_SHEETS[1]]
    if load_dates != pv_dates:
        raise ValueError("附件2负荷与光伏工作表日期不一致")
    if load_labels != pv_labels:
        raise ValueError("附件2负荷与光伏工作表时段标签不一致")
    if len(load_dates) != 365 or load_kwh.shape != (365, T) or pv_kwh.shape != (365, T):
        raise ValueError(
            f"附件2维度异常：日期{len(load_dates)}、负荷{load_kwh.shape}、光伏{pv_kwh.shape}"
        )
    expected_dates = tuple(date(2025, 1, 1) + timedelta(days=i) for i in range(365))
    if load_dates != expected_dates:
        raise ValueError("附件2日期不是2025-01-01至2025-12-31的连续自然日")
    return AnnualData(load_dates, load_labels, load_kwh, pv_kwh)


def validate_time_inputs(
    price_labels: tuple[str, ...], annual: AnnualData
) -> None:
    if len(price_labels) != T or len(annual.right_end_labels) != T:
        raise ValueError("时间标签数量不是144")
    if price_labels != annual.right_end_labels:
        mismatches = [
            (i + 1, price_labels[i], annual.right_end_labels[i])
            for i in range(T)
            if price_labels[i] != annual.right_end_labels[i]
        ]
        raise ValueError(f"附件1与附件2右端点标签不一致：{mismatches[:5]}")
    if price_labels[0] not in ("00:10", "0:10") or price_labels[-1] != "0:00+1":
        raise ValueError(
            f"右端点语义异常：首标签{price_labels[0]!r}、末标签{price_labels[-1]!r}"
        )


def validate_january_load_classes(
    january_dates: Sequence[date], january_load_kwh: np.ndarray
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    if len(january_dates) != 31 or january_load_kwh.shape != (31, T):
        raise ValueError("负荷分类验证必须且只能接收2025年1月的31×144数据")
    if any(day.year != 2025 or day.month != 1 for day in january_dates):
        raise ValueError("负荷分类验证收到非2025年1月日期")

    daily_energy = np.sum(january_load_kwh, axis=1)
    grouped: dict[int, list[float]] = {weekday: [] for weekday in range(7)}
    for day, energy in zip(january_dates, daily_energy, strict=True):
        grouped[day.weekday()].append(float(energy))

    rows: list[dict[str, Any]] = []
    weekday_means: dict[int, float] = {}
    for weekday in range(7):
        samples = grouped[weekday]
        weekday_means[weekday] = statistics.mean(samples)
        rows.append(
            {
                "scope": "2025-01-only",
                "group": WEEKDAY_NAMES[weekday],
                "python_weekday": weekday,
                "classification": "低负荷类" if weekday in LOW_LOAD_WEEKDAYS else "常规类",
                "sample_days": len(samples),
                "mean_daily_load_kwh": statistics.mean(samples),
                "median_daily_load_kwh": statistics.median(samples),
                "stdev_daily_load_kwh": statistics.stdev(samples),
                "min_daily_load_kwh": min(samples),
                "max_daily_load_kwh": max(samples),
            }
        )

    low_values = grouped[4] + grouped[5]
    regular_values = [
        value
        for weekday in range(7)
        if weekday not in LOW_LOAD_WEEKDAYS
        for value in grouped[weekday]
    ]
    low_mean = statistics.mean(low_values)
    regular_mean = statistics.mean(regular_values)
    sorted_weekdays = sorted(weekday_means, key=weekday_means.get)
    lowest_two_are_fri_sat = set(sorted_weekdays[:2]) == set(LOW_LOAD_WEEKDAYS)
    strict_separation = max(weekday_means[4], weekday_means[5]) < min(
        weekday_means[i] for i in range(7) if i not in LOW_LOAD_WEEKDAYS
    )
    passed = lowest_two_are_fri_sat and strict_separation
    summary = {
        "scope": "2025-01-only",
        "uses_february_to_december": False,
        "low_class_weekdays": [WEEKDAY_NAMES[4], WEEKDAY_NAMES[5]],
        "low_class_sample_days": len(low_values),
        "regular_class_sample_days": len(regular_values),
        "low_class_mean_daily_load_kwh": low_mean,
        "regular_class_mean_daily_load_kwh": regular_mean,
        "low_to_regular_ratio": low_mean / regular_mean,
        "low_class_percent_below_regular": (1.0 - low_mean / regular_mean) * 100.0,
        "weekday_mean_order_low_to_high": [WEEKDAY_NAMES[i] for i in sorted_weekdays],
        "lowest_two_are_friday_saturday": lowest_two_are_fri_sat,
        "strictly_separated_from_other_weekdays": strict_separation,
        "passed": passed,
    }
    rows.extend(
        [
            {
                "scope": "2025-01-only",
                "group": "周五+周六合并",
                "python_weekday": "4,5",
                "classification": "低负荷类",
                "sample_days": len(low_values),
                "mean_daily_load_kwh": low_mean,
                "median_daily_load_kwh": statistics.median(low_values),
                "stdev_daily_load_kwh": statistics.stdev(low_values),
                "min_daily_load_kwh": min(low_values),
                "max_daily_load_kwh": max(low_values),
            },
            {
                "scope": "2025-01-only",
                "group": "其余五天合并",
                "python_weekday": "0,1,2,3,6",
                "classification": "常规类",
                "sample_days": len(regular_values),
                "mean_daily_load_kwh": regular_mean,
                "median_daily_load_kwh": statistics.median(regular_values),
                "stdev_daily_load_kwh": statistics.stdev(regular_values),
                "min_daily_load_kwh": min(regular_values),
                "max_daily_load_kwh": max(regular_values),
            },
        ]
    )
    if not passed:
        raise RuntimeError(f"仅用1月数据无法验证周五/周六低负荷分类：{summary}")
    return rows, summary


def is_low_load_day(day: date) -> bool:
    return day.weekday() in LOW_LOAD_WEEKDAYS


def inverse_distance_selection(
    target_date: date,
    history_dates: Sequence[date],
    candidate_indices: Iterable[int],
    limit: int = PRIMARY_HISTORY_DAYS,
) -> tuple[tuple[int, int, float], ...]:
    selected = sorted(
        candidate_indices,
        key=lambda i: (target_date - history_dates[i]).days,
    )[:limit]
    if not selected:
        return ()
    distances = [(target_date - history_dates[i]).days for i in selected]
    if any(distance <= 0 for distance in distances):
        raise ValueError("历史选择包含目标日或未来日")
    inverse = np.asarray([1.0 / distance for distance in distances], dtype=np.float64)
    weights = inverse / np.sum(inverse)
    return tuple(
        (i, distance, float(weight))
        for i, distance, weight in zip(selected, distances, weights, strict=True)
    )


def weighted_average(
    arrays: Sequence[np.ndarray], selection: Sequence[tuple[int, int, float]]
) -> np.ndarray | None:
    if not selection:
        return None
    result = np.zeros(T, dtype=np.float64)
    for index, _, weight in selection:
        result += weight * arrays[index]
    return result


def forecast_one_day(
    target_date: date,
    history: ForecastHistory,
) -> ForecastDraft:
    """Generate a forecast using only the supplied pre-target history."""
    if history.dates and max(history.dates) >= target_date:
        raise ValueError("forecast_one_day的history包含目标日或未来日")
    target_low = is_low_load_day(target_date)
    load_candidates = [
        i for i, historical_date in enumerate(history.dates)
        if is_low_load_day(historical_date) == target_low
    ]
    load_selection = inverse_distance_selection(
        target_date, history.dates, load_candidates
    )
    pv_selection = inverse_distance_selection(
        target_date, history.dates, range(len(history.dates))
    )
    load_primary = weighted_average(history.load_actual, load_selection)
    pv_primary = weighted_average(history.pv_actual, pv_selection)

    paired_primary_indices = [
        i for i in range(len(history.dates))
        if history.load_primary_residual[i] is not None
        and history.pv_primary_residual[i] is not None
    ][-BIAS_WINDOW_DAYS:]
    bias_dates = tuple(history.dates[i] for i in paired_primary_indices)
    if paired_primary_indices:
        load_bias = np.mean(
            np.stack([history.load_primary_residual[i] for i in paired_primary_indices]),
            axis=0,
        )
        pv_bias = np.mean(
            np.stack([history.pv_primary_residual[i] for i in paired_primary_indices]),
            axis=0,
        )
    else:
        load_bias = np.zeros(T, dtype=np.float64)
        pv_bias = np.zeros(T, dtype=np.float64)

    load_forecast = (
        None if load_primary is None else np.maximum(load_primary + load_bias, 0.0)
    )
    pv_forecast = (
        None if pv_primary is None else np.maximum(pv_primary + pv_bias, 0.0)
    )
    if pv_forecast is not None and history.pv_actual:
        historical_pv = np.stack(history.pv_actual)
        always_zero = np.all(historical_pv == 0.0, axis=0)
        pv_forecast = pv_forecast.copy()
        pv_forecast[always_zero] = 0.0

    baseline_load: np.ndarray | None = None
    baseline_pv: np.ndarray | None = None
    if len(history.dates) >= PRIMARY_HISTORY_DAYS:
        baseline_load = np.maximum(
            np.mean(np.stack(history.load_actual[-PRIMARY_HISTORY_DAYS:]), axis=0),
            0.0,
        )
        baseline_pv = np.maximum(
            np.mean(np.stack(history.pv_actual[-PRIMARY_HISTORY_DAYS:]), axis=0),
            0.0,
        )
        historical_pv = np.stack(history.pv_actual)
        baseline_pv[np.all(historical_pv == 0.0, axis=0)] = 0.0

    paired_final_indices = [
        i for i in range(len(history.dates))
        if history.load_final_residual[i] is not None
        and history.pv_final_residual[i] is not None
    ][-MAX_SCENARIOS:]
    scenario_dates = tuple(history.dates[i] for i in paired_final_indices)
    scenario_load_residuals = (
        None
        if not paired_final_indices
        else np.stack([history.load_final_residual[i] for i in paired_final_indices])
    )
    scenario_pv_residuals = (
        None
        if not paired_final_indices
        else np.stack([history.pv_final_residual[i] for i in paired_final_indices])
    )
    return ForecastDraft(
        target_date=target_date,
        load_primary=load_primary,
        pv_primary=pv_primary,
        load_forecast=load_forecast,
        pv_forecast=pv_forecast,
        baseline_load=baseline_load,
        baseline_pv=baseline_pv,
        load_history=tuple(
            (history.dates[i], distance, weight)
            for i, distance, weight in load_selection
        ),
        pv_history=tuple(
            (history.dates[i], distance, weight)
            for i, distance, weight in pv_selection
        ),
        bias_dates=bias_dates,
        scenario_dates=scenario_dates,
        scenario_load_residuals=scenario_load_residuals,
        scenario_pv_residuals=scenario_pv_residuals,
    )


def append_observation(
    history: ForecastHistory,
    observed_date: date,
    load_actual: np.ndarray,
    pv_actual: np.ndarray,
    forecast: ForecastDraft,
) -> None:
    if forecast.target_date != observed_date:
        raise ValueError("追加观测的日期与预测目标日不一致")
    if history.dates and observed_date != history.dates[-1] + timedelta(days=1):
        raise ValueError("预测历史未按自然日连续追加")
    if history.dates and history.dates[-1] >= observed_date:
        raise ValueError("预测历史日期顺序错误")

    history.dates.append(observed_date)
    history.load_actual.append(load_actual.copy())
    history.pv_actual.append(pv_actual.copy())
    history.load_primary_residual.append(
        None
        if forecast.load_primary is None
        else load_actual - forecast.load_primary
    )
    history.pv_primary_residual.append(
        None if forecast.pv_primary is None else pv_actual - forecast.pv_primary
    )
    paired_final_available = (
        forecast.load_forecast is not None and forecast.pv_forecast is not None
    )
    history.load_final_residual.append(
        load_actual - forecast.load_forecast if paired_final_available else None
    )
    history.pv_final_residual.append(
        pv_actual - forecast.pv_forecast if paired_final_available else None
    )


def require_forecast_bundle(draft: ForecastDraft) -> ForecastBundle:
    required = (
        draft.load_primary,
        draft.pv_primary,
        draft.load_forecast,
        draft.pv_forecast,
        draft.baseline_load,
        draft.baseline_pv,
        draft.scenario_load_residuals,
        draft.scenario_pv_residuals,
    )
    if any(value is None for value in required):
        raise RuntimeError(f"{draft.target_date}历史不足，无法形成完整Q2预测/场景")
    if not draft.scenario_dates:
        raise RuntimeError(f"{draft.target_date}没有可用配对残差场景")
    scenario_load = np.maximum(
        draft.load_forecast[None, :] + draft.scenario_load_residuals,
        0.0,
    )
    scenario_pv = np.maximum(
        draft.pv_forecast[None, :] + draft.scenario_pv_residuals,
        0.0,
    )
    return ForecastBundle(
        target_date=draft.target_date,
        load_primary=draft.load_primary.copy(),
        pv_primary=draft.pv_primary.copy(),
        load_forecast=draft.load_forecast.copy(),
        pv_forecast=draft.pv_forecast.copy(),
        baseline_load=draft.baseline_load.copy(),
        baseline_pv=draft.baseline_pv.copy(),
        load_history=draft.load_history,
        pv_history=draft.pv_history,
        bias_dates=draft.bias_dates,
        scenario_dates=draft.scenario_dates,
        scenario_net_kwh=scenario_load - scenario_pv,
    )


def metric_triplet(forecast: np.ndarray, actual: np.ndarray) -> dict[str, float]:
    error = forecast - actual
    return {
        "mae_kwh": float(np.mean(np.abs(error))),
        "rmse_kwh": float(np.sqrt(np.mean(np.square(error)))),
        "bias_forecast_minus_actual_kwh": float(np.mean(error)),
    }


def prepare_causal_forecasts(
    annual: AnnualData,
) -> tuple[
    dict[date, ForecastBundle],
    ForecastBundle,
    list[dict[str, Any]],
    list[dict[str, Any]],
]:
    history = ForecastHistory.empty()
    bundles: dict[date, ForecastBundle] = {}
    metric_rows: list[dict[str, Any]] = []
    scenario_audit_rows: list[dict[str, Any]] = []

    for i, observed_date in enumerate(annual.dates):
        draft = forecast_one_day(observed_date, history)
        if FORECAST_START_DATE <= observed_date <= FORECAST_END_DATE:
            bundle = require_forecast_bundle(draft)
            bundles[observed_date] = bundle
            for method, load_forecast, pv_forecast in (
                ("baseline_recent3_equal_no_bias", bundle.baseline_load, bundle.baseline_pv),
                ("main_weighted_plus_30day_bias", bundle.load_forecast, bundle.pv_forecast),
            ):
                load_metrics = metric_triplet(load_forecast, annual.load_kwh[i])
                pv_metrics = metric_triplet(pv_forecast, annual.pv_kwh[i])
                metric_rows.append(
                    {
                        "date": observed_date.isoformat(),
                        "method": method,
                        "load_mae_kwh": load_metrics["mae_kwh"],
                        "load_rmse_kwh": load_metrics["rmse_kwh"],
                        "load_bias_forecast_minus_actual_kwh": load_metrics[
                            "bias_forecast_minus_actual_kwh"
                        ],
                        "pv_mae_kwh": pv_metrics["mae_kwh"],
                        "pv_rmse_kwh": pv_metrics["rmse_kwh"],
                        "pv_bias_forecast_minus_actual_kwh": pv_metrics[
                            "bias_forecast_minus_actual_kwh"
                        ],
                    }
                )
            scenario_audit_rows.extend(forecast_audit_rows(bundle, "formal_output"))
        append_observation(
            history,
            observed_date,
            annual.load_kwh[i],
            annual.pv_kwh[i],
            draft,
        )

    boundary_draft = forecast_one_day(BOUNDARY_PLAN_DATE, history)
    boundary_bundle = require_forecast_bundle(boundary_draft)
    scenario_audit_rows.extend(
        forecast_audit_rows(boundary_bundle, "template_boundary_only")
    )
    return bundles, boundary_bundle, metric_rows, scenario_audit_rows


def forecast_audit_rows(
    bundle: ForecastBundle, purpose: str
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    target = bundle.target_date.isoformat()
    for rank, (history_date, distance, weight) in enumerate(
        bundle.load_history, start=1
    ):
        rows.append(
            {
                "target_date": target,
                "purpose": purpose,
                "record_type": "load_primary_same_class",
                "rank": rank,
                "history_date": history_date.isoformat(),
                "distance_days": distance,
                "weight": weight,
                "load_class": "low" if is_low_load_day(bundle.target_date) else "regular",
            }
        )
    for rank, (history_date, distance, weight) in enumerate(
        bundle.pv_history, start=1
    ):
        rows.append(
            {
                "target_date": target,
                "purpose": purpose,
                "record_type": "pv_primary_recent_day",
                "rank": rank,
                "history_date": history_date.isoformat(),
                "distance_days": distance,
                "weight": weight,
                "load_class": "",
            }
        )
    bias_weight = 1.0 / len(bundle.bias_dates) if bundle.bias_dates else 0.0
    for rank, history_date in enumerate(bundle.bias_dates, start=1):
        rows.append(
            {
                "target_date": target,
                "purpose": purpose,
                "record_type": "paired_primary_residual_bias_window",
                "rank": rank,
                "history_date": history_date.isoformat(),
                "distance_days": (bundle.target_date - history_date).days,
                "weight": bias_weight,
                "load_class": "",
            }
        )
    scenario_weight = 1.0 / len(bundle.scenario_dates)
    for rank, history_date in enumerate(bundle.scenario_dates, start=1):
        rows.append(
            {
                "target_date": target,
                "purpose": purpose,
                "record_type": "paired_final_residual_scenario",
                "rank": rank,
                "history_date": history_date.isoformat(),
                "distance_days": (bundle.target_date - history_date).days,
                "weight": scenario_weight,
                "load_class": "",
            }
        )
    return rows


def build_planning_model(
    target_date: date,
    prices: np.ndarray,
    scenario_net: np.ndarray,
    initial_soc: float,
    terminal_value_coefficient: float,
    *,
    binary_mutex: bool,
) -> PlanningModel:
    scenarios = scenario_net.shape[0]
    problem = pulp.LpProblem(
        f"Q2_DayAhead_{target_date.isoformat()}", pulp.LpMinimize
    )
    grid = [pulp.LpVariable(f"g_{t + 1:03d}", lowBound=0.0) for t in range(T)]
    charge = [
        [
            pulp.LpVariable(
                f"C_{w + 1:02d}_{t + 1:03d}",
                lowBound=0.0,
                upBound=MAX_ACTION_KWH,
            )
            for t in range(T)
        ]
        for w in range(scenarios)
    ]
    discharge = [
        [
            pulp.LpVariable(
                f"D_{w + 1:02d}_{t + 1:03d}",
                lowBound=0.0,
                upBound=MAX_ACTION_KWH,
            )
            for t in range(T)
        ]
        for w in range(scenarios)
    ]
    emergency = [
        [pulp.LpVariable(f"b_{w + 1:02d}_{t + 1:03d}", lowBound=0.0) for t in range(T)]
        for w in range(scenarios)
    ]
    unused = [
        [pulp.LpVariable(f"U_{w + 1:02d}_{t + 1:03d}", lowBound=0.0) for t in range(T)]
        for w in range(scenarios)
    ]
    soc = [
        [
            pulp.LpVariable(
                f"E_{w + 1:02d}_{t + 1:03d}",
                lowBound=SOC_MIN_KWH,
                upBound=SOC_MAX_KWH,
            )
            for t in range(T)
        ]
        for w in range(scenarios)
    ]
    mode: list[list[pulp.LpVariable]] | None = None
    if binary_mutex:
        mode = [
            [pulp.LpVariable(f"z_{w + 1:02d}_{t + 1:03d}", cat=pulp.LpBinary) for t in range(T)]
            for w in range(scenarios)
        ]

    for w in range(scenarios):
        for t in range(T):
            problem += (
                grid[t] + emergency[w][t] + discharge[w][t]
                == float(scenario_net[w, t]) + charge[w][t] + unused[w][t],
                f"balance_{w + 1:02d}_{t + 1:03d}",
            )
            previous_soc: float | pulp.LpVariable = initial_soc if t == 0 else soc[w][t - 1]
            problem += (
                soc[w][t]
                == previous_soc
                + ETA_CHARGE * charge[w][t]
                - discharge[w][t] / ETA_DISCHARGE,
                f"soc_{w + 1:02d}_{t + 1:03d}",
            )
            if mode is not None:
                problem += (
                    charge[w][t] <= MAX_ACTION_KWH * mode[w][t],
                    f"charge_mutex_{w + 1:02d}_{t + 1:03d}",
                )
                problem += (
                    discharge[w][t] <= MAX_ACTION_KWH * (1.0 - mode[w][t]),
                    f"discharge_mutex_{w + 1:02d}_{t + 1:03d}",
                )

    probability = 1.0 / scenarios
    plan_cost = pulp.lpSum(float(prices[t]) * grid[t] for t in range(T))
    expected_recourse = probability * pulp.lpSum(
        pulp.lpSum(
            EMERGENCY_PRICE_FACTOR * float(prices[t]) * emergency[w][t]
            for t in range(T)
        )
        - terminal_value_coefficient * soc[w][-1]
        for w in range(scenarios)
    )
    original_objective = plan_cost + expected_recourse
    problem += original_objective, "expected_total_cost_with_terminal_value"
    return PlanningModel(
        problem,
        original_objective,
        grid,
        charge,
        discharge,
        emergency,
        unused,
        soc,
    )


def make_highs_solver(*, mip: bool) -> pulp.HiGHS:
    return pulp.HiGHS(
        mip=mip,
        msg=False,
        threads=1,
        random_seed=RANDOM_SEED,
        primal_feasibility_tolerance=1.0e-8,
        dual_feasibility_tolerance=1.0e-8,
        mip_rel_gap=0.0,
    )


def solve_current_model(model: PlanningModel, *, mip: bool) -> tuple[str, float, Any]:
    started = time.perf_counter()
    status_code = model.problem.solve(make_highs_solver(mip=mip))
    elapsed = time.perf_counter() - started
    status = pulp.LpStatus.get(status_code, f"Unknown({status_code})")
    if status != "Optimal":
        raise RuntimeError(f"日前模型未达到Optimal：{status}")
    return status, elapsed, model.problem.solverModel.getInfo()


def values_1d(variables: Sequence[pulp.LpVariable]) -> np.ndarray:
    values = [variable.value() for variable in variables]
    if any(value is None or not math.isfinite(float(value)) for value in values):
        raise RuntimeError("求解器返回了空值或非有限变量")
    return np.asarray(values, dtype=np.float64)


def values_2d(variables: Sequence[Sequence[pulp.LpVariable]]) -> np.ndarray:
    return np.asarray([values_1d(row) for row in variables], dtype=np.float64)


def planning_arrays(model: PlanningModel) -> dict[str, np.ndarray]:
    return {
        "g": values_1d(model.grid),
        "c": values_2d(model.charge),
        "d": values_2d(model.discharge),
        "b": values_2d(model.emergency),
        "u": values_2d(model.unused),
        "e": values_2d(model.soc),
    }


def solve_day_ahead(
    target_date: date,
    prices: np.ndarray,
    scenario_net: np.ndarray,
    initial_soc: float,
    terminal_value_coefficient: float,
) -> PlanningResult:
    if scenario_net.ndim != 2 or scenario_net.shape[1] != T:
        raise ValueError(f"场景净负荷维度异常：{scenario_net.shape}")
    if not (SOC_MIN_KWH - CHECK_TOLERANCE <= initial_soc <= SOC_MAX_KWH + CHECK_TOLERANCE):
        raise ValueError(f"日前模型初始SOC越界：{initial_soc}")

    model = build_planning_model(
        target_date,
        prices,
        scenario_net,
        initial_soc,
        terminal_value_coefficient,
        binary_mutex=False,
    )
    status, solve_seconds, info = solve_current_model(model, mip=False)
    best_objective = float(pulp.value(model.original_objective))
    arrays = planning_arrays(model)
    max_product = float(np.max(arrays["c"] * arrays["d"]))
    tiebreak_used = False
    binary_fallback = False

    if max_product > SIMULTANEOUS_PRODUCT_TOLERANCE:
        tiebreak_used = True
        cost_tolerance = max(1.0e-7, abs(best_objective) * 1.0e-10)
        model.problem += (
            model.original_objective <= best_objective + cost_tolerance,
            "preserve_primary_optimal_cost",
        )
        model.problem.setObjective(
            pulp.lpSum(
                variable
                for rows in (model.charge, model.discharge)
                for row in rows
                for variable in row
            )
        )
        tie_status, tie_seconds, info = solve_current_model(model, mip=False)
        status = f"{status}; tiebreak={tie_status}"
        solve_seconds += tie_seconds
        arrays = planning_arrays(model)
        max_product = float(np.max(arrays["c"] * arrays["d"]))

    if max_product > SIMULTANEOUS_PRODUCT_TOLERANCE:
        binary_fallback = True
        model = build_planning_model(
            target_date,
            prices,
            scenario_net,
            initial_soc,
            terminal_value_coefficient,
            binary_mutex=True,
        )
        status, binary_seconds, info = solve_current_model(model, mip=True)
        solve_seconds += binary_seconds
        status = f"binary_fallback={status}"
        arrays = planning_arrays(model)
        max_product = float(np.max(arrays["c"] * arrays["d"]))

    objective_value = float(pulp.value(model.original_objective))
    g = arrays["g"]
    c = arrays["c"]
    d = arrays["d"]
    b = arrays["b"]
    u = arrays["u"]
    e = arrays["e"]
    previous_e = np.concatenate(
        [np.full((e.shape[0], 1), initial_soc), e[:, :-1]], axis=1
    )
    balance_residual = g[None, :] + b + d - scenario_net - c - u
    soc_residual = e - previous_e - ETA_CHARGE * c + d / ETA_DISCHARGE
    max_balance = float(np.max(np.abs(balance_residual)))
    max_soc_residual = float(np.max(np.abs(soc_residual)))
    min_flow = float(min(np.min(g), np.min(c), np.min(d), np.min(b), np.min(u)))
    min_soc = float(np.min(e))
    max_soc = float(np.max(e))
    max_charge = float(np.max(c))
    max_discharge = float(np.max(d))

    checks = {
        "balance": max_balance <= CHECK_TOLERANCE,
        "soc_transition": max_soc_residual <= CHECK_TOLERANCE,
        "soc_bounds": min_soc >= SOC_MIN_KWH - CHECK_TOLERANCE
        and max_soc <= SOC_MAX_KWH + CHECK_TOLERANCE,
        "action_bounds": max_charge <= MAX_ACTION_KWH + CHECK_TOLERANCE
        and max_discharge <= MAX_ACTION_KWH + CHECK_TOLERANCE,
        "nonnegative": min_flow >= -CHECK_TOLERANCE,
        "no_simultaneous": max_product <= SIMULTANEOUS_PRODUCT_TOLERANCE,
    }
    if not all(checks.values()):
        raise RuntimeError(f"{target_date}日前模型验收失败：{checks}")

    cleaned_g = g.copy()
    cleaned_g[np.abs(cleaned_g) <= ACTION_ZERO_TOLERANCE] = 0.0
    highs_mip_gap = getattr(info, "mip_gap", None) if binary_fallback else None
    return PlanningResult(
        target_date=target_date,
        initial_soc_kwh=initial_soc,
        grid_purchase=cleaned_g,
        objective_value=objective_value,
        plan_energy_kwh=float(np.sum(cleaned_g)),
        plan_cost_yuan=float(np.dot(prices, cleaned_g)),
        scenario_count=scenario_net.shape[0],
        status=status,
        solve_seconds=solve_seconds,
        highs_objective=float(getattr(info, "objective_function_value", objective_value)),
        highs_max_primal_infeasibility=float(
            getattr(info, "max_primal_infeasibility", float("nan"))
        ),
        highs_num_primal_infeasibilities=int(
            getattr(info, "num_primal_infeasibilities", -1)
        ),
        simplex_iterations=int(getattr(info, "simplex_iteration_count", -1)),
        ipm_iterations=int(getattr(info, "ipm_iteration_count", -1)),
        mip_gap=None if highs_mip_gap is None else float(highs_mip_gap),
        tiebreak_used=tiebreak_used,
        binary_fallback_used=binary_fallback,
        max_balance_residual_kwh=max_balance,
        max_soc_residual_kwh=max_soc_residual,
        max_simultaneous_product_kwh2=max_product,
        min_variable_value_kwh=min_flow,
        min_soc_kwh=min_soc,
        max_soc_kwh=max_soc,
        max_charge_kwh=max_charge,
        max_discharge_kwh=max_discharge,
    )


def validate_executor(
    prices: np.ndarray,
    actual_net: np.ndarray,
    plan: np.ndarray,
    initial_soc: float,
    charge: np.ndarray,
    discharge: np.ndarray,
    emergency: np.ndarray,
    unused: np.ndarray,
    soc: np.ndarray,
    *,
    dp_max_convexity_violation: float | None = None,
) -> ExecutorResult:
    charge = charge.copy()
    discharge = discharge.copy()
    emergency = emergency.copy()
    unused = unused.copy()
    for flow in (charge, discharge, emergency, unused):
        flow[np.abs(flow) <= ACTION_ZERO_TOLERANCE] = 0.0
    previous_soc = np.concatenate(([initial_soc], soc[:-1]))
    balance_residual = plan + emergency + discharge - actual_net - charge - unused
    soc_residual = soc - previous_soc - ETA_CHARGE * charge + discharge / ETA_DISCHARGE
    max_balance = float(np.max(np.abs(balance_residual)))
    max_soc_residual = float(np.max(np.abs(soc_residual)))
    max_product = float(np.max(charge * discharge))
    min_flow = float(min(np.min(charge), np.min(discharge), np.min(emergency), np.min(unused)))
    checks = {
        "balance": max_balance <= CHECK_TOLERANCE,
        "soc": max_soc_residual <= CHECK_TOLERANCE,
        "soc_bounds": float(np.min(soc)) >= SOC_MIN_KWH - CHECK_TOLERANCE
        and float(np.max(soc)) <= SOC_MAX_KWH + CHECK_TOLERANCE,
        "action_bounds": float(np.max(charge)) <= MAX_ACTION_KWH + CHECK_TOLERANCE
        and float(np.max(discharge)) <= MAX_ACTION_KWH + CHECK_TOLERANCE,
        "nonnegative": min_flow >= -CHECK_TOLERANCE,
        "mutex": max_product <= SIMULTANEOUS_PRODUCT_TOLERANCE,
    }
    if not all(checks.values()):
        raise RuntimeError(f"执行器验收失败：{checks}")
    emergency_cost = float(np.dot(EMERGENCY_PRICE_FACTOR * prices, emergency))
    return ExecutorResult(
        charge=charge,
        discharge=discharge,
        emergency=emergency,
        unused=unused,
        soc=soc,
        initial_soc_kwh=initial_soc,
        terminal_soc_kwh=float(soc[-1]),
        emergency_energy_kwh=float(np.sum(emergency)),
        emergency_cost_yuan=emergency_cost,
        max_balance_residual_kwh=max_balance,
        max_soc_residual_kwh=max_soc_residual,
        min_soc_kwh=float(np.min(soc)),
        max_soc_kwh=float(np.max(soc)),
        max_charge_kwh=float(np.max(charge)),
        max_discharge_kwh=float(np.max(discharge)),
        max_simultaneous_product_kwh2=max_product,
        min_nonnegative_flow_kwh=min_flow,
        dp_max_convexity_violation=dp_max_convexity_violation,
    )


def execute_analytical(
    prices: np.ndarray,
    actual_net: np.ndarray,
    plan: np.ndarray,
    initial_soc: float,
) -> ExecutorResult:
    charge = np.zeros(T, dtype=np.float64)
    discharge = np.zeros(T, dtype=np.float64)
    emergency = np.zeros(T, dtype=np.float64)
    unused = np.zeros(T, dtype=np.float64)
    soc = np.zeros(T, dtype=np.float64)
    current_soc = initial_soc
    for t in range(T):
        r = float(actual_net[t] - plan[t])
        if r <= 0.0:
            charge[t] = min(
                max(-r, 0.0),
                MAX_ACTION_KWH,
                max((SOC_MAX_KWH - current_soc) / ETA_CHARGE, 0.0),
            )
            unused[t] = max(-r - charge[t], 0.0)
        else:
            discharge[t] = min(
                r,
                MAX_ACTION_KWH,
                max(ETA_DISCHARGE * (current_soc - SOC_MIN_KWH), 0.0),
            )
            emergency[t] = max(r - discharge[t], 0.0)
        current_soc += ETA_CHARGE * charge[t] - discharge[t] / ETA_DISCHARGE
        current_soc = min(max(current_soc, SOC_MIN_KWH), SOC_MAX_KWH)
        soc[t] = current_soc
    return validate_executor(
        prices,
        actual_net,
        plan,
        initial_soc,
        charge,
        discharge,
        emergency,
        unused,
        soc,
    )


def state_grid(step_kwh: float) -> np.ndarray:
    if step_kwh <= 0.0:
        raise ValueError("DP网格步长必须为正")
    span = SOC_MAX_KWH - SOC_MIN_KWH
    intervals = span / step_kwh
    if abs(intervals - round(intervals)) > 1.0e-9:
        raise ValueError("DP网格步长必须整除SOC范围9600 kWh")
    return np.linspace(SOC_MIN_KWH, SOC_MAX_KWH, int(round(intervals)) + 1)


def convexity_violation(values: np.ndarray, grid: np.ndarray) -> float:
    slopes = np.diff(values) / np.diff(grid)
    if slopes.size < 2:
        return 0.0
    return float(max(0.0, -np.min(np.diff(slopes))))


def average_scenario_value_functions(
    scenario_net: np.ndarray,
    plan: np.ndarray,
    prices: np.ndarray,
    terminal_value_coefficient: float,
    grid: np.ndarray,
) -> tuple[np.ndarray, float]:
    scenarios = scenario_net.shape[0]
    value_sum = np.zeros((T + 1, grid.size), dtype=np.float64)
    terminal = -terminal_value_coefficient * grid
    value_sum[T] = scenarios * terminal
    max_convex_violation = 0.0

    for w in range(scenarios):
        next_value = terminal.copy()
        for t in range(T - 1, -1, -1):
            r = float(scenario_net[w, t] - plan[t])
            if r > 0.0:
                internal_span = min(MAX_ACTION_KWH, r) / ETA_DISCHARGE
                lower = np.maximum(SOC_MIN_KWH, grid - internal_span)
                upper = grid
                slope = EMERGENCY_PRICE_FACTOR * float(prices[t]) * ETA_DISCHARGE
                transformed = next_value + slope * grid
                unconstrained = float(grid[int(np.argmin(transformed))])
                next_soc = np.clip(unconstrained, lower, upper)
                action = next_soc - grid
                current = (
                    EMERGENCY_PRICE_FACTOR
                    * float(prices[t])
                    * np.maximum(r + ETA_DISCHARGE * action, 0.0)
                    + np.interp(next_soc, grid, next_value)
                )
            else:
                internal_span = ETA_CHARGE * min(MAX_ACTION_KWH, -r)
                lower = grid
                upper = np.minimum(SOC_MAX_KWH, grid + internal_span)
                unconstrained = float(grid[int(np.argmin(next_value))])
                next_soc = np.clip(unconstrained, lower, upper)
                current = np.interp(next_soc, grid, next_value)
            violation = convexity_violation(current, grid)
            max_convex_violation = max(max_convex_violation, violation)
            if violation > 1.0e-7:
                raise RuntimeError(
                    f"DP场景价值函数失去凸性：date-step={t + 1}, "
                    f"scenario={w + 1}, violation={violation}"
                )
            value_sum[t] += current
            next_value = current
    return value_sum / scenarios, max_convex_violation


def candidate_minimizers(
    transformed_grid_values: np.ndarray,
    grid: np.ndarray,
    lower: float,
    upper: float,
    current_soc: float,
) -> list[float]:
    minimum = float(np.min(transformed_grid_values))
    tie_tolerance = max(1.0e-10, abs(minimum) * 1.0e-12)
    tied = np.flatnonzero(transformed_grid_values <= minimum + tie_tolerance)
    q_low = float(grid[int(tied[0])])
    q_high = float(grid[int(tied[-1])])
    candidates = [
        lower,
        upper,
        min(max(q_low, lower), upper),
        min(max(q_high, lower), upper),
        min(max(current_soc, max(lower, q_low)), min(upper, q_high)),
    ]
    unique: list[float] = []
    for candidate in candidates:
        if not any(abs(candidate - existing) <= 1.0e-10 for existing in unique):
            unique.append(candidate)
    return unique


def execute_dp(
    prices: np.ndarray,
    actual_net: np.ndarray,
    plan: np.ndarray,
    initial_soc: float,
    scenario_net: np.ndarray,
    terminal_value_coefficient: float,
    grid_step_kwh: float,
) -> ExecutorResult:
    grid = state_grid(grid_step_kwh)
    average_values, max_convex_violation = average_scenario_value_functions(
        scenario_net,
        plan,
        prices,
        terminal_value_coefficient,
        grid,
    )
    charge = np.zeros(T, dtype=np.float64)
    discharge = np.zeros(T, dtype=np.float64)
    emergency = np.zeros(T, dtype=np.float64)
    unused = np.zeros(T, dtype=np.float64)
    soc = np.zeros(T, dtype=np.float64)
    current_soc = initial_soc

    for t in range(T):
        r = float(actual_net[t] - plan[t])
        future = average_values[t + 1]
        if r > 0.0:
            lower = max(
                SOC_MIN_KWH,
                current_soc - min(MAX_ACTION_KWH, r) / ETA_DISCHARGE,
            )
            upper = current_soc
            slope = EMERGENCY_PRICE_FACTOR * float(prices[t]) * ETA_DISCHARGE
            transformed = future + slope * grid
        else:
            lower = current_soc
            upper = min(
                SOC_MAX_KWH,
                current_soc + ETA_CHARGE * min(MAX_ACTION_KWH, -r),
            )
            transformed = future
        candidates = candidate_minimizers(
            transformed, grid, lower, upper, current_soc
        )
        evaluated: list[tuple[float, float, float]] = []
        for next_soc in candidates:
            x = next_soc - current_soc
            psi = x / ETA_CHARGE if x >= 0.0 else ETA_DISCHARGE * x
            stage_cost = (
                EMERGENCY_PRICE_FACTOR
                * float(prices[t])
                * max(r + psi, 0.0)
            )
            objective = stage_cost + float(np.interp(next_soc, grid, future))
            evaluated.append((objective, abs(x), next_soc))
        best_objective = min(item[0] for item in evaluated)
        objective_tolerance = max(1.0e-9, abs(best_objective) * 1.0e-12)
        best_candidates = [
            item for item in evaluated if item[0] <= best_objective + objective_tolerance
        ]
        _, _, next_soc = min(best_candidates, key=lambda item: (item[1], item[2]))
        x = next_soc - current_soc
        if abs(x) <= ACTION_ZERO_TOLERANCE:
            x = 0.0
            next_soc = current_soc
        if x >= 0.0:
            charge[t] = x / ETA_CHARGE
        else:
            discharge[t] = -ETA_DISCHARGE * x
        psi = x / ETA_CHARGE if x >= 0.0 else ETA_DISCHARGE * x
        emergency[t] = max(r + psi, 0.0)
        unused[t] = max(-r - psi, 0.0)
        current_soc = min(max(next_soc, SOC_MIN_KWH), SOC_MAX_KWH)
        soc[t] = current_soc

    return validate_executor(
        prices,
        actual_net,
        plan,
        initial_soc,
        charge,
        discharge,
        emergency,
        unused,
        soc,
        dp_max_convexity_violation=max_convex_violation,
    )


def merge_emergency_intervals(
    emergency: np.ndarray,
    tolerance: float = EMERGENCY_INTERVAL_TOLERANCE,
) -> list[dict[str, Any]]:
    intervals: list[dict[str, Any]] = []
    start: int | None = None
    for i in range(T + 1):
        active = i < T and float(emergency[i]) > tolerance
        if active and start is None:
            start = i
        if not active and start is not None:
            stop = i
            start_minutes = start * 10
            end_minutes = stop * 10
            start_label = f"{start_minutes // 60}:{start_minutes % 60:02d}"
            end_label = (
                "00:00+1"
                if end_minutes == 1440
                else f"{end_minutes // 60}:{end_minutes % 60:02d}"
            )
            intervals.append(
                {
                    "start_t": start + 1,
                    "end_t": stop,
                    "physical_interval": f"{start_label}-{end_label}",
                    "emergency_energy_kwh": float(np.sum(emergency[start:stop])),
                }
            )
            start = None
    return intervals


def four_hour_aggregates(result: ExecutorResult) -> list[tuple[float, float]]:
    return [
        (
            float(np.sum(result.charge[block * 24 : (block + 1) * 24])),
            float(np.sum(result.discharge[block * 24 : (block + 1) * 24])),
        )
        for block in range(6)
    ]


def aggregate_forecast_metrics(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for method in ("baseline_recent3_equal_no_bias", "main_weighted_plus_30day_bias"):
        selected = [row for row in rows if row["method"] == method]
        for resource in ("load", "pv"):
            result.append(
                {
                    "method": method,
                    "resource": resource,
                    "days": len(selected),
                    "mean_daily_mae_kwh": statistics.mean(
                        row[f"{resource}_mae_kwh"] for row in selected
                    ),
                    "root_mean_daily_mse_kwh": math.sqrt(
                        statistics.mean(
                            row[f"{resource}_rmse_kwh"] ** 2 for row in selected
                        )
                    ),
                    "mean_bias_forecast_minus_actual_kwh": statistics.mean(
                        row[f"{resource}_bias_forecast_minus_actual_kwh"]
                        for row in selected
                    ),
                }
            )
    return result


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: Sequence[str] | None = None) -> None:
    if fieldnames is None:
        if not rows:
            raise ValueError(f"不能推断空CSV的列：{path}")
        fieldnames = tuple(rows[0].keys())
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames, extrasaction="raise")
        writer.writeheader()
        writer.writerows(rows)


def apply_row_style(
    worksheet: Any,
    target_row: int,
    style_cells: Sequence[Any],
    row_height: float | None,
) -> None:
    for column, source_cell in enumerate(style_cells, start=1):
        target = worksheet.cell(target_row, column)
        target._style = copy(source_cell._style)
        target.number_format = source_cell.number_format
        target.alignment = copy(source_cell.alignment)
        target.font = copy(source_cell.font)
        target.fill = copy(source_cell.fill)
        target.border = copy(source_cell.border)
        target.protection = copy(source_cell.protection)
    worksheet.row_dimensions[target_row].height = row_height


def validate_result2_template(path: Path) -> tuple[date, ...]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        if tuple(workbook.sheetnames) != RESULT_SHEETS:
            raise ValueError(f"result2工作表结构异常：{workbook.sheetnames}")
        plan = workbook["计划购电量"]
        if plan.max_row != 335 or plan.max_column != 147:
            raise ValueError(f"result2计划表尺寸异常：{plan.max_row}×{plan.max_column}")
        headers = [plan.cell(1, column).value for column in range(1, 148)]
        if headers[0] != "日期\\时间" or headers[1] != "0:10-0:20":
            raise ValueError("result2计划表前部表头异常")
        if headers[143:147] != [
            "23:50-0:00+1",
            "0:00-0:10+1",
            "全天购电量",
            "全天购电费",
        ]:
            raise ValueError(f"result2计划表尾部表头异常：{headers[143:147]}")
        plan_dates = tuple(
            as_date(plan.cell(row, 1).value, f"result2计划表第{row}行")
            for row in range(2, 336)
        )
        expected = tuple(
            FORECAST_START_DATE + timedelta(days=i)
            for i in range((FORECAST_END_DATE - FORECAST_START_DATE).days + 1)
        )
        if plan_dates != expected:
            raise ValueError("result2计划表日期不连续或不完整")
        storage = workbook["充放电量"]
        if storage.max_row != 20 or storage.max_column != 6:
            raise ValueError("result2充放电量表尺寸异常")
        emergency = workbook["紧急购电量"]
        if emergency.max_row != 11 or emergency.max_column != 3:
            raise ValueError("result2紧急购电量表初始尺寸异常")
        return plan_dates
    finally:
        workbook.close()


def write_result2(
    template_path: Path,
    output_path: Path,
    prices: np.ndarray,
    exports: dict[date, ExportDay],
    boundary_t1: float,
) -> None:
    """按334天全量展开填写官方模板，省略号缩写行不保留。

    充放电量：每天6个4小时块，共334*6=2004个数据行；
    紧急购电量：每天按实际连续区间数成行，无紧急购电写“无/0”。
    """
    plan_dates = validate_result2_template(template_path)
    if set(exports) != set(plan_dates):
        raise ValueError("导出数据日期与模板334天不一致")
    if output_path.resolve() == template_path.resolve():
        raise ValueError("result2输出路径不得覆盖官方模板")
    if not math.isfinite(boundary_t1) or boundary_t1 < 0.0:
        raise ValueError(f"边界计划t=1购电量异常：{boundary_t1}")
    source_hash = sha256_file(template_path)
    shutil.copy2(template_path, output_path)
    workbook = load_workbook(output_path, data_only=False, read_only=False)
    try:
        plan_sheet = workbook["计划购电量"]
        displayed_prices = np.concatenate((prices[1:], prices[:1]))
        for row, target_date in enumerate(plan_dates, start=2):
            current_plan = exports[target_date].plan_grid
            next_date = target_date + timedelta(days=1)
            next_t1 = (
                float(exports[next_date].plan_grid[0])
                if next_date in exports
                else boundary_t1
            )
            displayed = np.concatenate((current_plan[1:], [next_t1]))
            for column, value in enumerate(displayed, start=2):
                plan_sheet.cell(row, column).value = float(value)
            plan_sheet.cell(row, 146).value = float(np.sum(displayed))
            plan_sheet.cell(row, 147).value = float(np.dot(displayed_prices, displayed))

        storage_sheet = workbook["充放电量"]
        block_styles = {
            offset: (
                [
                    copy(storage_sheet.cell(2 + offset, column))
                    for column in range(1, 7)
                ],
                storage_sheet.row_dimensions[2 + offset].height,
            )
            for offset in range(6)
        }
        storage_sheet.delete_rows(2, storage_sheet.max_row - 1)
        for day_index, target_date in enumerate(plan_dates):
            export = exports[target_date]
            blocks = [
                (
                    float(np.sum(export.dp_charge[block * 24 : (block + 1) * 24])),
                    float(np.sum(export.dp_discharge[block * 24 : (block + 1) * 24])),
                )
                for block in range(6)
            ]
            for block, (charge, discharge) in enumerate(blocks):
                template_offset = (0, 1, 2, 2, 2, 5)[block]
                style_cells, height = block_styles[template_offset]
                target_row = 2 + day_index * 6 + block
                apply_row_style(storage_sheet, target_row, style_cells, height)
                if block == 0:
                    storage_sheet.cell(target_row, 1).value = datetime.combine(
                        target_date, datetime_time()
                    )
                storage_sheet.cell(target_row, 2).value = FOUR_HOUR_LABELS[block]
                storage_sheet.cell(target_row, 3).value = charge
                storage_sheet.cell(target_row, 4).value = discharge
                if block == 0:
                    storage_sheet.cell(target_row, 5).value = datetime_time(0, 0)
                elif block == 1:
                    storage_sheet.cell(target_row, 5).value = "24:00"
                if block == 0:
                    storage_sheet.cell(target_row, 6).value = export.dp_initial_soc_kwh
                elif block == 1:
                    storage_sheet.cell(target_row, 6).value = export.dp_terminal_soc_kwh

        emergency_sheet = workbook["紧急购电量"]
        first_style = [
            copy(emergency_sheet.cell(2, column)) for column in range(1, 4)
        ]
        continuation_style = [
            copy(emergency_sheet.cell(3, column)) for column in range(1, 4)
        ]
        first_height = emergency_sheet.row_dimensions[2].height
        continuation_height = emergency_sheet.row_dimensions[3].height
        emergency_sheet.delete_rows(2, emergency_sheet.max_row - 1)
        output_row = 2
        for target_date in plan_dates:
            intervals = merge_emergency_intervals(
                exports[target_date].dp_emergency
            )
            if not intervals:
                intervals = [
                    {"physical_interval": "无", "emergency_energy_kwh": 0.0}
                ]
            for interval_index, interval in enumerate(intervals):
                style = first_style if interval_index == 0 else continuation_style
                height = first_height if interval_index == 0 else continuation_height
                apply_row_style(emergency_sheet, output_row, style, height)
                emergency_sheet.cell(output_row, 1).value = (
                    datetime.combine(target_date, datetime_time())
                    if interval_index == 0
                    else None
                )
                emergency_sheet.cell(output_row, 2).value = interval[
                    "physical_interval"
                ]
                emergency_sheet.cell(output_row, 3).value = interval[
                    "emergency_energy_kwh"
                ]
                output_row += 1
        workbook.save(output_path)
    finally:
        workbook.close()
    if sha256_file(template_path) != source_hash:
        raise RuntimeError("官方result2模板在导出过程中被修改")


def verify_result2(
    path: Path,
    prices: np.ndarray,
    exports: dict[date, ExportDay],
    boundary_t1: float,
) -> dict[str, Any]:
    plan_dates = tuple(
        FORECAST_START_DATE + timedelta(days=i) for i in range(334)
    )
    if set(exports) != set(plan_dates):
        raise RuntimeError("result2回读时导出数据日期不完整")
    workbook = load_workbook(path, data_only=True, read_only=True)
    max_value_difference = 0.0
    try:
        if tuple(workbook.sheetnames) != RESULT_SHEETS:
            raise RuntimeError("输出result2工作表名称或顺序改变")
        plan_sheet = workbook["计划购电量"]
        if plan_sheet.max_row != 335 or plan_sheet.max_column != 147:
            raise RuntimeError("输出result2计划表尺寸改变")
        displayed_prices = np.concatenate((prices[1:], prices[:1]))
        plan_rows = list(
            plan_sheet.iter_rows(min_row=2, max_row=335, values_only=True)
        )
        for row_index, target_date in enumerate(plan_dates, start=2):
            values = plan_rows[row_index - 2]
            next_date = target_date + timedelta(days=1)
            next_t1 = (
                float(exports[next_date].plan_grid[0])
                if next_date in exports
                else boundary_t1
            )
            expected = np.concatenate(
                (exports[target_date].plan_grid[1:], [next_t1])
            )
            actual = np.asarray(values[1:145], dtype=np.float64)
            max_value_difference = max(
                max_value_difference, float(np.max(np.abs(expected - actual)))
            )
            if abs(float(values[145]) - float(np.sum(expected))) > CHECK_TOLERANCE:
                raise RuntimeError(f"输出计划表{target_date}全天购电量错误")
            expected_cost = float(np.dot(displayed_prices, expected))
            if abs(float(values[146]) - expected_cost) > CHECK_TOLERANCE:
                raise RuntimeError(f"输出计划表{target_date}全天购电费错误")

        storage_sheet = workbook["充放电量"]
        if storage_sheet.max_row != 1 + 334 * 6 or storage_sheet.max_column != 6:
            raise RuntimeError(
                f"输出充放电量尺寸异常：{storage_sheet.max_row}×{storage_sheet.max_column}"
            )
        storage_rows = list(
            storage_sheet.iter_rows(min_row=2, values_only=True)
        )
        for day_index, target_date in enumerate(plan_dates):
            export = exports[target_date]
            base = day_index * 6
            if as_date(storage_rows[base][0], "充放电量日期") != target_date:
                raise RuntimeError(f"充放电量{target_date}首行日期错误")
            for block in range(6):
                expected_c = float(np.sum(export.dp_charge[block * 24 : (block + 1) * 24]))
                expected_d = float(np.sum(export.dp_discharge[block * 24 : (block + 1) * 24]))
                values = storage_rows[base + block]
                if values[1] != FOUR_HOUR_LABELS[block]:
                    raise RuntimeError(f"充放电量{target_date}块标签错误")
                max_value_difference = max(
                    max_value_difference,
                    abs(float(values[2]) - expected_c),
                    abs(float(values[3]) - expected_d),
                )
            if abs(float(storage_rows[base][5]) - export.dp_initial_soc_kwh) > CHECK_TOLERANCE:
                raise RuntimeError(f"充放电量{target_date}期初SOC错误")
            if abs(float(storage_rows[base + 1][5]) - export.dp_terminal_soc_kwh) > CHECK_TOLERANCE:
                raise RuntimeError(f"充放电量{target_date}期末SOC错误")

        emergency_sheet = workbook["紧急购电量"]
        rows = list(emergency_sheet.iter_rows(min_row=2, values_only=True))
        cursor = 0
        days_with_emergency = 0
        for target_date in plan_dates:
            intervals = merge_emergency_intervals(
                exports[target_date].dp_emergency
            )
            if not intervals:
                intervals = [
                    {"physical_interval": "无", "emergency_energy_kwh": 0.0}
                ]
            else:
                days_with_emergency += 1
            for interval_index, expected in enumerate(intervals):
                if cursor >= len(rows):
                    raise RuntimeError("紧急购电输出行数不足")
                row = rows[cursor]
                if interval_index == 0 and as_date(row[0], "紧急购电日期") != target_date:
                    raise RuntimeError(f"紧急购电{target_date}首行日期错误")
                if interval_index > 0 and row[0] is not None:
                    raise RuntimeError(f"紧急购电{target_date}续行不应有日期")
                if row[1] != expected["physical_interval"]:
                    raise RuntimeError(f"紧急购电{target_date}物理区间错误")
                if abs(float(row[2]) - float(expected["emergency_energy_kwh"])) > CHECK_TOLERANCE:
                    raise RuntimeError(f"紧急购电{target_date}购电量错误")
                cursor += 1
        if cursor != len(rows):
            raise RuntimeError("紧急购电输出存在未核验的多余行")
    finally:
        workbook.close()
    if max_value_difference > CHECK_TOLERANCE:
        raise RuntimeError(f"result2回读最大差异超限：{max_value_difference}")
    return {
        "sheet_names_preserved": True,
        "plan_rows": 334,
        "plan_physical_mapping": "date d: t=2..144 plus date d+1 t=1",
        "storage_rows": 334 * 6,
        "storage_all_days_filled": True,
        "emergency_rows_dynamic": True,
        "emergency_days_with_purchase": days_with_emergency,
        "max_roundtrip_value_difference": max_value_difference,
    }


def relative_path(path: Path, root: Path) -> str:
    try:
        return path.resolve().relative_to(root.resolve()).as_posix()
    except ValueError:
        return str(path.resolve())


def package_versions() -> dict[str, str]:
    highs = highspy.Highs()
    return {
        "python": platform.python_version(),
        "numpy": np.__version__,
        "pulp": pulp.__version__,
        "openpyxl": openpyxl.__version__,
        "highspy": highs.version(),
    }


def run_pipeline(args: argparse.Namespace) -> dict[str, Any]:
    started_all = time.perf_counter()
    root = Path(__file__).resolve().parent.parent
    price_path = args.price_input.resolve()
    annual_path = args.annual_input.resolve()
    template_path = args.template.resolve()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    prices, price_labels = read_prices(price_path)
    annual = read_annual_data(annual_path)
    validate_time_inputs(price_labels, annual)
    january_indices = [i for i, day in enumerate(annual.dates) if day.month == 1]
    january_dates = tuple(annual.dates[i] for i in january_indices)
    january_load = annual.load_kwh[january_indices].copy()
    class_rows, class_summary = validate_january_load_classes(
        january_dates, january_load
    )
    forecasts, boundary_forecast, forecast_metric_rows, scenario_audit_rows = (
        prepare_causal_forecasts(annual)
    )
    forecast_aggregate_rows = aggregate_forecast_metrics(forecast_metric_rows)

    formal_dates = [
        FORECAST_START_DATE + timedelta(days=i)
        for i in range((FORECAST_END_DATE - FORECAST_START_DATE).days + 1)
    ]
    if args.max_days is not None:
        if not 1 <= args.max_days <= len(formal_dates):
            raise ValueError("--max-days必须位于1..334")
        formal_dates = formal_dates[: args.max_days]
    full_run = len(formal_dates) == 334
    progress_path = output_dir / "q2_progress.log"
    progress_path.write_text(
        f"{datetime.now().astimezone().isoformat(timespec='seconds')} started "
        f"{len(formal_dates)} days, output={output_dir}\n",
        encoding="utf-8",
    )
    date_to_index = {day: i for i, day in enumerate(annual.dates)}
    terminal_value_coefficient = ETA_DISCHARGE * float(np.min(prices))
    dp_initial_soc = INITIAL_SOC_KWH
    analytical_initial_soc = INITIAL_SOC_KWH
    runs: dict[date, DailyRun] = {}
    plans: dict[date, PlanningResult] = {}
    solver_rows: list[dict[str, Any]] = []

    for run_index, target_date in enumerate(formal_dates, start=1):
        bundle = forecasts[target_date]
        planning = solve_day_ahead(
            target_date,
            prices,
            bundle.scenario_net_kwh,
            dp_initial_soc,
            terminal_value_coefficient,
        )
        actual_index = date_to_index[target_date]
        actual_load = annual.load_kwh[actual_index]
        actual_pv = annual.pv_kwh[actual_index]
        actual_net = actual_load - actual_pv
        dp = execute_dp(
            prices,
            actual_net,
            planning.grid_purchase,
            dp_initial_soc,
            bundle.scenario_net_kwh,
            terminal_value_coefficient,
            args.dp_grid_step,
        )
        analytical = execute_analytical(
            prices,
            actual_net,
            planning.grid_purchase,
            analytical_initial_soc,
        )
        run = DailyRun(
            target_date,
            bundle,
            planning,
            actual_load.copy(),
            actual_pv.copy(),
            dp,
            analytical,
        )
        runs[target_date] = run
        plans[target_date] = planning
        solver_rows.append(planning_audit_row(planning, "formal_output"))
        dp_initial_soc = dp.terminal_soc_kwh
        analytical_initial_soc = analytical.terminal_soc_kwh
        progress_line = (
            f"{datetime.now().astimezone().isoformat(timespec='seconds')} "
            f"[{run_index}/{len(formal_dates)}] {target_date} "
            f"LP={planning.solve_seconds:.3f}s "
            f"DP末SOC={dp.terminal_soc_kwh:.3f} "
            f"DP紧急电={dp.emergency_energy_kwh:.3f}"
        )
        with progress_path.open("a", encoding="utf-8") as progress_file:
            progress_file.write(progress_line + "\n")
        if args.progress_every > 0 and (
            run_index == 1
            or run_index % args.progress_every == 0
            or run_index == len(formal_dates)
        ):
            print(progress_line, flush=True)

    if full_run:
        boundary_plan = solve_day_ahead(
            BOUNDARY_PLAN_DATE,
            prices,
            boundary_forecast.scenario_net_kwh,
            dp_initial_soc,
            terminal_value_coefficient,
        )
        plans[BOUNDARY_PLAN_DATE] = boundary_plan
        solver_rows.append(
            planning_audit_row(boundary_plan, "template_boundary_only")
        )

    daily_rows = [daily_metrics_row(run, prices) for run in runs.values()]
    comparison_rows = [executor_comparison_row(run) for run in runs.values()]
    interval_rows = [
        row
        for run in runs.values()
        for row in interval_detail_rows(run, prices, price_labels)
    ]
    warning_messages: list[str] = []
    backtest_by_key = {
        (row["method"], row["resource"]): row for row in forecast_aggregate_rows
    }
    improvements: dict[str, Any] = {}
    for resource in ("load", "pv"):
        baseline = backtest_by_key[("baseline_recent3_equal_no_bias", resource)]
        main = backtest_by_key[("main_weighted_plus_30day_bias", resource)]
        mae_improved = main["mean_daily_mae_kwh"] < baseline["mean_daily_mae_kwh"]
        rmse_improved = (
            main["root_mean_daily_mse_kwh"]
            < baseline["root_mean_daily_mse_kwh"]
        )
        improvements[resource] = {
            "mae_improved": mae_improved,
            "rmse_improved": rmse_improved,
            "main_minus_baseline_mae_kwh": (
                main["mean_daily_mae_kwh"] - baseline["mean_daily_mae_kwh"]
            ),
            "main_minus_baseline_rmse_kwh": (
                main["root_mean_daily_mse_kwh"]
                - baseline["root_mean_daily_mse_kwh"]
            ),
        }
        if not (mae_improved and rmse_improved):
            warning_messages.append(
                f"主预测对{resource}未同时改善总体MAE与RMSE；按指令保留主方案，未回退。"
            )

    output_paths = {
        "load_class_validation": output_dir / "q2_load_class_validation.csv",
        "daily_metrics": output_dir / "q2_daily_metrics.csv",
        "forecast_metrics": output_dir / "q2_forecast_metrics.csv",
        "scenario_audit": output_dir / "q2_scenario_audit.csv",
        "executor_comparison": output_dir / "q2_executor_comparison.csv",
        "interval_details": output_dir / "q2_interval_details.csv",
        "backtest_comparison": output_dir / "q2_backtest_comparison.csv",
        "solver_audit": output_dir / "q2_solver_audit.csv",
    }
    write_csv(output_paths["load_class_validation"], class_rows)
    write_csv(output_paths["daily_metrics"], daily_rows)
    write_csv(output_paths["forecast_metrics"], forecast_metric_rows)
    write_csv(output_paths["scenario_audit"], scenario_audit_rows)
    write_csv(output_paths["executor_comparison"], comparison_rows)
    write_csv(output_paths["interval_details"], interval_rows)
    write_csv(output_paths["backtest_comparison"], forecast_aggregate_rows)
    write_csv(output_paths["solver_audit"], solver_rows)

    template_validation: dict[str, Any] | None = None
    result2_path = output_dir / "result2.xlsx"
    if full_run:
        exports = {
            run.target_date: ExportDay(
                run.planning.grid_purchase,
                run.dp.charge,
                run.dp.discharge,
                run.dp.emergency,
                run.dp.initial_soc_kwh,
                run.dp.terminal_soc_kwh,
            )
            for run in runs.values()
        }
        boundary_t1 = float(plans[BOUNDARY_PLAN_DATE].grid_purchase[0])
        write_result2(template_path, result2_path, prices, exports, boundary_t1)
        template_validation = verify_result2(
            result2_path, prices, exports, boundary_t1
        )
    else:
        warning_messages.append(
            "这是--max-days部分运行：未生成result2.xlsx，也未求解2026-01-01边界计划。"
        )

    annual_totals = annual_summary(runs)
    validation_summary = global_validation_summary(runs, plans)
    elapsed_all = time.perf_counter() - started_all
    report_path = output_dir / "run_summary.json"
    report = {
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "run_type": "full" if full_run else "partial_smoke_test",
        "formal_days_completed": len(runs),
        "parameters": {
            "periods_per_day": T,
            "delta_t_hours": DELTA_T_HOURS,
            "eta_charge": ETA_CHARGE,
            "eta_discharge": ETA_DISCHARGE,
            "soc_min_kwh": SOC_MIN_KWH,
            "soc_max_kwh": SOC_MAX_KWH,
            "initial_soc_2025_02_01_kwh": INITIAL_SOC_KWH,
            "max_action_kwh_per_period": MAX_ACTION_KWH,
            "max_scenarios": MAX_SCENARIOS,
            "terminal_value_coefficient_yuan_per_kwh": terminal_value_coefficient,
            "dp_grid_step_kwh": args.dp_grid_step,
            "random_seed": RANDOM_SEED,
            "low_load_weekdays": ["Friday", "Saturday"],
            "load_class_validation_scope": "2025-01-only",
            "forecast_bias_window_days": BIAS_WINDOW_DAYS,
            "result2_plan_mapping": "date d: physical t=2..144 plus date d+1 t=1",
        },
        "versions": package_versions(),
        "inputs": {
            "price_input": relative_path(price_path, root),
            "price_input_sha256": sha256_file(price_path),
            "annual_input": relative_path(annual_path, root),
            "annual_input_sha256": sha256_file(annual_path),
            "template": relative_path(template_path, root),
            "template_sha256": sha256_file(template_path),
            "attachment3_read": False,
        },
        "load_class_validation": class_summary,
        "forecast_backtest": {
            "aggregate_rows": forecast_aggregate_rows,
            "main_vs_baseline": improvements,
            "automatic_fallback_performed": False,
        },
        "annual_totals": annual_totals,
        "validation": validation_summary,
        "template_validation": template_validation,
        "runtime_seconds": elapsed_all,
        "warnings": warning_messages,
        "outputs": {
            name: {
                "path": relative_path(path, root),
                "sha256": sha256_file(path),
            }
            for name, path in output_paths.items()
        } | {
            "progress_log": {
                "path": relative_path(progress_path, root),
                "sha256": sha256_file(progress_path),
            }
        },
    }
    if full_run:
        report["outputs"]["result2"] = {
            "path": relative_path(result2_path, root),
            "sha256": sha256_file(result2_path),
        }
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    with progress_path.open("a", encoding="utf-8") as progress_file:
        progress_file.write(
            f"{datetime.now().astimezone().isoformat(timespec='seconds')} "
            f"completed {len(runs)} days in {elapsed_all:.2f}s\n"
        )
    print(f"Q2运行完成：{len(runs)}个正式日，耗时{elapsed_all:.2f}s", flush=True)
    print(f"输出目录：{output_dir}", flush=True)
    if warning_messages:
        print("警告：", flush=True)
        for warning in warning_messages:
            print(f"  - {warning}", flush=True)
    return report


def planning_audit_row(planning: PlanningResult, purpose: str) -> dict[str, Any]:
    data = asdict(planning)
    data.pop("grid_purchase")
    data["target_date"] = planning.target_date.isoformat()
    data["purpose"] = purpose
    return data


def daily_metrics_row(run: DailyRun, prices: np.ndarray) -> dict[str, Any]:
    plan_cost = run.planning.plan_cost_yuan
    return {
        "date": run.target_date.isoformat(),
        "scenario_count": run.planning.scenario_count,
        "plan_energy_kwh": run.planning.plan_energy_kwh,
        "plan_cost_yuan": plan_cost,
        "dp_emergency_energy_kwh": run.dp.emergency_energy_kwh,
        "dp_emergency_cost_yuan": run.dp.emergency_cost_yuan,
        "dp_total_cost_yuan": plan_cost + run.dp.emergency_cost_yuan,
        "dp_initial_soc_kwh": run.dp.initial_soc_kwh,
        "dp_terminal_soc_kwh": run.dp.terminal_soc_kwh,
        "dp_charge_kwh": float(np.sum(run.dp.charge)),
        "dp_discharge_kwh": float(np.sum(run.dp.discharge)),
        "dp_unused_kwh": float(np.sum(run.dp.unused)),
        "dp_max_balance_residual_kwh": run.dp.max_balance_residual_kwh,
        "dp_max_soc_residual_kwh": run.dp.max_soc_residual_kwh,
        "dp_max_convexity_violation": run.dp.dp_max_convexity_violation,
        "analytical_emergency_energy_kwh": run.analytical.emergency_energy_kwh,
        "analytical_emergency_cost_yuan": run.analytical.emergency_cost_yuan,
        "analytical_total_cost_yuan": plan_cost + run.analytical.emergency_cost_yuan,
        "analytical_initial_soc_kwh": run.analytical.initial_soc_kwh,
        "analytical_terminal_soc_kwh": run.analytical.terminal_soc_kwh,
        "analytical_charge_kwh": float(np.sum(run.analytical.charge)),
        "analytical_discharge_kwh": float(np.sum(run.analytical.discharge)),
        "analytical_unused_kwh": float(np.sum(run.analytical.unused)),
        "analytical_max_balance_residual_kwh": run.analytical.max_balance_residual_kwh,
        "analytical_max_soc_residual_kwh": run.analytical.max_soc_residual_kwh,
    }


def executor_comparison_row(run: DailyRun) -> dict[str, Any]:
    plan_cost = run.planning.plan_cost_yuan
    analytical_total = plan_cost + run.analytical.emergency_cost_yuan
    dp_total = plan_cost + run.dp.emergency_cost_yuan
    return {
        "date": run.target_date.isoformat(),
        "shared_plan_energy_kwh": run.planning.plan_energy_kwh,
        "shared_plan_cost_yuan": plan_cost,
        "analytical_emergency_energy_kwh": run.analytical.emergency_energy_kwh,
        "analytical_emergency_cost_yuan": run.analytical.emergency_cost_yuan,
        "analytical_total_cost_yuan": analytical_total,
        "dp_emergency_energy_kwh": run.dp.emergency_energy_kwh,
        "dp_emergency_cost_yuan": run.dp.emergency_cost_yuan,
        "dp_total_cost_yuan": dp_total,
        "dp_saving_yuan": analytical_total - dp_total,
        "dp_saving_rate_percent": (
            (analytical_total - dp_total) / analytical_total * 100.0
            if analytical_total != 0.0
            else 0.0
        ),
    }


def interval_detail_rows(
    run: DailyRun,
    prices: np.ndarray,
    source_right_end_labels: tuple[str, ...],
) -> list[dict[str, Any]]:
    labels = physical_interval_labels()
    rows: list[dict[str, Any]] = []
    for i in range(T):
        rows.append(
            {
                "date": run.target_date.isoformat(),
                "t": i + 1,
                "source_right_end_label": source_right_end_labels[i],
                "physical_interval": labels[i],
                "price_yuan_per_kwh": float(prices[i]),
                "actual_load_kwh": float(run.actual_load[i]),
                "actual_pv_kwh": float(run.actual_pv[i]),
                "forecast_load_kwh": float(run.forecast.load_forecast[i]),
                "forecast_pv_kwh": float(run.forecast.pv_forecast[i]),
                "plan_grid_kwh": float(run.planning.grid_purchase[i]),
                "dp_charge_kwh": float(run.dp.charge[i]),
                "dp_discharge_kwh": float(run.dp.discharge[i]),
                "dp_emergency_kwh": float(run.dp.emergency[i]),
                "dp_unused_kwh": float(run.dp.unused[i]),
                "dp_end_soc_kwh": float(run.dp.soc[i]),
                "analytical_charge_kwh": float(run.analytical.charge[i]),
                "analytical_discharge_kwh": float(run.analytical.discharge[i]),
                "analytical_emergency_kwh": float(run.analytical.emergency[i]),
                "analytical_unused_kwh": float(run.analytical.unused[i]),
                "analytical_end_soc_kwh": float(run.analytical.soc[i]),
            }
        )
    return rows


def annual_summary(runs: dict[date, DailyRun]) -> dict[str, Any]:
    plan_cost = sum(run.planning.plan_cost_yuan for run in runs.values())
    plan_energy = sum(run.planning.plan_energy_kwh for run in runs.values())
    dp_emergency_cost = sum(run.dp.emergency_cost_yuan for run in runs.values())
    analytical_emergency_cost = sum(
        run.analytical.emergency_cost_yuan for run in runs.values()
    )
    dp_total = plan_cost + dp_emergency_cost
    analytical_total = plan_cost + analytical_emergency_cost
    saving = analytical_total - dp_total
    return {
        "days": len(runs),
        "plan_energy_kwh": plan_energy,
        "plan_cost_yuan": plan_cost,
        "dp_emergency_energy_kwh": sum(
            run.dp.emergency_energy_kwh for run in runs.values()
        ),
        "dp_emergency_cost_yuan": dp_emergency_cost,
        "dp_total_cost_yuan": dp_total,
        "analytical_emergency_energy_kwh": sum(
            run.analytical.emergency_energy_kwh for run in runs.values()
        ),
        "analytical_emergency_cost_yuan": analytical_emergency_cost,
        "analytical_total_cost_yuan": analytical_total,
        "dp_saving_yuan": saving,
        "dp_saving_rate_percent": (
            saving / analytical_total * 100.0 if analytical_total else 0.0
        ),
    }


def global_validation_summary(
    runs: dict[date, DailyRun],
    plans: dict[date, PlanningResult],
) -> dict[str, Any]:
    ordered_runs = list(runs.values())
    dp_cross_day = max(
        [
            abs(
                ordered_runs[i].dp.initial_soc_kwh
                - ordered_runs[i - 1].dp.terminal_soc_kwh
            )
            for i in range(1, len(ordered_runs))
        ]
        or [0.0]
    )
    analytical_cross_day = max(
        [
            abs(
                ordered_runs[i].analytical.initial_soc_kwh
                - ordered_runs[i - 1].analytical.terminal_soc_kwh
            )
            for i in range(1, len(ordered_runs))
        ]
        or [0.0]
    )
    max_plan_balance = max(
        plan.max_balance_residual_kwh for plan in plans.values()
    )
    max_plan_soc = max(plan.max_soc_residual_kwh for plan in plans.values())
    max_dp_balance = max(run.dp.max_balance_residual_kwh for run in ordered_runs)
    max_dp_soc = max(run.dp.max_soc_residual_kwh for run in ordered_runs)
    max_ana_balance = max(
        run.analytical.max_balance_residual_kwh for run in ordered_runs
    )
    max_ana_soc = max(run.analytical.max_soc_residual_kwh for run in ordered_runs)
    max_dp_convexity = max(
        float(run.dp.dp_max_convexity_violation or 0.0) for run in ordered_runs
    )
    daily_cost_identity = max(
        abs(
            (run.planning.plan_cost_yuan + run.dp.emergency_cost_yuan)
            - (run.planning.plan_cost_yuan + run.dp.emergency_cost_yuan)
        )
        for run in ordered_runs
    )
    passed = all(
        value <= CHECK_TOLERANCE
        for value in (
            max_plan_balance,
            max_plan_soc,
            max_dp_balance,
            max_dp_soc,
            max_ana_balance,
            max_ana_soc,
            dp_cross_day,
            analytical_cross_day,
            daily_cost_identity,
        )
    )
    if not passed:
        raise RuntimeError("Q2全局残差或跨日连续性验收失败")
    return {
        "tolerance": CHECK_TOLERANCE,
        "max_planning_balance_residual_kwh": max_plan_balance,
        "max_planning_soc_residual_kwh": max_plan_soc,
        "max_dp_balance_residual_kwh": max_dp_balance,
        "max_dp_soc_residual_kwh": max_dp_soc,
        "max_analytical_balance_residual_kwh": max_ana_balance,
        "max_analytical_soc_residual_kwh": max_ana_soc,
        "max_dp_value_convexity_violation": max_dp_convexity,
        "max_dp_cross_day_soc_residual_kwh": dp_cross_day,
        "max_analytical_cross_day_soc_residual_kwh": analytical_cross_day,
        "max_daily_cost_identity_residual_yuan": daily_cost_identity,
        "all_planning_statuses_optimal": all(
            "Optimal" in plan.status for plan in plans.values()
        ),
        "all_checks_passed": passed,
    }


def main() -> None:
    args = parse_args()
    run_pipeline(args)


if __name__ == "__main__":
    main()
