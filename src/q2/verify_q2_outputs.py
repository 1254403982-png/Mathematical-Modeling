"""独立复核 output/q2 的全部Q2输出（不依赖 q2.py 内部函数）。

用法：python scr/verify_q2_outputs.py [--output-dir output/q2]
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook

TOL = 1.0e-6
TEMPLATE_SHA256 = "1c26494cfc6d754e0bd9bff7e13e1126a73d2d2da6c5336eb251d89b9a1a1a47"
INPUT1_SHA256 = "66b87134f5ecccd68184d3539bb1293ef039f9e0fdd955a589b9bfa7f227c377"
INPUT2_SHA256 = "2e95fd446bfafa0d8c59577b5c2e2ea8b3f1def20dde54a3062556f4da9b4c72"

results: list[tuple[str, bool, str]] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    results.append((name, ok, detail))
    print(("PASS " if ok else "FAIL ") + name + ("  " + detail if detail else ""))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def f(value: Any) -> float:
    return float(value)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as file:
        return list(csv.DictReader(file))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=None)
    args = parser.parse_args()
    root = Path(__file__).resolve().parent.parent
    out = (args.output_dir or root / "output" / "q2").resolve()

    # ---------- 1 输入与模板哈希 ----------
    check("附件1哈希未变", sha256_file(root / "data/附件/附件1.xlsx") == INPUT1_SHA256)
    check("附件2哈希未变", sha256_file(root / "data/附件/附件2.xlsx") == INPUT2_SHA256)
    check(
        "官方result2模板哈希未变",
        sha256_file(root / "data/附件/附件5/result2.xlsx") == TEMPLATE_SHA256,
    )

    # ---------- 2 电价与明细 ----------
    wb = load_workbook(root / "data/附件/附件1.xlsx", data_only=True, read_only=True)
    ws = wb.active
    prices = [float(ws.cell(row, 2).value) for row in range(2, 146)]
    price_labels = [
        ws.cell(row, 1).value.strftime("%H:%M")
        if hasattr(ws.cell(row, 1).value, "strftime")
        else str(ws.cell(row, 1).value)
        for row in range(2, 146)
    ]
    wb.close()
    shifted_prices = np.concatenate((prices[1:], prices[:1]))
    check("附件1电价144个", len(prices) == 144)
    check(
        "凌晨窗口严格止于05:00且不含05:10",
        price_labels[29] == "05:00" and price_labels[30] == "05:10",
    )
    terminal_price_mean = float(np.mean(prices[:30]))
    terminal_value_coefficient = terminal_price_mean / 0.9
    check(
        f"凌晨00:00-05:00均价={terminal_price_mean:.12f}",
        np.isclose(
            terminal_price_mean,
            0.43339333333333335,
            atol=1.0e-12,
        ),
    )
    check(
        f"充电边际终值系数={terminal_value_coefficient:.12f}",
        np.isclose(
            terminal_value_coefficient,
            0.4815481481481482,
            atol=1.0e-12,
        )
        and not np.isclose(terminal_value_coefficient, 0.33417, atol=1.0e-8),
    )

    summary = json.loads((out / "run_summary.json").read_text(encoding="utf-8"))
    parameters = summary.get("parameters", {})
    check(
        "run_summary终值窗口与效率口径",
        parameters.get("terminal_price_window") == "physical 00:00-05:00"
        and parameters.get("terminal_price_slot_start") == 0
        and parameters.get("terminal_price_slot_end_exclusive") == 30
        and parameters.get("terminal_efficiency_basis") == "charge",
    )
    check(
        "run_summary凌晨均价与终值系数",
        np.isclose(
            f(parameters.get("terminal_price_mean_yuan_per_kwh", float("nan"))),
            terminal_price_mean,
            atol=1.0e-12,
        )
        and np.isclose(
            f(
                parameters.get(
                    "terminal_value_coefficient_yuan_per_kwh", float("nan")
                )
            ),
            terminal_value_coefficient,
            atol=1.0e-12,
        ),
    )
    check(
        "run_summary all_checks_passed=true",
        summary.get("validation", {}).get("all_checks_passed") is True,
    )

    interval_rows = read_csv(out / "q2_interval_details.csv")
    detail: dict[tuple[str, int], dict[str, str]] = {
        (row["date"], int(row["t"])): row for row in interval_rows
    }
    dates = sorted({day for day, _ in detail})
    expected_dates = [
        (date(2025, 2, 1) + timedelta(days=i)).isoformat() for i in range(334)
    ]
    check("明细48096行=334天×144", len(detail) == 48096)
    check(
        f"明细日期连续 {dates[0]}..{dates[-1]}",
        dates == expected_dates,
    )
    day_rows = {day: [detail[(day, t)] for t in range(1, 145)] for day in dates}
    daily = read_csv(out / "q2_daily_metrics.csv")
    dm = {row["date"]: row for row in daily}

    # ---------- 3 逐时段残差、SOC、互斥 ----------
    max_dp_bal = max_ana_bal = max_dp_soc = max_ana_soc = 0.0
    max_charge = max_discharge = max_product = 0.0
    min_nonnegative_flow = float("inf")
    min_soc = float("inf")
    max_soc = float("-inf")
    for day in dates:
        rows = day_rows[day]
        for t in range(1, 145):
            row = rows[t - 1]
            net = f(row["actual_load_kwh"]) - f(row["actual_pv_kwh"])
            for prefix, ck, dk, bk, uk, sk in (
                ("dp", "dp_charge_kwh", "dp_discharge_kwh", "dp_emergency_kwh", "dp_unused_kwh", "dp_end_soc_kwh"),
                ("ana", "analytical_charge_kwh", "analytical_discharge_kwh", "analytical_emergency_kwh", "analytical_unused_kwh", "analytical_end_soc_kwh"),
            ):
                c = f(row[ck]); d = f(row[dk]); b = f(row[bk]); u = f(row[uk]); e = f(row[sk])
                plan = f(row["plan_grid_kwh"])
                balance = plan + b + d - net - c - u
                if t == 1:
                    initial_key = (
                        "dp_initial_soc_kwh" if prefix == "dp" else "analytical_initial_soc_kwh"
                    )
                    previous = f(dm[day][initial_key])
                else:
                    previous = f(rows[t - 2][sk])
                soc_residual = e - previous - 0.9 * c + d / 0.9
                if prefix == "dp":
                    max_dp_bal = max(max_dp_bal, abs(balance))
                    max_dp_soc = max(max_dp_soc, abs(soc_residual))
                else:
                    max_ana_bal = max(max_ana_bal, abs(balance))
                    max_ana_soc = max(max_ana_soc, abs(soc_residual))
                max_charge = max(max_charge, c, d)
                max_discharge = max(max_discharge, d)
                max_product = max(max_product, c * d)
                min_nonnegative_flow = min(
                    min_nonnegative_flow, plan, c, d, b, u
                )
                min_soc = min(min_soc, e)
                max_soc = max(max_soc, e)
    check(f"DP平衡残差<=1e-6 (max={max_dp_bal:.3g})", max_dp_bal <= TOL)
    check(f"解析平衡残差<=1e-6 (max={max_ana_bal:.3g})", max_ana_bal <= TOL)
    check(f"DP SOC残差<=1e-6 (max={max_dp_soc:.3g})", max_dp_soc <= TOL)
    check(f"解析SOC残差<=1e-6 (max={max_ana_soc:.3g})", max_ana_soc <= TOL)
    check(
        f"C/D<=833.3333且互斥 (max={max_charge:.6f}, prod={max_product:.3g})",
        max_charge <= 5000.0 / 6.0 + TOL and max_product <= 1.0e-8,
    )
    check(f"购电/充/放/紧急/弃用非负 (min={min_nonnegative_flow:.3g})", min_nonnegative_flow >= -TOL)
    check(f"SOC位于[1200,10800] (min={min_soc:.6f}, max={max_soc:.6f})", min_soc >= 1200.0 - TOL and max_soc <= 10800.0 + TOL)

    # ---------- 4 daily_metrics ----------
    check("daily_metrics 334行", len(daily) == 334 and sorted(dm) == expected_dates)
    identity_ok = True
    for day in dates:
        row = dm[day]
        if abs(f(row["dp_total_cost_yuan"]) - (f(row["plan_cost_yuan"]) + f(row["dp_emergency_cost_yuan"]))) > TOL:
            identity_ok = False
            break
        if abs(f(row["analytical_total_cost_yuan"]) - (f(row["plan_cost_yuan"]) + f(row["analytical_emergency_cost_yuan"]))) > TOL:
            identity_ok = False
            break
        part = day_rows[day]
        if abs(sum(f(x["dp_charge_kwh"]) for x in part) - f(row["dp_charge_kwh"])) > 1.0e-6:
            identity_ok = False
            break
        if abs(sum(f(x["dp_discharge_kwh"]) for x in part) - f(row["dp_discharge_kwh"])) > 1.0e-6:
            identity_ok = False
            break
    check("334日费用恒等式与充放电汇总", identity_ok)
    cross_ok = True
    for i in range(1, len(dates)):
        if abs(f(dm[dates[i]]["dp_initial_soc_kwh"]) - f(dm[dates[i - 1]]["dp_terminal_soc_kwh"])) > TOL:
            cross_ok = False
            break
        if abs(f(dm[dates[i]]["analytical_initial_soc_kwh"]) - f(dm[dates[i - 1]]["analytical_terminal_soc_kwh"])) > TOL:
            cross_ok = False
            break
    check("DP/解析跨日SOC连续", cross_ok)
    check(
        "2025-02-01期初SOC=6000",
        abs(f(dm[dates[0]]["dp_initial_soc_kwh"]) - 6000.0) <= TOL
        and abs(f(dm[dates[0]]["analytical_initial_soc_kwh"]) - 6000.0) <= TOL,
    )

    # ---------- 5 计划表跨日映射 ----------
    wb = load_workbook(out / "result2.xlsx", data_only=True, read_only=True)
    plan_sheet = wb["计划购电量"]
    plan_rows = list(plan_sheet.iter_rows(min_row=2, max_row=335, values_only=True))
    max_map_diff = 0.0
    ep_eq_bad = None
    for row_index, values in enumerate(plan_rows, start=2):
        target = values[0].date().isoformat()
        displayed = np.asarray(values[1:145], dtype=np.float64)
        current = [f(x["plan_grid_kwh"]) for x in day_rows[target]]
        following = (date.fromisoformat(target) + timedelta(days=1)).isoformat()
        if following in day_rows:
            expected = np.concatenate(
                (np.asarray(current[1:]), [f(day_rows[following][0]["plan_grid_kwh"])])
            )
        else:
            expected = np.concatenate((np.asarray(current[1:]), [displayed[-1]]))
        max_map_diff = max(
            max_map_diff, float(np.max(np.abs(displayed - expected)))
        )
        ep = float(values[145])
        eq = float(values[146])
        if abs(ep - float(np.sum(displayed))) > TOL:
            ep_eq_bad = (target, "EP")
            break
        if abs(eq - float(np.dot(shifted_prices, displayed))) > TOL:
            ep_eq_bad = (target, "EQ")
            break
    check(
        f"计划表跨日映射(B:EN=t2..144,EO=次日t1) maxdiff={max_map_diff:.3g}",
        ep_eq_bad is None and max_map_diff <= TOL,
    )
    check(f"334行EP=列和、EQ=位移价点积 bad={ep_eq_bad}", ep_eq_bad is None)
    boundary_t1 = plan_rows[-1][144]
    check(
        f"12-31 EO格=边界计划t1={boundary_t1:.6f} 且有限非负",
        boundary_t1 is not None and boundary_t1 >= 0,
    )

    # ---------- 6 充放电量表（334天全量展开） ----------
    storage_sheet = wb["充放电量"]
    storage_rows = list(storage_sheet.iter_rows(min_row=2, values_only=True))
    storage_ok = len(storage_rows) == 334 * 6
    if storage_ok:
        for day_index, day in enumerate(dates):
            part = day_rows[day]
            base = day_index * 6
            if storage_rows[base][0] is None or storage_rows[base][0].date().isoformat() != day:
                storage_ok = False
                break
            for block in range(6):
                expected_c = sum(f(x["dp_charge_kwh"]) for x in part[24 * block : 24 * (block + 1)])
                expected_d = sum(f(x["dp_discharge_kwh"]) for x in part[24 * block : 24 * (block + 1)])
                if storage_rows[base + block][1] != ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"][block]:
                    storage_ok = False
                    break
                if abs(storage_rows[base + block][2] - expected_c) > TOL or abs(storage_rows[base + block][3] - expected_d) > TOL:
                    storage_ok = False
                    break
            if abs(storage_rows[base][5] - f(dm[day]["dp_initial_soc_kwh"])) > TOL:
                storage_ok = False
                break
            if abs(storage_rows[base + 1][5] - f(dm[day]["dp_terminal_soc_kwh"])) > TOL:
                storage_ok = False
                break
    check("充放电量表334天×6块+SOC全量一致", storage_ok)

    # ---------- 7 紧急购电表（334天全量展开，无省略号） ----------
    emergency_sheet = wb["紧急购电量"]
    emergency_rows = list(emergency_sheet.iter_rows(min_row=2, values_only=True))
    cursor = 0
    emergency_ok = True
    days_with_purchase = 0
    for day in dates:
        part = day_rows[day]
        intervals: list[tuple[str, float]] = []
        start_t: int | None = None
        for t in range(1, 145):
            value = f(part[t - 1]["dp_emergency_kwh"])
            if value > 1.0e-7 and start_t is None:
                start_t = t
            if value <= 1.0e-7 and start_t is not None:
                start_minutes = (start_t - 1) * 10
                end_minutes = (t - 1) * 10
                label_start = f"{start_minutes // 60}:{start_minutes % 60:02d}"
                label_end = "00:00+1" if end_minutes == 1440 else f"{end_minutes // 60}:{end_minutes % 60:02d}"
                intervals.append(
                    (f"{label_start}-{label_end}", sum(f(x["dp_emergency_kwh"]) for x in part[start_t - 1 : t - 1]))
                )
                start_t = None
        if start_t is not None:
            start_minutes = (start_t - 1) * 10
            intervals.append(
                (f"{start_minutes // 60}:{start_minutes % 60:02d}-00:00+1",
                 sum(f(x["dp_emergency_kwh"]) for x in part[start_t - 1 :]))
            )
        if not intervals:
            intervals = [("无", 0.0)]
        else:
            days_with_purchase += 1
        for interval_index, (expected_label, expected_value) in enumerate(intervals):
            if cursor >= len(emergency_rows):
                emergency_ok = False
                break
            row = emergency_rows[cursor]
            if interval_index == 0 and (row[0] is None or row[0].date().isoformat() != day):
                emergency_ok = False
                break
            if interval_index > 0 and row[0] is not None:
                emergency_ok = False
                break
            if row[1] != expected_label or abs(float(row[2]) - expected_value) > TOL:
                emergency_ok = False
                break
            cursor += 1
        if not emergency_ok:
            break
    if emergency_ok and cursor != len(emergency_rows):
        emergency_ok = False
    check(
        f"紧急购电表334天逐行一致(有购电{days_with_purchase}天)",
        emergency_ok,
    )
    wb.close()

    # ---------- 8 solver_audit ----------
    solver = read_csv(out / "q2_solver_audit.csv")
    solver_dates = [row["target_date"] for row in solver]
    check(
        "solver_audit 335行且含2026-01-01边界",
        len(solver) == 335 and solver_dates[-1] == "2026-01-01",
    )
    check("全部LP为Optimal", all("Optimal" in row["status"] for row in solver))
    check(
        "LP残差全部<=1e-6",
        all(f(row["max_balance_residual_kwh"]) <= TOL and f(row["max_soc_residual_kwh"]) <= TOL for row in solver),
    )
    check(
        "无同时充放电",
        all(f(row["max_simultaneous_product_kwh2"]) <= 1.0e-8 for row in solver),
    )

    # ---------- 9 场景审计因果性 ----------
    audit = read_csv(out / "q2_scenario_audit.csv")
    causal_bad = 0
    for row in audit:
        if date.fromisoformat(row["history_date"]) >= date.fromisoformat(row["target_date"]):
            causal_bad += 1
            break
    check("场景审计所有历史日期<目标日(因果)", causal_bad == 0)
    groups: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    for row in audit:
        groups[(row["target_date"], row["record_type"])].append(row)
    distance_ok = True
    for target in ("2025-12-29", "2025-12-30", "2025-12-31", "2026-01-01"):
        distances = sorted(int(row["distance_days"]) for row in groups[(target, "pv_primary_recent_day")])
        if distances != [1, 2, 3]:
            distance_ok = False
            break
    check("光伏主预测=最近3自然日(距离1,2,3)", distance_ok)

    def weekday_class(day: str) -> str:
        return "low" if date.fromisoformat(day).weekday() in (4, 5) else "regular"

    class_ok = True
    for day in dates:
        rows = groups[(day, "load_primary_same_class")]
        if rows and any(row["load_class"] != weekday_class(day) for row in rows):
            class_ok = False
            break
    check("负荷主预测历史与目标同类", class_ok)
    weight_ok = True
    for (_, record_type), rows in groups.items():
        if not rows:
            continue
        if abs(sum(f(row["weight"]) for row in rows) - 1.0) > 1.0e-9:
            weight_ok = False
            break
    check("审计各组权重和=1", weight_ok)

    # ---------- 10 预测指标抽检 ----------
    forecast_metrics = read_csv(out / "q2_forecast_metrics.csv")
    forecast_ok = True
    for day in ("2025-03-15", "2025-07-01", "2025-12-20"):
        part = day_rows[day]
        row = [x for x in forecast_metrics if x["date"] == day and x["method"] == "main_weighted_plus_30day_bias"][0]
        forecast_load = [f(x["forecast_load_kwh"]) for x in part]
        actual_load = [f(x["actual_load_kwh"]) for x in part]
        mae = float(np.mean(np.abs(np.asarray(forecast_load) - np.asarray(actual_load))))
        rmse = float(np.sqrt(np.mean((np.asarray(forecast_load) - np.asarray(actual_load)) ** 2)))
        if abs(mae - f(row["load_mae_kwh"])) > 1.0e-9 or abs(rmse - f(row["load_rmse_kwh"])) > 1.0e-9:
            forecast_ok = False
            break
    check("主预测MAE/RMSE抽样复算", forecast_ok)

    # ---------- 11 模板非目标单元格未变 ----------
    template = load_workbook(root / "data/附件/附件5/result2.xlsx", data_only=False)
    output = load_workbook(out / "result2.xlsx", data_only=False)
    mismatch = 0
    plan_template = template["计划购电量"]
    plan_output = output["计划购电量"]
    for row in range(1, 336):
        for column in range(1, 148):
            if row >= 2 and 2 <= column <= 147:
                continue
            if (plan_template.cell(row, column).value != plan_output.cell(row, column).value
                    or plan_template.cell(row, column).style_id != plan_output.cell(row, column).style_id):
                mismatch += 1
    # 充放电量/紧急购电量已按334天全量重建（尺寸必然大于模板），仅校验表头行不变。
    for sheet_name, max_column in (("充放电量", 6), ("紧急购电量", 3)):
        template_sheet = template[sheet_name]
        output_sheet = output[sheet_name]
        for column in range(1, max_column + 1):
            if (template_sheet.cell(1, column).value != output_sheet.cell(1, column).value
                    or template_sheet.cell(1, column).style_id != output_sheet.cell(1, column).style_id):
                mismatch += 1
    template.close()
    output.close()
    check("模板非目标单元格与样式未变(充放电量/紧急表已全量展开)", mismatch == 0, f"mismatch={mismatch}")

    # ---------- 12 全年费用 ----------
    plan_cost = sum(f(row["plan_cost_yuan"]) for row in daily)
    dp_emg = sum(f(row["dp_emergency_cost_yuan"]) for row in daily)
    ana_emg = sum(f(row["analytical_emergency_cost_yuan"]) for row in daily)
    dp_total = plan_cost + dp_emg
    ana_total = plan_cost + ana_emg
    print("\n全年费用汇总：")
    print(f"  计划费 = {plan_cost:.2f} 元")
    print(f"  DP紧急费 = {dp_emg:.2f} 元, DP总 = {dp_total:.2f} 元")
    print(f"  解析紧急费 = {ana_emg:.2f} 元, 解析总 = {ana_total:.2f} 元")
    print(f"  DP节省 = {ana_emg - dp_emg:.2f} 元 ({(ana_emg - dp_emg) / ana_total * 100:.4f}%)")
    print(f"  实际费用未含终值抵扣：计划+紧急=总成立（费用恒等式检查项）")

    fails = [item for item in results if not item[1]]
    print(f"\n总结：{len(results) - len(fails)}/{len(results)} 项通过")
    if fails:
        print("未通过项：")
        for name, _, detail in fails:
            print(f"  - {name} {detail}")


if __name__ == "__main__":
    main()
